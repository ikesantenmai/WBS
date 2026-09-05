import datetime as dt
import io
import zipfile

import openpyxl
import pytest

pytest.importorskip("fastapi", reason="Web アプリの依存 (pip install 'wbsgen[web]')")
from fastapi.testclient import TestClient  # noqa: E402

from wbsgen.web.app import app  # noqa: E402

D = dt.date


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


SPEC = {"start": "2026-02-01", "end": "2027-01-31", "unit": "week", "rows": 40}


# ---------------------------------------------------------------- 画面
def test_index_serves_the_app_shell(client):
    response = client.get("/")
    assert response.status_code == 200
    assert "WBS ジェネレータ" in response.text
    assert "/static/app.js" in response.text


def test_the_assets_carry_a_stamp(client):
    """入れ替えたあとに前の版が使われないよう、読み込みに印を付ける。"""
    import re

    text = client.get("/").text
    assert re.search(r"/static/app\.js\?v=[0-9a-f]{8}", text)
    assert re.search(r"/static/style\.css\?v=[0-9a-f]{8}", text)


def test_the_stamp_follows_the_contents(client):
    """印は中身から作る。版番号だと、上げ忘れたときに切り替わらない。"""
    import re

    from wbsgen.web.app import STATIC_DIR

    def stamp():
        found = re.search(r"/static/app\.js\?v=([0-9a-f]{8})", client.get("/").text)
        return found.group(1)

    path = STATIC_DIR / "app.js"
    before = path.read_bytes()
    was = stamp()
    try:
        path.write_bytes(before + b"\n// changed\n")
        assert stamp() != was
    finally:
        path.write_bytes(before)
    assert stamp() == was


def test_the_page_itself_is_never_cached(client):
    """版を書き換えるのはこのページなので、毎回取り直させる。"""
    assert "no-store" in client.get("/").headers["cache-control"]


@pytest.mark.parametrize("path", ["/static/app.js", "/static/style.css"])
def test_static_assets_are_served(client, path):
    assert client.get(path).status_code == 200


def test_meta_offers_the_choices(client):
    body = client.get("/api/meta").json()
    assert [u["value"] for u in body["units"]] == ["day", "week", "month"]
    assert [w["label"] for w in body["weekdays"]] == list("月火水木金土日")
    assert body["default_rows"] == 40
    # 既定は年度 (4 月始まり)
    assert body["suggested"]["start"].endswith("-04-01")
    assert body["suggested"]["end"].endswith("-03-31")
    # 動いている版 (入れ替えたかどうかを画面で確かめられるように返す)
    from wbsgen import __version__
    assert body["version"] == __version__


# ---------------------------------------------------------------- preview
def test_preview_returns_the_timeline(client):
    body = client.post("/api/preview", json=SPEC).json()
    assert body["spec"]["period_days"] == 365
    assert body["spec"]["end"] == "2027-01-31"
    timeline = body["timeline"]
    assert len(timeline["columns"]) == 53
    assert [c["label"] for c in timeline["columns"][:3]] == ["2/1", "2/8", "2/15"]
    assert [b["label"] for b in timeline["bands"]][:3] == ["2月", "3月", "4月"]
    assert [b["start"] for b in timeline["bands"]][:3] == [0, 4, 9]


def test_preview_marks_rest_columns_in_day_view(client):
    body = client.post("/api/preview",
                       json={**SPEC, "start": "2026-04-01", "months": 1,
                             "end": None, "unit": "day"}).json()
    columns = {c["start"]: c for c in body["timeline"]["columns"]}
    assert columns["2026-04-04"]["rest"] and columns["2026-04-04"]["saturday"]
    assert columns["2026-04-29"]["rest"]          # 昭和の日
    assert not columns["2026-04-30"]["rest"]
    assert columns["2026-04-01"]["weekday"] == "水"


def test_preview_lists_the_holidays(client):
    body = client.post("/api/preview", json=SPEC).json()
    assert "2026-02-11" in body["holidays"]
    assert len(body["holidays"]) == 18


def test_preview_honours_the_calendar_settings(client):
    body = client.post("/api/preview", json={
        **SPEC, "japanese_holidays": False, "holidays": ["2026-12-30"],
    }).json()
    assert body["holidays"] == ["2026-12-30"]


@pytest.mark.parametrize("payload,message", [
    ({"start": "bad"}, "start"),
    ({"start": "2026-04-01", "end": "2020-01-01"}, "終了日"),
    ({"start": "2026-04-01", "unit": "hour"}, "表示単位"),
    ({"start": "2026-04-01", "rows": 99999}, "行数"),
    ({"start": "2026-04-01", "workdays": ["someday"]}, "曜日"),
])
def test_preview_rejects_bad_input(client, payload, message):
    response = client.post("/api/preview", json=payload)
    assert response.status_code == 422
    assert message in response.json()["detail"]


# ---------------------------------------------------------------- build
def test_build_returns_a_real_workbook(client):
    response = client.post("/api/build", json={**SPEC, "members": ["設計"]})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml")

    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["スケジュール", "担当者一覧", "設定"]
    ws = book["スケジュール"]
    assert ws["S3"].value == dt.datetime(2026, 2, 1)
    assert ws["S4"].value == dt.datetime(2026, 2, 1)
    assert ws.max_row == 44
    assert book["担当者一覧"]["B4"].value == "設計"


def test_build_has_no_drawings(client):
    """空の WBS なので図形は入らない。"""
    response = client.post("/api/build", json=SPEC)
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        assert not [n for n in zf.namelist() if "drawings" in n]


def test_build_filename_carries_the_title(client):
    response = client.post("/api/build", json={**SPEC, "title": "年度計画/2026"})
    disposition = response.headers["content-disposition"]
    assert "filename*=UTF-8''" in disposition
    assert "%2F" not in disposition             # 使えない文字は落としてある
    assert "20260201" in disposition


def test_build_rejects_bad_input(client):
    assert client.post("/api/build", json={"start": "bad"}).status_code == 422


# ---------------------------------------------------------------- import
def _upload(client, path, name="filled.xlsx", params=None):
    with open(path, "rb") as handle:
        return client.post("/api/import", params=params or {}, files={
            "file": (name, handle.read(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        })


def test_import_returns_a_chart(client, filled_book):
    body = _upload(client, filled_book).json()
    assert body["title"] == "2026年度 開発スケジュール"
    assert len(body["rows"]) == 5
    assert body["timeline"]["unit"] == "week"
    assert body["totals"]["done"] == 2
    assert body["rows"][0]["plan"]["x2"] > body["rows"][0]["plan"]["x1"]


def test_import_returns_the_workload_check(client, filled_book):
    """要員稼働チェックも一緒に返す (画面で切り替えて見られるように)。"""
    body = _upload(client, filled_book).json()
    months = body["workload"]
    assert [m["label"] for m in months] == ["要員稼働チェック_4月",
                                            "要員稼働チェック_5月",
                                            "要員稼働チェック_6月",
                                            "要員稼働チェック_7月",
                                            "要員稼働チェック_8月"]
    april = months[0]
    assert april["title"].startswith("◆要員稼働チェック（2026/4/1")
    assert april["days"][0] == {"date": "2026-04-01", "label": "4/1",
                                "weekday": "水", "working": True}
    assert april["days"][4]["working"] is False          # 4/5 は日曜
    assert april["workdays"] == 21                       # 昭和の日 4/29 を除く
    names = [m["name"] for m in april["members"]]
    assert names == ["設計", "製造", "テスト"]
    design = april["members"][0]
    assert design["counts"][0] == 1
    assert design["busy"] + design["free"] == april["workdays"]
    # 空き日があるときだけ、その日付を並べる
    for member in april["members"]:
        assert bool(member["free_days"]) == (member["free"] > 0)
    assert any("4/" in member["free_days"] for member in april["members"])


def test_the_workload_is_empty_without_owners(client, make_filled):
    path = make_filled("noowner.xlsx", rows=[
        ("開発", "", "1", "A", D(2026, 4, 1), 10, D(2026, 4, 14),
         None, None, None, None, None, "", "", ""),
    ])
    assert _upload(client, path).json()["workload"] == []


@pytest.mark.parametrize("unit,columns", [("day", 365), ("week", 53), ("month", 12)])
def test_import_can_override_the_unit(client, filled_book, unit, columns):
    body = _upload(client, filled_book, params={"unit": unit}).json()
    assert body["timeline"]["unit"] == unit
    assert len(body["timeline"]["columns"]) == columns


def test_import_rejects_an_unknown_unit(client, filled_book):
    response = _upload(client, filled_book, params={"unit": "hour"})
    assert response.status_code == 422
    assert "表示単位" in response.json()["detail"]


def test_import_rejects_a_non_excel_file(client):
    response = client.post("/api/import", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert response.status_code == 415
    assert "対応していない形式" in response.json()["detail"]


def test_import_rejects_a_broken_file(client):
    response = client.post("/api/import", files={"file": ("broken.xlsx", b"nope", "x")})
    assert response.status_code == 422
    assert "Excel として読めません" in response.json()["detail"]


def test_import_rejects_an_empty_workbook(client, tmp_path):
    """記入されていない空の WBS を読ませたら、その旨を返す。"""
    from wbsgen.blank import build as build_spec
    from wbsgen.workbook import write as write_book

    path = tmp_path / "blank.xlsx"
    write_book(build_spec(dt.date(2026, 4, 1), months=3, rows=10), path)
    response = _upload(client, path, name="blank.xlsx")
    assert response.status_code == 422
    assert "記入された行が見つかりません" in response.json()["detail"]


def test_import_reports_unreadable_cells_without_dropping_rows(client, make_filled):
    import openpyxl

    path = make_filled("bad.xlsx")
    book = openpyxl.load_workbook(path)
    book["スケジュール"]["F6"] = "来週くらい"
    book.save(path)

    body = _upload(client, path, name="bad.xlsx").json()
    assert len(body["rows"]) == 5                  # 行は消えない
    assert body["rows"][1]["start"] is None        # 読めなかった項目だけ空
    assert any("6 行目" in w and "開始日" in w for w in body["warnings"])


# ---------------------------------------------------------------- export
def test_export_returns_a_workbook_with_the_chart(client, filled_book):
    with open(filled_book, "rb") as handle:
        response = client.post("/api/export", files={"file": ("filled.xlsx", handle.read(), "x")})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml")

    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        assert "xl/drawings/drawing1.xml" in zf.namelist()   # ガントの図形
    ws = openpyxl.load_workbook(io.BytesIO(response.content))["スケジュール"]
    assert ws["E5"].value == "要件定義"


def test_export_filename_marks_the_chart(client, filled_book):
    with open(filled_book, "rb") as handle:
        response = client.post("/api/export", files={"file": ("filled.xlsx", handle.read(), "x")})
    disposition = response.headers["content-disposition"]
    assert "filename*=UTF-8''" in disposition
    assert "%E3%82%AC%E3%83%B3%E3%83%88" in disposition      # 「ガント」


@pytest.mark.parametrize("unit", ["day", "week", "month"])
def test_export_follows_the_requested_unit(client, filled_book, unit):
    with open(filled_book, "rb") as handle:
        response = client.post("/api/export", params={"unit": unit},
                               files={"file": ("filled.xlsx", handle.read(), "x")})
    assert response.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(response.content))["スケジュール"]
    assert ws["S4"].number_format == {"day": "d", "week": "m/d", "month": 'm"月"'}[unit]


def test_export_rejects_a_broken_file(client):
    response = client.post("/api/export", files={"file": ("broken.xlsx", b"nope", "x")})
    assert response.status_code == 422


def test_export_rejects_an_empty_workbook(client, tmp_path):
    from wbsgen.blank import build as build_spec
    from wbsgen.workbook import write as write_book

    path = tmp_path / "blank.xlsx"
    write_book(build_spec(dt.date(2026, 4, 1), months=3, rows=10), path)
    with open(path, "rb") as handle:
        response = client.post("/api/export", files={"file": ("blank.xlsx", handle.read(), "x")})
    assert response.status_code == 422
    assert "記入された行が見つかりません" in response.json()["detail"]


# ---------------------------------------------------------------- 言語
def test_meta_is_japanese_by_default(client):
    body = client.get("/api/meta").json()
    assert body["language"] == "ja"
    assert [u["label"] for u in body["units"]] == ["日単位", "週単位", "月単位"]
    assert [w["label"] for w in body["weekdays"]][:3] == ["月", "火", "水"]
    assert [l["value"] for l in body["languages"]] == ["ja", "en"]


def test_meta_can_be_english(client):
    body = client.get("/api/meta", params={"lang": "en"}).json()
    assert body["language"] == "en"
    assert [u["label"] for u in body["units"]] == ["Daily", "Weekly", "Monthly"]
    assert [w["label"] for w in body["weekdays"]][:3] == ["Mon", "Tue", "Wed"]
    assert body["suggested"]["title"].startswith("FY")


def test_an_unknown_language_falls_back_to_japanese(client):
    assert client.get("/api/meta", params={"lang": "fr"}).json()["language"] == "ja"


@pytest.mark.parametrize("params,months", [
    ({}, ["4月", "5月"]),
    ({"lang": "en"}, ["Apr", "May"]),
])
def test_preview_headings_follow_the_language(client, params, months):
    body = client.post("/api/preview", params=params,
                       json={"start": "2026-04-01", "months": 3}).json()
    assert [b["label"] for b in body["timeline"]["bands"]][:2] == months


def test_preview_weekdays_follow_the_language(client):
    body = client.post("/api/preview", params={"lang": "en"},
                       json={"start": "2026-04-01", "months": 1, "unit": "day"}).json()
    assert body["timeline"]["columns"][0]["weekday"] == "Wed"


@pytest.mark.parametrize("params,fragment", [
    ({}, "終了日が開始日より前"),
    ({"lang": "en"}, "end date is before"),
])
def test_api_errors_follow_the_language(client, params, fragment):
    response = client.post("/api/preview", params=params,
                           json={"start": "2026-04-01", "end": "2020-01-01"})
    assert response.status_code == 422
    assert fragment in response.json()["detail"]


@pytest.mark.parametrize("params,sheets", [
    ({}, ["スケジュール", "担当者一覧", "設定"]),
    ({"lang": "en"}, ["Schedule", "Members", "Settings"]),
])
def test_build_writes_the_requested_language(client, params, sheets):
    response = client.post("/api/build", params=params, json=SPEC)
    assert openpyxl.load_workbook(io.BytesIO(response.content)).sheetnames == sheets


def test_the_language_can_come_from_the_body(client):
    response = client.post("/api/build", json={**SPEC, "language": "en"})
    assert openpyxl.load_workbook(io.BytesIO(response.content)).sheetnames[0] == "Schedule"


@pytest.mark.parametrize("lang", ["ja", "en"])
def test_import_works_in_either_language(client, filled_book, lang):
    with open(filled_book, "rb") as handle:
        body = client.post("/api/import", params={"lang": lang},
                           files={"file": ("filled.xlsx", handle.read(), "x")}).json()
    assert body["language"] == lang
    assert len(body["rows"]) == 5


@pytest.mark.parametrize("lang,marker", [("ja", "%E3%82%AC%E3%83%B3%E3%83%88"), ("en", "gantt")])
def test_export_filename_follows_the_language(client, filled_book, lang, marker):
    with open(filled_book, "rb") as handle:
        response = client.post("/api/export", params={"lang": lang},
                               files={"file": ("filled.xlsx", handle.read(), "x")})
    assert marker in response.headers["content-disposition"]
