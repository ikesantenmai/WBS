"""要員稼働チェックの数え方のテスト。

添付いただいた元表 (要員稼働チェック_9月) と同じ数え方になることを確かめる。
"""

import datetime as dt

import openpyxl
import pytest

from wbsgen.importer import read
from wbsgen.workcal import WorkCalendar
from wbsgen.workload import (MAX_MONTHS, build, member_names, split_names)
from wbsgen.workbook import SHEET_PLAN, export

D = dt.date
BASE = D(2026, 6, 10)


class _Row:
    """稼働チェックが見るぶんだけの行。"""

    def __init__(self, start=None, end=None, member=""):
        self.start, self.end, self.member = start, end, member


def _calendar(**kwargs):
    return WorkCalendar(**kwargs)


# ---------------------------------------------------------------- 名前の切り出し
@pytest.mark.parametrize("text, expected", [
    ("高瀬", ["高瀬"]),
    ("吉田/菊池/虎岩", ["吉田", "菊池", "虎岩"]),
    ("長見、本家、池田、増田", ["長見", "本家", "池田", "増田"]),
    ("佐々木（高瀬）", ["佐々木", "高瀬"]),
    ("木村(吉田)", ["木村", "吉田"]),
    ("鈴木さん、塩谷さん", ["鈴木", "塩谷"]),
    ("SBI/DS", ["SBI", "DS"]),
    ("", []),
    ("  ", []),
])
def test_names_are_split_out_of_the_owner_cell(text, expected):
    assert split_names(text) == expected


def test_a_long_sentence_is_not_taken_as_a_name():
    """名前にしては長すぎるものは拾わない。"""
    assert split_names("担当は追って調整のうえ別途連絡します" * 2) == []


def test_names_keep_the_order_they_first_appear(make_filled=None):
    rows = [_Row(member="吉田/菊池"), _Row(member="高瀬"), _Row(member="菊池")]
    assert member_names(rows) == ["吉田", "菊池", "高瀬"]


# ---------------------------------------------------------------- 数え方
def test_counts_the_tasks_that_cover_each_day():
    rows = [
        _Row(D(2026, 6, 1), D(2026, 6, 3), "高瀬"),
        _Row(D(2026, 6, 2), D(2026, 6, 2), "高瀬"),
        _Row(D(2026, 6, 2), D(2026, 6, 5), "吉田"),
    ]
    month = build(rows, _calendar())[0]
    takase, yoshida = month.members
    assert takase.name == "高瀬"
    # 6/1 は 1 件、6/2 は 2 件、6/3 は 1 件、6/4 以降は 0 件
    assert takase.counts[:5] == [1, 2, 1, 0, 0]
    assert yoshida.counts[:5] == [0, 1, 1, 1, 1]


def test_a_name_inside_another_owner_cell_counts_too():
    """「佐々木（高瀬）」は佐々木にも高瀬にも数える (元表と同じ)。"""
    rows = [_Row(D(2026, 6, 1), D(2026, 6, 1), "佐々木（高瀬）")]
    month = build(rows, _calendar())[0]
    assert [m.name for m in month.members] == ["佐々木", "高瀬"]
    assert [m.counts[0] for m in month.members] == [1, 1]


def test_days_off_are_marked_and_never_counted_as_busy():
    """土日はタスクがあっても「タスク有日数」に数えない。"""
    rows = [_Row(D(2026, 6, 1), D(2026, 6, 30), "高瀬")]      # 月末まで通し
    month = build(rows, _calendar())[0]
    load = month.members[0]
    assert month.workdays[5] is False and month.workdays[6] is False   # 6/6 土, 6/7 日
    assert load.busy_days == month.workday_count
    assert load.free_days == []


def test_free_days_are_the_working_days_with_no_task():
    rows = [
        _Row(D(2026, 6, 1), D(2026, 6, 3), "高瀬"),
        _Row(D(2026, 6, 1), D(2026, 6, 30), "吉田"),
    ]
    month = build(rows, _calendar())[0]
    takase = month.members[0]
    assert takase.busy_days == 3
    assert takase.free_days[0] == D(2026, 6, 4)
    assert len(takase.free_days) == month.workday_count - 3
    assert takase.busy_days + len(takase.free_days) == month.workday_count


def test_extra_days_off_come_from_the_calendar():
    rows = [_Row(D(2026, 6, 1), D(2026, 6, 5), "高瀬")]
    plain = build(rows, _calendar())[0]
    with_holiday = build(rows, _calendar(holidays=[D(2026, 6, 2)]))[0]
    assert with_holiday.workday_count == plain.workday_count - 1
    assert with_holiday.workdays[1] is False


# ---------------------------------------------------------------- 月の分け方
def test_one_sheet_per_month_clipped_to_the_plan():
    rows = [_Row(D(2026, 6, 20), D(2026, 8, 10), "高瀬")]
    months = build(rows, _calendar())
    assert [(m.year, m.month) for m in months] == [(2026, 6), (2026, 7), (2026, 8)]
    assert (months[0].start, months[0].end) == (D(2026, 6, 20), D(2026, 6, 30))
    assert (months[1].start, months[1].end) == (D(2026, 7, 1), D(2026, 7, 31))
    assert (months[2].start, months[2].end) == (D(2026, 8, 1), D(2026, 8, 10))


def test_months_without_anyone_assigned_are_still_listed():
    """担当の無い行しかない月も並べる (空いていることが判るように)。"""
    rows = [
        _Row(D(2026, 6, 1), D(2026, 6, 5), "高瀬"),
        _Row(D(2026, 8, 1), D(2026, 8, 5), ""),      # 担当なし
    ]
    months = build(rows, _calendar())
    assert [m.month for m in months] == [6, 7, 8]
    assert all(count == 0 for count in months[2].members[0].counts)


def test_a_very_long_plan_is_capped():
    rows = [_Row(D(2026, 1, 1), D(2032, 12, 31), "高瀬")]
    assert len(build(rows, _calendar())) == MAX_MONTHS


def test_nothing_is_built_without_owners_or_dates():
    assert build([_Row(D(2026, 6, 1), D(2026, 6, 2), "")], _calendar()) == []
    assert build([_Row(None, None, "高瀬")], _calendar()) == []
    assert build([], _calendar()) == []


# ---------------------------------------------------------------- 書き出し
def _filled(make_filled, name="load.xlsx"):
    return make_filled(name, rows=[
        ("開発", "", "1", "要件定義", D(2026, 4, 1), 10, D(2026, 4, 14),
         None, None, None, None, None, "", "高瀬", ""),
        ("開発", "", "2", "基本設計", D(2026, 4, 8), 10, D(2026, 4, 21),
         None, None, None, None, None, "", "高瀬/吉田", ""),
        ("開発", "", "3", "詳細設計", D(2026, 5, 1), 10, D(2026, 5, 14),
         None, None, None, None, None, "", "吉田", ""),
    ])


def test_a_sheet_is_added_for_every_month(make_filled, tmp_path):
    out = export(read(_filled(make_filled), base_date=BASE),
                 tmp_path / "out.xlsx", BASE)
    names = openpyxl.load_workbook(out).sheetnames
    assert names == [SHEET_PLAN, "担当者一覧", "設定",
                     "要員稼働チェック_4月", "要員稼働チェック_5月"]


def test_the_sheet_matches_the_layout_of_the_original(make_filled, tmp_path):
    out = export(read(_filled(make_filled), base_date=BASE),
                 tmp_path / "out.xlsx", BASE)
    sheet = openpyxl.load_workbook(out)["要員稼働チェック_4月"]

    assert sheet["B1"].value == "◆要員稼働チェック（2026/4/1〜4/30）"
    assert (sheet["B4"].value, sheet["B5"].value, sheet["B6"].value) == \
        ("日付", "曜日", "稼働判定")
    assert sheet["C4"].value.date() == D(2026, 4, 1)
    assert sheet["C4"].number_format == "m/d"
    assert (sheet["C5"].value, sheet["C6"].value) == ("水", "稼")
    assert (sheet["G5"].value, sheet["G6"].value) == ("日", "休")   # 4/5 は日曜
    assert [sheet.cell(row=r, column=2).value for r in (7, 8)] == ["高瀬", "吉田"]
    assert [sheet.cell(row=4, column=c).value for c in range(34, 38)] == \
        ["稼働日数", "タスク有日数", "空き日数", "空き日の内訳"]
    assert "凡例" in sheet.cell(row=10, column=2).value


def test_the_counts_and_totals_are_written(make_filled, tmp_path):
    out = export(read(_filled(make_filled), base_date=BASE),
                 tmp_path / "out.xlsx", BASE)
    sheet = openpyxl.load_workbook(out)["要員稼働チェック_4月"]

    # 4/1〜4/7 は高瀬だけ 1 件、4/8 からは 2 件 (2 つのタスクが重なる)
    assert [sheet.cell(row=7, column=c).value for c in range(3, 11)] == \
        [1, 1, 1, 1, 1, 1, 1, 2]
    assert [sheet.cell(row=8, column=c).value for c in range(3, 11)] == \
        [0, 0, 0, 0, 0, 0, 0, 1]

    # 2026 年 4 月の稼働日は 21 日 (平日 22 日 − 昭和の日 4/29)
    workdays, busy, free = (sheet.cell(row=7, column=c).value for c in (34, 35, 36))
    assert workdays == 21
    assert busy + free == workdays
    assert busy == 15                       # 高瀬は 4/1〜4/21 の稼働日ぶん
    assert "4/22(水)" in sheet.cell(row=7, column=37).value

    # 吉田は 4/8 から。4/1 は空き
    assert sheet.cell(row=8, column=36).value == \
        workdays - sheet.cell(row=8, column=35).value
    assert sheet.cell(row=8, column=37).value.startswith("4/1(水)")


def test_the_colours_are_conditional_formatting(make_filled, tmp_path):
    out = export(read(_filled(make_filled), base_date=BASE),
                 tmp_path / "out.xlsx", BASE)
    sheet = openpyxl.load_workbook(out)["要員稼働チェック_4月"]
    colours = {
        rule.dxf.fill.bgColor.rgb[-6:]
        for rules in sheet.conditional_formatting._cf_rules.values()
        for rule in rules
    }
    # 灰=非稼働 / 赤=0 件 / 橙=3 件以上 / 緑=1〜2 件
    assert {"F2F2F2", "FF7C80", "FFC000", "C6E0B4"} <= colours


def test_exporting_twice_does_not_pile_up_sheets(make_filled, tmp_path):
    once = export(read(_filled(make_filled), base_date=BASE),
                  tmp_path / "a.xlsx", BASE)
    twice = export(read(once, base_date=BASE), tmp_path / "b.xlsx", BASE)
    assert openpyxl.load_workbook(twice).sheetnames == \
        openpyxl.load_workbook(once).sheetnames


def test_the_sheet_names_carry_the_year_when_months_repeat(make_filled, tmp_path):
    """1 年を超える計画では、同じ月が 2 度出るので年を付ける。"""
    path = make_filled("long.xlsx", start=D(2026, 4, 1), end=D(2027, 9, 30), rows=[
        ("開発", "", "1", "長期", D(2026, 4, 1), None, D(2027, 6, 30),
         None, None, None, None, None, "", "高瀬", ""),
    ])
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)
    names = [n for n in openpyxl.load_workbook(out).sheetnames if "稼働" in n]
    assert names[0] == "要員稼働チェック_2026年4月"
    assert "要員稼働チェック_2027年4月" in names


def test_no_sheet_is_added_without_owners(make_filled, tmp_path):
    path = make_filled("none.xlsx", rows=[
        ("開発", "", "1", "担当なし", D(2026, 4, 1), 10, D(2026, 4, 14),
         None, None, None, None, None, "", "", ""),
    ])
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)
    assert openpyxl.load_workbook(out).sheetnames == \
        [SHEET_PLAN, "担当者一覧", "設定"]


def test_the_sheets_are_translated(make_filled, tmp_path):
    path = _filled(make_filled, "en.xlsx")
    imported = read(path, language="en", base_date=BASE)
    imported.spec.language = "en"
    out = export(imported, tmp_path / "out.xlsx", BASE)

    book = openpyxl.load_workbook(out)
    assert "Workload_04" in book.sheetnames
    sheet = book["Workload_04"]
    assert sheet["B4"].value == "Date"
    assert [sheet.cell(row=4, column=c).value for c in range(34, 38)] == \
        ["Working days", "Days with tasks", "Free days", "Which days are free"]
