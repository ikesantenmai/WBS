"""ブラウザ経由の疎通確認。

Playwright と Chromium がある環境でだけ動く。無い場合は skip する。
JavaScript が実際に動いてプレビューを描き、Excel を落とせることを確かめる。
"""

from __future__ import annotations

import contextlib
import os
import socket
import threading
import time
from pathlib import Path

import pytest

from conftest import D

pytest.importorskip("fastapi", reason="Web アプリの依存が必要")
sync_playwright = pytest.importorskip(
    "playwright.sync_api", reason="playwright が必要"
).sync_playwright


def _chromium_path():
    """環境が用意した Chromium を探す。"""
    explicit = os.environ.get("WBSGEN_CHROMIUM")
    if explicit and Path(explicit).exists():
        return explicit
    root = Path(os.environ.get("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers"))
    if root.is_dir():
        for candidate in sorted(root.glob("chromium-*/chrome-linux/chrome"), reverse=True):
            return str(candidate)
    return None


@pytest.fixture(scope="module")
def server():
    import uvicorn

    from wbsgen.web.app import app

    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]

    instance = uvicorn.Server(
        uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
    thread = threading.Thread(target=instance.run, daemon=True)
    thread.start()
    for _ in range(100):
        if instance.started:
            break
        time.sleep(0.05)
    else:
        pytest.skip("テスト用サーバを起動できませんでした")

    yield f"http://127.0.0.1:{port}/"

    instance.should_exit = True
    thread.join(timeout=5)


@pytest.fixture(scope="module")
def browser():
    with sync_playwright() as playwright:
        launch = {"args": ["--no-sandbox"]}
        path = _chromium_path()
        if path:
            launch["executable_path"] = path
        try:
            instance = playwright.chromium.launch(**launch)
        except Exception as exc:  # noqa: BLE001 - ブラウザが無い環境では skip
            pytest.skip(f"Chromium を起動できません: {exc}")
        yield instance
        with contextlib.suppress(Exception):
            instance.close()


def _open(browser, server, **context_options):
    """指定した画面の大きさでアプリを開き、最初の描画まで待つ。"""
    context = browser.new_context(accept_downloads=True, **context_options)
    view = context.new_page()
    errors = []
    view.on("pageerror", lambda e: errors.append(str(e)))
    view.errors = errors
    view.goto(server, wait_until="networkidle")
    # 狭い画面では表が畳まれていることがあるので、出来ていることだけ確かめる
    view.wait_for_selector("table.wbs tbody tr", state="attached")
    return view


@pytest.fixture(scope="module")
def page(browser, server):
    return _open(browser, server, viewport={"width": 1440, "height": 900})


@pytest.fixture(scope="module")
def phone(browser, server):
    """スマートフォン相当の画面 (iPhone くらいの幅) で開いたページ。"""
    return _open(browser, server, viewport={"width": 390, "height": 844},
                 device_scale_factor=2, is_mobile=True, has_touch=True)


def _reset(page):
    """新規作成の既定 (年度・週単位) に戻し、プレビューが描き終わるまで待つ。"""
    if page.is_visible("#btn-back"):
        page.click("#btn-back")
        page.wait_for_selector("#spec-form:not([hidden])")
    page.select_option("#unit", "week")
    page.check("input[name=mode][value=end]")
    # 指定が通っていればダウンロードできる状態になる
    page.wait_for_function("() => !document.querySelector('#btn-build').disabled")
    page.wait_for_selector("table.wbs tbody tr")


# ----------------------------------------------------------------------
def _months(page):
    return [m for m in page.eval_on_selector_all(
        "svg.chart-head text[font-weight='700']", "n => n.map(x => x.textContent)") if m]


def test_preview_renders_the_sheet(page):
    _reset(page)
    assert _months(page)[:3] == ["4月", "5月", "6月"]
    assert page.eval_on_selector_all("table.wbs tbody tr", "n => n.length") > 10
    assert "空行" in page.text_content("#preview-info")
    assert page.errors == []


def test_month_field_is_hidden_until_selected(page):
    _reset(page)
    assert page.is_hidden("#field-months")
    page.check("input[name=mode][value=months]")
    page.wait_for_selector("#field-months:not([hidden])")
    assert page.is_hidden("#field-end")
    page.check("input[name=mode][value=end]")
    page.wait_for_selector("#field-end:not([hidden])")
    assert page.errors == []


def test_changing_the_unit_redraws_the_axis(page):
    _reset(page)
    page.select_option("#unit", "day")
    page.wait_for_selector("svg.chart-head[data-unit=day]")
    labels = page.eval_on_selector_all(
        "svg.chart-head text", "n => n.map(x => x.textContent)")
    assert "土" in labels and "日" in labels        # 曜日の行が出る

    page.select_option("#unit", "week")
    page.wait_for_selector("svg.chart-head[data-unit=week]")
    assert page.errors == []


def test_changing_the_row_count_changes_the_preview(page):
    _reset(page)
    page.fill("#rows", "7")
    page.dispatch_event("#rows", "change")
    page.wait_for_function(
        "() => document.querySelectorAll('table.wbs tbody tr').length === 7")
    assert "空行 7 行" in page.text_content("#preview-info")
    page.fill("#rows", "40")
    page.dispatch_event("#rows", "change")
    assert page.errors == []


def test_extra_holidays_are_counted(page):
    _reset(page)
    before = page.text_content("#holiday-note")
    page.fill("#holidays", "2026-12-30, 2026-12-31")
    page.wait_for_function(
        f"() => document.querySelector('#holiday-note').textContent !== {before!r}")
    assert "祝日・休業日" in page.text_content("#holiday-note")
    page.fill("#holidays", "")
    assert page.errors == []


def test_invalid_period_shows_a_message_and_blocks_download(page):
    _reset(page)
    page.fill("#end", "2020-01-01")
    page.dispatch_event("#end", "change")
    page.wait_for_selector("#banner:not([hidden])")
    assert "終了日" in page.text_content("#banner")
    assert page.is_disabled("#btn-build")

    page.fill("#end", "2027-03-31")
    page.dispatch_event("#end", "change")
    page.wait_for_function("() => !document.querySelector('#btn-build').disabled")
    assert page.errors == []


def test_downloading_the_workbook(page, tmp_path):
    _reset(page)
    with page.expect_download() as download:
        page.click("#btn-build")
    saved = tmp_path / "out.xlsx"
    download.value.save_as(saved)
    assert saved.stat().st_size > 8000
    assert download.value.suggested_filename.endswith(".xlsx")
    assert page.errors == []


# ---------------------------------------------------------------- 読み込み
def test_importing_a_workbook_draws_a_gantt_chart(page, filled_book):
    page.set_input_files("#import-file", str(filled_book))
    page.wait_for_function(
        "() => document.querySelectorAll('svg.chart-body rect[rx=\"2\"]').length > 0")

    # 表に記入内容が並ぶ
    first = page.eval_on_selector_all(
        "table.wbs tbody tr:first-child td", "n => n.map(x => x.textContent)")
    assert "要件定義" in first
    assert "完了" in first

    # チャートにバーと現在日線が描かれる
    assert page.eval_on_selector_all(
        "svg.chart-body rect[rx='2']", "n => n.length") >= 7
    assert page.eval_on_selector_all(
        "svg.chart-body line[stroke-dasharray]", "n => n.length") == 1

    # 集計が出て、指定フォームは畳まれる
    assert "行数" in page.text_content("#totals")
    assert page.is_hidden("#spec-form")
    assert page.errors == []


def test_the_chart_opens_in_day_units(page, filled_book):
    """読み込んだ WBS は、まず日単位で見せる。"""
    page.set_input_files("#import-file", str(filled_book))
    page.wait_for_selector("svg.chart-head[data-unit=day]")
    assert page.eval_on_selector("#chart-unit", "n => n.value") == "day"


def test_the_chart_unit_can_be_switched(page, filled_book):
    page.set_input_files("#import-file", str(filled_book))
    page.wait_for_selector("svg.chart-head[data-unit=day]")

    page.select_option("#chart-unit", "month")
    page.wait_for_selector("svg.chart-head[data-unit=month]")
    assert _months(page)[:2] == ["2026年", "2027年"]

    page.select_option("#chart-unit", "week")
    page.wait_for_selector("svg.chart-head[data-unit=week]")

    page.select_option("#chart-unit", "day")
    page.wait_for_selector("svg.chart-head[data-unit=day]")
    assert page.errors == []


def test_going_back_restores_the_form(page, filled_book):
    page.set_input_files("#import-file", str(filled_book))
    page.wait_for_selector("#btn-back:not([hidden])")
    page.click("#btn-back")
    page.wait_for_selector("#spec-form:not([hidden])")
    assert page.is_visible("#period-section")
    assert page.is_hidden("#totals")
    assert page.errors == []


def test_importing_a_non_excel_file_shows_a_message(page, tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hello", encoding="utf-8")
    page.set_input_files("#import-file", str(path))
    page.wait_for_selector("#banner:not([hidden])")
    assert "対応していない形式" in page.text_content("#banner")
    assert page.errors == []


def test_exporting_the_chart(page, filled_book, tmp_path):
    page.set_input_files("#import-file", str(filled_book))
    page.wait_for_selector("#btn-export:not([hidden])")

    with page.expect_download() as download:
        page.click("#btn-export")
    saved = tmp_path / "gantt.xlsx"
    download.value.save_as(saved)

    import zipfile
    with zipfile.ZipFile(saved) as zf:
        assert "xl/drawings/drawing1.xml" in zf.namelist()
    assert "ガント" in download.value.suggested_filename
    assert page.errors == []


def test_the_export_button_is_hidden_before_importing(page):
    _reset(page)
    assert page.is_hidden("#btn-export")
    assert page.is_visible("#btn-build")
    assert page.errors == []


# ---------------------------------------------------------------- 言語
def _set_language(page, value):
    page.select_option("#language", value)
    page.wait_for_function(f"() => document.documentElement.lang === '{value}'")
    page.wait_for_selector("table.wbs tbody tr")


def test_the_page_starts_in_japanese(page):
    _reset(page)
    _set_language(page, "ja")
    assert page.text_content("h1") == "WBS ジェネレータ"
    assert page.text_content("#btn-build") == "Excel をダウンロード"
    headers = page.eval_on_selector_all(
        "table.wbs thead tr:nth-child(2) th", "n => n.map(x => x.textContent)")
    assert headers[:4] == ["大項目", "中項目", "項番", "項目"]


def test_switching_to_english_translates_the_page(page):
    _reset(page)
    _set_language(page, "en")
    assert page.text_content("h1") == "WBS Generator"
    assert page.text_content("#btn-build") == "Download Excel"
    headers = page.eval_on_selector_all(
        "table.wbs thead tr:nth-child(2) th", "n => n.map(x => x.textContent)")
    assert headers[:4] == ["Group", "Sub-group", "No.", "Task"]
    assert _months(page)[:2] == ["Apr", "May"]
    assert "columns" in page.text_content("#preview-info")
    _set_language(page, "ja")
    assert page.errors == []


def test_the_suggested_title_follows_the_language(page):
    _reset(page)
    _set_language(page, "ja")
    assert "スケジュール" in page.input_value("#title")
    _set_language(page, "en")
    assert page.input_value("#title").startswith("FY")

    # 自分で書いた名前は言語を変えても残す
    page.fill("#title", "My Project")
    _set_language(page, "ja")
    assert page.input_value("#title") == "My Project"
    _set_language(page, "en")
    assert page.input_value("#title") == "My Project"
    assert page.errors == []


def test_the_language_survives_a_reload(page, server):
    _reset(page)
    _set_language(page, "en")
    page.reload(wait_until="networkidle")
    page.wait_for_selector("table.wbs tbody tr")
    assert page.input_value("#language") == "en"
    assert page.text_content("h1") == "WBS Generator"
    _set_language(page, "ja")
    assert page.errors == []


def test_downloading_in_english(page, tmp_path):
    _reset(page)
    _set_language(page, "en")
    with page.expect_download() as download:
        page.click("#btn-build")
    saved = tmp_path / "en.xlsx"
    download.value.save_as(saved)

    import openpyxl

    assert openpyxl.load_workbook(saved).sheetnames == ["Schedule", "Members", "Settings"]
    _set_language(page, "ja")
    assert page.errors == []


def test_the_indent_in_a_task_name_is_shown(page, make_filled):
    """項目名の字下げを画面でもそのまま見せる (HTML は空白を詰めるため)。"""
    book = make_filled("indent-ui.xlsx", rows=[
        ("開発", "", "1", "テスト実施", D(2026, 4, 1), 10, D(2026, 4, 14),
         None, None, None, None, None, "", "設計", ""),
        ("", "", "2", "\u3000WEB口座開設システム", D(2026, 4, 1), 10, D(2026, 4, 14),
         None, None, None, None, None, "", "設計", ""),
    ])
    page.set_input_files("#import-file", str(book))
    page.wait_for_function(
        "() => document.querySelectorAll('svg.chart-body rect[rx=\"2\"]').length > 0")

    names = page.eval_on_selector_all(
        "table.wbs tbody td.name", "n => n.map(x => x.textContent)")
    assert names[:2] == ["テスト実施", "\u3000WEB口座開設システム"]
    # 詰められていないこと (字下げが見た目にも残る)
    lefts = page.eval_on_selector_all(
        "table.wbs tbody td.name",
        "n => n.slice(0, 2).map(x => x.getBoundingClientRect().left)")
    assert lefts[0] == lefts[1]              # セル自体は同じ位置
    assert page.eval_on_selector(
        "table.wbs tbody td.name",
        "n => getComputedStyle(n).whiteSpace") == "pre"
    assert page.errors == []



def test_the_font_colour_is_shown_as_written(page, make_filled):
    """記入した文字色は、画面でもそのまま見せる。"""
    book = make_filled("color-ui.xlsx")
    import openpyxl
    from wbsgen.workbook import SHEET_PLAN
    wb = openpyxl.load_workbook(book)
    for coordinate, colour in (("E5", "FFFF0000"),     # 項目
                               ("H9", "FF7030A0")):    # 導き出す終了日
        cell = wb[SHEET_PLAN][coordinate]
        cell.font = openpyxl.styles.Font(name=cell.font.name, size=cell.font.sz,
                                         color=colour)
    wb.save(book)

    page.set_input_files("#import-file", str(book))
    # 直前のテストの表が残っていることがあるので、この色になるまで待つ
    page.wait_for_function("""() => {
        const cell = document.querySelector('table.wbs tbody td.name');
        return cell && getComputedStyle(cell).color === 'rgb(255, 0, 0)';
    }""")
    # 導き出した値のセルも、書かれた色のまま
    assert page.eval_on_selector(
        "table.wbs tbody tr:nth-child(5) td:nth-child(7)",
        "n => getComputedStyle(n).color") == "rgb(112, 48, 160)"
    assert page.errors == []


# ---------------------------------------------------------------- 処理中の表示
class _Held:
    """通信を握って離さないでおく。処理中の画面をゆっくり確かめるため。

    ハンドラの中で待つと Playwright 自体が止まってしまうので、受け取った
    まま返さずに置いておき、確認が済んでから流す。
    """

    def __init__(self, page, pattern):
        self.page, self.pattern, self.routes = page, pattern, []
        self.handler = lambda route: self.routes.append(route)
        page.route(pattern, self.handler)

    def wait(self):
        for _ in range(100):
            if self.routes:
                return
            self.page.wait_for_timeout(50)
        raise AssertionError(f"{self.pattern} への通信が来ませんでした")

    def release(self):
        # 先に握っていたぶんを流してから外す (外すと自動で流れてしまうため)
        for route in self.routes:
            route.continue_()
        self.routes = []
        self.page.unroute(self.pattern, self.handler)


def _busy_state(page, button):
    """処理中の見え方を、ひと呼びでまとめて取る (途中で終わらないように)。"""
    return page.evaluate("""(id) => ({
        shown: !document.querySelector('#busy').hidden,
        text: document.querySelector('#busy-text').textContent,
        spinner: !!document.querySelector('.spinner'),
        marked: document.body.classList.contains('is-busy'),
        locked: getComputedStyle(document.querySelector(id)).pointerEvents === 'none',
    })""", button)


def test_importing_shows_what_is_happening(page, filled_book):
    held = _Held(page, "**/api/import*")
    try:
        page.set_input_files("#import-file", str(filled_book))
        held.wait()
        state = _busy_state(page, "#btn-build")
    finally:
        held.release()

    assert state["shown"] and state["spinner"] and state["marked"]
    assert "読み込んでいます" in state["text"]
    assert filled_book.name in state["text"]
    assert state["locked"]                   # 処理中はボタンを押せない

    page.wait_for_selector("#busy", state="hidden")
    assert not page.evaluate("document.body.classList.contains('is-busy')")
    assert page.errors == []


def test_exporting_shows_what_is_happening(page, filled_book):
    page.set_input_files("#import-file", str(filled_book))
    page.wait_for_function(
        "() => document.querySelectorAll('svg.chart-body rect[rx=\"2\"]').length > 0")

    held = _Held(page, "**/api/export*")
    with page.expect_download():
        try:
            page.click("#btn-export")
            held.wait()
            state = _busy_state(page, "#btn-export")
        finally:
            held.release()

    assert state["shown"] and state["marked"] and state["locked"]
    assert "ガントチャート" in state["text"]
    page.wait_for_selector("#busy", state="hidden")
    assert page.errors == []


def test_building_a_blank_wbs_shows_what_is_happening(page):
    _reset(page)
    held = _Held(page, "**/api/build*")
    with page.expect_download():
        try:
            page.click("#btn-build")
            held.wait()
            state = _busy_state(page, "#btn-build")
        finally:
            held.release()

    assert state["shown"] and state["marked"] and state["locked"]
    assert "作っています" in state["text"]
    page.wait_for_selector("#busy", state="hidden")
    assert page.errors == []


# ------------------------------------------------ スマートフォン (狭い画面)
def _phone_reset(phone):
    """新規作成の指定タブに戻す。"""
    if phone.is_visible("#btn-back"):
        phone.click("#btn-back")
    phone.wait_for_selector("#pane-tabs:not([hidden])")
    phone.click('.pane-tab[data-pane="form"]')


def test_the_phone_layout_switches_between_the_form_and_the_preview(phone):
    _phone_reset(phone)
    assert phone.is_visible("#spec-form")
    assert phone.is_hidden("#sheet")

    phone.click('.pane-tab[data-pane="preview"]')
    assert phone.is_visible("#sheet")
    assert phone.is_hidden("#spec-form")
    assert phone.errors == []


def test_the_phone_page_never_scrolls_sideways(phone):
    """本文がはみ出さないこと (横に振れる画面は使いにくい)。"""
    _phone_reset(phone)
    width, client = phone.evaluate(
        "[document.body.scrollWidth, document.body.clientWidth]")
    assert width <= client


def test_the_phone_shows_the_chart_next_to_the_names(phone, filled_book):
    """既定では項目だけを出して、日程表を画面に入れる。"""
    phone.set_input_files("#import-file", str(filled_book))
    phone.wait_for_function(
        "() => document.querySelectorAll('svg.chart-body rect[rx=\"2\"]').length > 0")

    headers = phone.eval_on_selector_all(
        "table.wbs thead tr:nth-child(2) th", "n => n.map(x => x.textContent)")
    assert headers == ["項目"]
    # 日程表が画面の中に入っている
    left = phone.eval_on_selector("svg.chart-body", "n => n.getBoundingClientRect().left")
    assert left < 390
    assert phone.errors == []


def test_the_phone_keeps_the_names_in_view_while_scrolling(phone, filled_book):
    phone.set_input_files("#import-file", str(filled_book))
    phone.wait_for_function(
        "() => document.querySelectorAll('svg.chart-body rect[rx=\"2\"]').length > 0")
    phone.select_option("#columns-mode", "min")

    phone.eval_on_selector(".sheet", "n => n.scrollLeft = 600")
    phone.wait_for_timeout(100)
    left = phone.eval_on_selector("table.wbs", "n => n.getBoundingClientRect().left")
    assert 0 <= left < 40                    # 貼り付いたまま
    assert phone.eval_on_selector("table.wbs tbody td", "n => n.textContent") == "要件定義"


def test_the_phone_can_show_more_columns(phone, filled_book):
    phone.set_input_files("#import-file", str(filled_book))
    phone.wait_for_function(
        "() => document.querySelectorAll('svg.chart-body rect[rx=\"2\"]').length > 0")

    phone.select_option("#columns-mode", "key")
    assert phone.eval_on_selector_all(
        "table.wbs thead tr:nth-child(2) th", "n => n.length") == 7
    phone.select_option("#columns-mode", "all")
    assert phone.eval_on_selector_all(
        "table.wbs thead tr:nth-child(2) th", "n => n.length") == 16
    phone.select_option("#columns-mode", "min")
    assert phone.errors == []


def test_the_desktop_keeps_every_column(page, filled_book):
    """広い画面はこれまでどおり (列を減らす選択も出さない)。"""
    page.set_input_files("#import-file", str(filled_book))
    page.wait_for_function(
        "() => document.querySelectorAll('svg.chart-body rect[rx=\"2\"]').length > 0")
    assert page.eval_on_selector_all(
        "table.wbs thead tr:nth-child(2) th", "n => n.length") == 16
    assert page.is_hidden("#columns-field")
    assert page.is_hidden("#pane-tabs")
