"""セルの数式を、書き出しでも残すテスト。"""

import datetime as dt

import openpyxl
import pytest

from wbsgen import formulas
from wbsgen.importer import apply_edits, read
from wbsgen.inject import inject_cached_values, sheet_part
from wbsgen.workbook import export

D = dt.date
BASE = D(2026, 6, 10)


# ---------------------------------------------------------------- 参照の書き換え
def test_references_are_moved_to_the_new_positions():
    cols, rows = {6: 7, 8: 9}, {5: 6, 6: 7}
    assert formulas.translate("=F5+13", cols, rows, "S") == "=G6+13"
    assert formulas.translate("=NETWORKDAYS($F5,H6)+SUM(F5:H6)", cols, rows, "S") \
        == "=NETWORKDAYS($G6,I7)+SUM(G6:I7)"


def test_other_sheets_and_names_are_left_alone():
    cols, rows = {6: 7}, {5: 6}
    assert formulas.translate("=Sheet2!A1+F5", cols, rows, "S") == "=Sheet2!A1+G6"
    assert formulas.translate("='S'!F5", cols, rows, "S") == "='S'!G6"      # 自分のシート
    assert formulas.translate("=TaxRate*F5", cols, rows, "S") == "=TaxRate*G6"


def test_a_formula_pointing_outside_the_table_cannot_be_moved():
    assert formulas.translate("=F9+1", {6: 7}, {5: 6}, "S") is None        # 書き出さない行
    assert formulas.translate("=SUM(F:F)", {6: 7}, {5: 6}, "S") is None    # 列まるごと


def test_nothing_moves_when_positions_stay():
    assert not formulas.moves({6: 6, 8: 8}, {5: 5, 6: 6})
    assert formulas.moves({6: 7}, {5: 5})


# ---------------------------------------------------------------- 書き出し
def _with_formulas(make_filled, tmp_path, formula_cells, cached):
    """``formula_cells`` (セル -> 数式) を入れ、計算結果 ``cached`` を添えたファイル。"""
    path = make_filled()
    book = openpyxl.load_workbook(path)
    sheet = book["スケジュール"]
    for ref, formula in formula_cells.items():
        sheet[ref] = formula
    book.save(path)
    inject_cached_values(path, sheet_part(path, "スケジュール"), cached)
    return path


def test_formulas_are_kept_in_the_export(make_filled, tmp_path):
    path = _with_formulas(
        make_filled, tmp_path,
        {"H5": "=F5+13", "G5": "=NETWORKDAYS(F5,H5)",
         "R5": '=IF(M5>=1,"完了","実行中")'},
        {"H5": D(2026, 4, 14), "G5": 10, "R5": "完了"})
    imported = read(path, base_date=BASE)
    assert imported.rows[0].formulas == {
        "end": "=F5+13", "days": "=NETWORKDAYS(F5,H5)", "status": '=IF(M5>=1,"完了","実行中")'}

    out = export(imported, tmp_path / "out.xlsx", BASE)
    sheet = openpyxl.load_workbook(out)["スケジュール"]
    assert sheet["H5"].value == "=F5+13"
    assert sheet["G5"].value == "=NETWORKDAYS(F5,H5)"
    assert sheet["R5"].value == '=IF(M5>=1,"完了","実行中")'
    assert sheet["H6"].value == D(2026, 5, 7) or sheet["H6"].value.date() == D(2026, 5, 7)  # 数式でない行は値


def test_the_exported_formulas_carry_their_results(make_filled, tmp_path):
    """数式を計算し直さない読み手 (このツールの読み込みを含む) にも値が見える。"""
    path = _with_formulas(
        make_filled, tmp_path, {"H5": "=F5+13", "G5": "=NETWORKDAYS(F5,H5)"},
        {"H5": D(2026, 4, 14), "G5": 10})
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)
    cached = openpyxl.load_workbook(out, data_only=True)["スケジュール"]
    assert cached["H5"].value.date() == D(2026, 4, 14)
    assert cached["G5"].value == 10

    again = read(out, base_date=BASE)                    # 読み直しても値が失われない
    assert again.rows[0].end == D(2026, 4, 14)
    assert again.warnings == []
    assert again.rows[0].formulas["end"] == "=F5+13"     # 数式も残っている


def test_excel_is_told_to_recalculate_on_open(make_filled, tmp_path):
    import zipfile

    path = _with_formulas(make_filled, tmp_path, {"H5": "=F5+13"}, {"H5": D(2026, 4, 14)})
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)
    with zipfile.ZipFile(out) as zf:
        assert b'fullCalcOnLoad="1"' in zf.read("xl/workbook.xml")


def test_references_follow_when_rows_move(make_filled, tmp_path):
    """空行が詰められて行の位置が変わっても、数式は自分の行を指したままになる。"""
    path = make_filled()
    book = openpyxl.load_workbook(path)
    sheet = book["スケジュール"]
    sheet.insert_rows(6)                          # 5 行目の下に空行
    sheet["H5"], sheet["H7"] = "=F5+13", "=F7+13"
    sheet["O8"] = "=D7"                           # 上の行 (空行をまたぐ) の項番
    book.save(path)
    inject_cached_values(path, sheet_part(path, "スケジュール"),
                         {"H5": D(2026, 4, 14), "H7": D(2026, 4, 28), "O8": "102"})
    imported = read(path, base_date=BASE)
    assert [r.row for r in imported.rows][:3] == [5, 7, 8]

    out = export(imported, tmp_path / "out.xlsx", BASE)
    written = openpyxl.load_workbook(out)["スケジュール"]
    assert written["H5"].value == "=F5+13"
    assert written["H6"].value == "=F6+13"        # 7 行目が 6 行目に詰まる
    assert written["O7"].value == "=D6"


def test_a_formula_that_cannot_be_moved_is_written_as_its_value(make_filled, tmp_path):
    path = make_filled()
    book = openpyxl.load_workbook(path)
    sheet = book["スケジュール"]
    sheet.insert_rows(6)
    sheet["H7"] = "=F5+F6"                        # 空行 (書き出さない) を指す
    book.save(path)
    inject_cached_values(path, sheet_part(path, "スケジュール"), {"H7": D(2026, 5, 8)})
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)
    cell = openpyxl.load_workbook(out)["スケジュール"]["H6"]
    assert not str(cell.value).startswith("=")    # 数式のかわりに値
    assert cell.value is not None


def test_an_edited_cell_replaces_its_formula(make_filled, tmp_path):
    path = _with_formulas(make_filled, tmp_path, {"H5": "=F5+13"}, {"H5": D(2026, 4, 14)})
    imported = read(path, base_date=BASE)
    assert "end" in imported.rows[0].formulas
    apply_edits(imported, {"5": {"end": "2026-04-30"}}, base_date=BASE)
    assert "end" not in imported.rows[0].formulas
    out = export(imported, tmp_path / "out.xlsx", BASE)
    value = openpyxl.load_workbook(out)["スケジュール"]["H5"].value
    assert not str(value).startswith("=")
    assert value.date() == D(2026, 4, 30)


def test_a_file_without_formulas_is_unchanged(make_filled, tmp_path):
    imported = read(make_filled(), base_date=BASE)
    assert all(not row.formulas for row in imported.rows)
    out = export(imported, tmp_path / "out.xlsx", BASE)
    sheet = openpyxl.load_workbook(out)["スケジュール"]
    assert not any(str(c.value).startswith("=") for r in sheet.iter_rows(min_row=5, max_row=9)
                   for c in r if c.value is not None)
