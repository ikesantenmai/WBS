"""空の WBS (期間だけを指定した WBS) のテスト。"""

from __future__ import annotations

import datetime as dt
import zipfile

import openpyxl
import pytest

from wbsgen.blank import DEFAULT_ROWS, add_months, build_blank
from wbsgen.cli import main
from wbsgen.loader import ProjectError, from_dict
from wbsgen.render import WorkbookRenderer
from wbsgen.scheduling import resolve_project
from wbsgen.serialize import to_dict


# ---------------------------------------------------------------- 期間
def test_period_from_end_date():
    project = build_blank(dt.date(2026, 4, 1), end=dt.date(2027, 3, 31))
    assert project.chart.period_days == 365


def test_period_from_months():
    project = build_blank(dt.date(2026, 4, 1), months=6)
    assert project.chart.period_days == (dt.date(2026, 10, 1) - dt.date(2026, 4, 1)).days


def test_period_defaults_to_a_year():
    project = build_blank(dt.date(2026, 4, 1))
    assert project.chart.period_days == 365
    assert project.blank_rows == DEFAULT_ROWS


@pytest.mark.parametrize("start,months,expected", [
    ((2026, 1, 31), 1, (2026, 2, 28)),
    ((2028, 1, 31), 1, (2028, 2, 29)),
    ((2026, 12, 1), 3, (2027, 3, 1)),
])
def test_add_months_clamps_to_the_month_end(start, months, expected):
    assert add_months(dt.date(*start), months) == dt.date(*expected)


def test_end_before_start_is_rejected():
    with pytest.raises(ProjectError, match="終了日"):
        build_blank(dt.date(2026, 4, 1), end=dt.date(2026, 3, 1))


def test_too_many_rows_is_rejected():
    with pytest.raises(ProjectError, match="行数"):
        build_blank(dt.date(2026, 4, 1), rows=99999)


# ---------------------------------------------------------------- 解決
def test_a_project_with_no_task_resolves():
    project = build_blank(dt.date(2026, 4, 1), months=3)
    resolve_project(project)
    assert project.tasks == []
    assert project.chart.start == dt.date(2026, 4, 1)


def test_definition_without_tasks_needs_a_chart_start():
    with pytest.raises(ProjectError, match="tasks"):
        from_dict({"project": {"title": "x"}, "tasks": []})

    project = from_dict({"chart": {"start": "2026-04-01"}, "tasks": [], "blank_rows": 10})
    resolve_project(project)
    assert project.blank_rows == 10


def test_blank_rows_round_trip():
    project = build_blank(dt.date(2026, 4, 1), rows=25)
    assert from_dict(to_dict(project)).blank_rows == 25


@pytest.mark.parametrize("value", [-1, 5000, "abc"])
def test_invalid_blank_rows_is_rejected(value):
    with pytest.raises(ProjectError, match="blank_rows"):
        from_dict({"chart": {"start": "2026-04-01"}, "tasks": [], "blank_rows": value})


# ---------------------------------------------------------------- 出力
@pytest.fixture
def blank_book(tmp_path):
    project = build_blank(dt.date(2026, 4, 1), end=dt.date(2027, 3, 31),
                          rows=30, title="2026年度 スケジュール")
    path = tmp_path / "blank.xlsx"
    WorkbookRenderer(project).save(path)
    return path


def test_blank_workbook_has_headers_and_empty_rows(blank_book):
    ws = openpyxl.load_workbook(blank_book)["スケジュール"]
    assert ws["B1"].value == "2026年度 スケジュール"
    assert [ws[f"{c}4"].value for c in "BCDE"] == ["大項目", "中項目", "項番", "項目"]
    assert ws.freeze_panes == "S5"
    # 5 行目から 30 行ぶんの記入枠
    assert ws.max_row == 34
    assert ws["E5"].value is None
    assert ws["E34"].value is None


def test_blank_rows_keep_borders_and_formats(blank_book):
    ws = openpyxl.load_workbook(blank_book)["スケジュール"]
    for row in (5, 20, 34):
        assert ws[f"E{row}"].border.left.style == "thin"
        assert ws[f"F{row}"].number_format == "m/dd"
        assert ws[f"G{row}"].number_format == '0\\ "日"'
        assert ws[f"M{row}"].number_format == "0%"
        assert ws[f"F{row}"].fill.fgColor.rgb.endswith("FFFFCC")   # 予定欄の色


def test_blank_workbook_still_draws_the_now_line(blank_book):
    """タスクが無くても、期間内なら現在日線は引かれる。"""
    with zipfile.ZipFile(blank_book) as zf:
        assert "xl/drawings/drawing1.xml" in zf.namelist()
        xml = zf.read("xl/drawings/drawing1.xml").decode()
    assert "now-line" in xml


def test_blank_workbook_outside_the_period_has_no_drawing(tmp_path):
    project = build_blank(dt.date(2040, 4, 1), months=3, rows=5)
    path = tmp_path / "future.xlsx"
    WorkbookRenderer(project).save(path)
    with zipfile.ZipFile(path) as zf:
        assert "xl/drawings/drawing1.xml" not in zf.namelist()


# ---------------------------------------------------------------- CLI
def test_cli_blank_writes_a_workbook(tmp_path, capsys):
    output = tmp_path / "blank.xlsx"
    assert main(["blank", "--start", "2026-04-01", "--end", "2027-03-31",
                 "--rows", "50", "--title", "年度計画", "-o", str(output)]) == 0
    out = capsys.readouterr().out
    assert "2026-04-01 〜 2027-03-31 (365 日)" in out
    assert "50 行" in out

    ws = openpyxl.load_workbook(output)["スケジュール"]
    assert ws["B1"].value == "年度計画"
    assert ws.max_row == 54


def test_cli_blank_writes_a_definition_file(tmp_path):
    output = tmp_path / "blank.yaml"
    assert main(["blank", "--start", "2026-04-01", "--months", "6",
                 "--unit", "day", "-o", str(output)]) == 0
    import yaml

    data = yaml.safe_load(output.read_text(encoding="utf-8"))
    assert data["chart"]["unit"] == "day"
    assert data["tasks"] == []
    assert data["blank_rows"] == DEFAULT_ROWS

    # 書き出した定義がそのまま build に通る
    assert main(["build", str(output), "-o", str(tmp_path / "from-def.xlsx")]) == 0


def test_cli_blank_registers_members(tmp_path):
    output = tmp_path / "m.xlsx"
    assert main(["blank", "--start", "2026-04-01", "--member", "設計",
                 "--member", "製造", "-o", str(output)]) == 0
    ws = openpyxl.load_workbook(output)["担当者一覧"]
    assert [ws[f"B{r}"].value for r in (4, 5)] == ["設計", "製造"]


def test_cli_blank_rejects_a_bad_date(capsys):
    assert main(["blank", "--start", "2026/04/01"]) == 2
    assert "YYYY-MM-DD" in capsys.readouterr().err
