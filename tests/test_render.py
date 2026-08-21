import datetime as dt
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pytest

from wbsgen.loader import from_dict
from wbsgen.model import UNIT_DAY, UNIT_MONTH, UNIT_WEEK
from wbsgen.render import WorkbookRenderer
from wbsgen.render.timeline import Timeline
from wbsgen.workcal import WorkCalendar

NS_XDR = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.fixture
def project():
    return from_dict({
        "project": {"title": "テスト計画"},
        "chart": {"start": "2026-04-01", "period_days": 90, "unit": "week",
                  "base_date": "2026-04-20",
                  "show": {"progress": True}},
        "members": [{"name": "DS", "color": "#4472C4"}],
        "tasks": [
            {"group": "開発", "no": "10", "name": "開発", "kind": "summary", "member": "DS"},
            {"subgroup": "設計", "no": "11", "name": "画面設計", "start": "2026-04-01",
             "days": 5, "effort": 5, "actual_start": "2026-04-01", "progress": 1.0,
             "member": "DS", "comment": "レビュー済み"},
            {"no": "12", "name": "実装", "predecessor": "11", "days": 10, "effort": 20,
             "actual_start": "2026-04-08", "progress": 0.5, "member": "DS"},
            {"group": "試験", "no": "20", "name": "テスト", "predecessor": "12", "days": 5,
             "effort": 5, "member": "DS"},
            {"no": "M1", "name": "リリース", "kind": "milestone", "shape": "diamond",
             "predecessor": "20", "days": 1, "member": "DS"},
        ],
    })


@pytest.fixture
def built(project, tmp_path):
    path = tmp_path / "out.xlsx"
    WorkbookRenderer(project).save(path)
    return path


# ----------------------------------------------------------------------
def test_workbook_has_the_four_sheets(built):
    wb = openpyxl.load_workbook(built)
    assert wb.sheetnames == ["スケジュール", "担当者一覧", "標準工程", "設定"]


def test_schedule_sheet_headers_and_values(built):
    ws = openpyxl.load_workbook(built)["スケジュール"]
    assert ws["B1"].value == "テスト計画"
    assert [ws[f"{c}4"].value for c in "BCDE"] == ["大項目", "中項目", "項番", "項目"]
    assert [ws[f"{c}4"].value for c in "FGH"] == ["開始日", "日数", "終了日"]
    assert ws["F3"].value == "予定"
    assert ws["I3"].value == "実績"
    # 1 行目のタスク (工程バー) から順に並ぶ
    assert ws["E5"].value == "開発"
    assert ws["E6"].value == "画面設計"
    assert ws["F6"].value == dt.datetime(2026, 4, 1)
    assert ws["H6"].value == dt.datetime(2026, 4, 7)
    assert ws["R6"].value == "完了"


def test_group_separator_row_shows_total_effort(built):
    ws = openpyxl.load_workbook(built)["スケジュール"]
    values = [ws.cell(row=r, column=2).value for r in range(5, 20)]
    assert "[ 25 人日 ]" in values      # 開発グループ = 5 + 20 人日


def test_panes_are_frozen_at_the_chart_origin(built):
    ws = openpyxl.load_workbook(built)["スケジュール"]
    assert ws.freeze_panes == "S5"
    assert ws.sheet_view.showGridLines is False


def test_comment_is_attached(built):
    ws = openpyxl.load_workbook(built)["スケジュール"]
    assert ws["E6"].comment is not None
    assert "レビュー済み" in ws["E6"].comment.text


def test_hidden_columns_follow_display_flags(tmp_path):
    project = from_dict({
        "chart": {"start": "2026-04-01", "show": {"manpower": False, "status": False}},
        "tasks": [{"name": "A", "start": "2026-04-01", "days": 1}],
    })
    path = tmp_path / "hidden.xlsx"
    WorkbookRenderer(project).save(path)
    ws = openpyxl.load_workbook(path)["スケジュール"]
    assert ws.column_dimensions["N"].hidden is True
    assert ws.column_dimensions["R"].hidden is True
    assert ws.column_dimensions["F"].hidden is False


# ----------------------------------------------------------------------
def test_drawing_part_is_wired_into_the_package(built):
    with zipfile.ZipFile(built) as zf:
        names = zf.namelist()
        assert "xl/drawings/drawing1.xml" in names
        rels = zf.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
        assert "../drawings/drawing1.xml" in rels
        types = zf.read("[Content_Types].xml").decode()
        assert "/xl/drawings/drawing1.xml" in types
        sheet = zf.read("xl/worksheets/sheet1.xml").decode()
        assert "<drawing r:id=" in sheet
        assert 'xmlns:r="' in sheet
        ET.fromstring(sheet)          # 名前空間が解決できること


def test_drawing_contains_bars_milestone_and_lines(built):
    with zipfile.ZipFile(built) as zf:
        root = ET.fromstring(zf.read("xl/drawings/drawing1.xml"))
    anchors = root.findall(f"{NS_XDR}twoCellAnchor")
    assert anchors, "図形が 1 つも出力されていない"

    names = [
        sp.find(f"{NS_XDR}nvSpPr/{NS_XDR}cNvPr").get("name")
        for sp in root.iter(f"{NS_XDR}sp")
    ]
    assert any(n.startswith("summary-") for n in names)
    assert any(n.startswith("plan-") for n in names)
    assert any(n.startswith("actual-") for n in names)
    assert any(n.startswith("progress-") for n in names)
    assert any(n.startswith("milestone-") for n in names)
    assert any(n.startswith("now-line") for n in names)
    assert any(n.startswith("link") for n in names)

    presets = {
        g.get("prst") for g in root.iter(f"{NS_A}prstGeom")
    }
    assert {"rect", "chevron", "diamond", "line"} <= presets


def test_every_shape_has_a_positive_extent(built):
    with zipfile.ZipFile(built) as zf:
        root = ET.fromstring(zf.read("xl/drawings/drawing1.xml"))
    for sp in root.iter(f"{NS_XDR}sp"):
        name = sp.find(f"{NS_XDR}nvSpPr/{NS_XDR}cNvPr").get("name")
        ext = sp.find(f"{NS_XDR}spPr/{NS_A}xfrm/{NS_A}ext")
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        assert cx >= 0 and cy >= 0, name
        if not name.startswith(("link", "now-line")):
            assert cx > 0 and cy > 0, f"{name} の大きさが 0"


def test_bar_left_edge_matches_the_start_date_column(built, project):
    """予定バーの左端が開始日の列に載っていること。"""
    with zipfile.ZipFile(built) as zf:
        root = ET.fromstring(zf.read("xl/drawings/drawing1.xml"))
    for anchor in root.findall(f"{NS_XDR}twoCellAnchor"):
        name = anchor.find(f"{NS_XDR}sp/{NS_XDR}nvSpPr/{NS_XDR}cNvPr").get("name")
        if name.startswith("plan-11:画面設計"):
            col = int(anchor.find(f"{NS_XDR}from/{NS_XDR}col").text)
            # 2026-04-01 は表示開始週 (S 列 = index 18) に含まれる
            assert col == 18
            assert int(anchor.find(f"{NS_XDR}from/{NS_XDR}colOff").text) > 0
            return
    pytest.fail("予定バーが見つからない")


# ----------------------------------------------------------------------
@pytest.mark.parametrize("unit,expected_first_label", [
    (UNIT_DAY, "1"), (UNIT_WEEK, "3/30"), (UNIT_MONTH, "4月"),
])
def test_timeline_units(unit, expected_first_label):
    cal = WorkCalendar.build()
    tl = Timeline(dt.date(2026, 4, 1), 30, unit, cal)
    assert tl.column_label(tl.columns[0]) == expected_first_label


def test_timeline_position_maps_dates_into_columns():
    cal = WorkCalendar.build()
    tl = Timeline(dt.date(2026, 4, 1), 28, UNIT_WEEK, cal)
    # 週表示: 2026-03-30(月) が 0 列目の先頭
    assert tl.position(dt.date(2026, 3, 30)) == 0.0
    assert tl.position(dt.date(2026, 4, 2)) == pytest.approx(3 / 7)
    assert tl.position(dt.date(2026, 4, 5), end_of_day=True) == pytest.approx(1.0)
    assert tl.position(dt.date(2026, 4, 6)) == pytest.approx(1.0)


def test_timeline_clamps_out_of_range_dates():
    cal = WorkCalendar.build()
    tl = Timeline(dt.date(2026, 4, 1), 28, UNIT_DAY, cal)
    assert tl.position(dt.date(2020, 1, 1)) == 0.0
    assert tl.position(dt.date(2030, 1, 1)) == len(tl)
    assert tl.in_range(dt.date(2026, 4, 10), dt.date(2026, 4, 12))
    assert not tl.in_range(dt.date(2027, 1, 1), dt.date(2027, 2, 1))


def test_day_unit_marks_rest_columns():
    cal = WorkCalendar.build(holidays=[dt.date(2026, 4, 29)])
    tl = Timeline(dt.date(2026, 4, 1), 30, UNIT_DAY, cal)
    rest = [c.start for c in tl.columns if tl.is_rest_column(c)]
    assert dt.date(2026, 4, 4) in rest      # 土
    assert dt.date(2026, 4, 5) in rest      # 日
    assert dt.date(2026, 4, 29) in rest     # 祝日
    assert dt.date(2026, 4, 30) not in rest


def test_out_of_range_task_draws_no_bar(tmp_path):
    project = from_dict({
        "chart": {"start": "2026-04-01", "period_days": 30, "unit": "day"},
        "tasks": [
            {"name": "範囲内", "start": "2026-04-01", "days": 3},
            {"name": "範囲外", "start": "2030-01-01", "days": 3},
        ],
    })
    path = tmp_path / "range.xlsx"
    WorkbookRenderer(project).save(path)
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("xl/drawings/drawing1.xml"))
    names = [
        sp.find(f"{NS_XDR}nvSpPr/{NS_XDR}cNvPr").get("name")
        for sp in root.iter(f"{NS_XDR}sp")
    ]
    assert any("範囲内" in n for n in names)
    assert not any("範囲外" in n for n in names)
