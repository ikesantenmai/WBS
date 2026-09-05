"""Excel の色を画面用の色に直すテスト。

セルの文字色は RGB のほか、色番号 (indexed) やテーマ色 (theme) でも
指定される。どれも画面に出せる ``RRGGBB`` になること。
"""

import openpyxl
import pytest
from openpyxl.styles.colors import Color

from wbsgen.palette import theme_colors, to_rgb


def test_an_rgb_colour_passes_through():
    assert to_rgb(Color(rgb="FFFF0000")) == "FF0000"
    assert to_rgb(Color(rgb="FF0070C0")) == "0070C0"


def test_an_indexed_colour_is_looked_up():
    """古いファイルでよく使われる色番号を、実際の色に直す。"""
    assert to_rgb(Color(indexed=10)) == "FF0000"      # 赤
    assert to_rgb(Color(indexed=0)) == "000000"       # 黒


def test_an_out_of_range_index_is_automatic():
    """64 番以降は「自動」なので、色としては扱わない。"""
    assert to_rgb(Color(indexed=64)) is None
    assert to_rgb(Color(indexed=200)) is None


def test_a_theme_colour_uses_the_theme_of_the_book():
    theme = ["FFFFFF", "000000", "EEECE1", "1F497D", "4F81BD", "C0504D"]
    assert to_rgb(Color(theme=4), theme) == "4F81BD"
    assert to_rgb(Color(theme=3), theme) == "1F497D"


def test_a_theme_colour_without_a_theme_is_automatic():
    assert to_rgb(Color(theme=4), []) is None


def test_a_tint_lightens_or_darkens():
    """濃淡 (tint) は明度を動かす。正なら明るく、負なら暗く。"""
    theme = ["FFFFFF", "000000", "EEECE1", "1F497D", "4F81BD"]
    plain = to_rgb(Color(theme=4), theme)
    lighter = to_rgb(Color(theme=4, tint=0.6), theme)
    darker = to_rgb(Color(theme=4, tint=-0.5), theme)

    def brightness(rgb):
        return sum(int(rgb[i:i + 2], 16) for i in (0, 2, 4))

    assert brightness(darker) < brightness(plain) < brightness(lighter)


def test_no_colour_is_automatic():
    assert to_rgb(None) is None


@pytest.mark.parametrize("color", [Color(), Color(auto=True)])
def test_colours_that_cannot_be_resolved_are_automatic(color):
    assert to_rgb(color, ["FFFFFF"]) in (None, "000000")


# ---------------------------------------------------------------- テーマの読み取り
def test_the_theme_of_a_workbook_is_read(tmp_path):
    path = tmp_path / "themed.xlsx"
    openpyxl.Workbook().save(path)
    theme = theme_colors(openpyxl.load_workbook(path))

    assert len(theme) == 12
    assert all(len(color) == 6 for color in theme)
    # 先頭の 2 つは明暗の組 (背景・文字)
    assert theme[0] != theme[1]


def test_a_book_without_a_theme_gives_nothing():
    assert theme_colors(openpyxl.Workbook()) == []
