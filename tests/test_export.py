"""ガントチャート付き Excel の書き出しのテスト。"""

import datetime as dt
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pytest

from wbsgen.chart import build as build_chart
from wbsgen.importer import read
from wbsgen.workbook import SHEET_PLAN, export

NS_XDR = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
BASE = dt.date(2026, 6, 10)


@pytest.fixture
def exported(filled_book, tmp_path):
    path = tmp_path / "gantt.xlsx"
    export(read(filled_book, "filled.xlsx"), path, base_date=BASE)
    return path


def _shapes(path):
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("xl/drawings/drawing1.xml"))
    return {
        sp.find(f"{NS_XDR}nvSpPr/{NS_XDR}cNvPr").get("name").rsplit(" ", 1)[0]: sp
        for sp in root.iter(f"{NS_XDR}sp")
    }


def _anchors(path):
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("xl/drawings/drawing1.xml"))
    out = {}
    for anchor in root.findall(f"{NS_XDR}twoCellAnchor"):
        name = anchor.find(
            f"{NS_XDR}sp/{NS_XDR}nvSpPr/{NS_XDR}cNvPr").get("name").rsplit(" ", 1)[0]
        out[name] = anchor
    return out


# ---------------------------------------------------------------- 表
def test_rows_are_written(exported):
    ws = openpyxl.load_workbook(exported)[SHEET_PLAN]
    assert [ws[f"E{r}"].value for r in range(5, 10)] == [
        "要件定義", "基本設計", "詳細設計", "コーディング", "結合テスト"]
    assert ws["F5"].value == dt.datetime(2026, 4, 1)
    assert ws["G5"].value == 10
    assert ws["M5"].value == 1.0
    assert ws["Q5"].value == "設計"


def test_group_and_subgroup_appear_only_when_they_change(exported):
    ws = openpyxl.load_workbook(exported)[SHEET_PLAN]
    assert [ws[f"B{r}"].value for r in range(5, 10)] == [
        "開発", None, None, None, "テスト"]
    assert [ws[f"C{r}"].value for r in range(5, 10)] == [
        "要件", None, "製造", None, "結合"]


def test_a_missing_end_date_is_filled_in(exported):
    """終了日が空だった行は、日数から求めた値が書き込まれる。"""
    ws = openpyxl.load_workbook(exported)[SHEET_PLAN]
    assert ws["E9"].value == "結合テスト"
    assert ws["H9"].value == dt.datetime(2026, 8, 17)


def test_the_status_is_calculated_and_coloured(exported):
    """状態は基準日から計算して書き込む (配色は添付ファイルと同じ)。"""
    ws = openpyxl.load_workbook(exported)[SHEET_PLAN]
    # 101/102 は実績終了日が入っているので完了
    assert ws["R5"].value == "完了"
    assert ws["R5"].fill.fgColor.rgb.endswith("C0C0C0")
    # 201 は予定終了 5/22 を過ぎて未完了 (基準日 6/10)
    assert ws["R7"].value == "遅れ 13 日"
    assert ws["L7"].value == 13                              # 遅れ列も揃える
    assert ws["R7"].fill.fgColor.rgb.endswith("FF99CC")
    # 202 は着手済みで予定終了 6/19 まで
    assert ws["R8"].value == "残り 7 日"
    assert ws["R8"].fill.fgColor.rgb.endswith("FFCC00")
    # 301 は未着手
    assert ws["R9"].value.startswith("あと")
    assert ws["R9"].fill.fgColor.rgb.endswith("CCFFFF")


def test_the_timeline_header_is_kept(exported):
    ws = openpyxl.load_workbook(exported)[SHEET_PLAN]
    assert ws["S3"].value == dt.datetime(2026, 4, 1)
    assert ws["S3"].number_format == 'm"月"'
    assert ws["S4"].number_format == "m/d"
    assert ws.freeze_panes == "S5"


# ---------------------------------------------------------------- 図形
def test_bars_are_drawn_for_every_row(exported):
    names = _shapes(exported)
    for no in ("101", "102", "201", "202", "301"):
        assert f"plan-{no}" in names
    # 実績のバーを描くのは、終了日まで入っている 101 / 102 の 2 行
    assert sum(1 for n in names if n.startswith("actual-")) == 2
    assert "now-line" in names


def test_progress_is_drawn_inside_the_plan_bar(exported):
    anchors = _anchors(exported)
    plan = anchors["plan-201"]
    progress = anchors["progress-201"]

    def left(anchor):
        node = anchor.find(f"{NS_XDR}from")
        return (int(node.find(f"{NS_XDR}col").text),
                int(node.find(f"{NS_XDR}colOff").text))

    assert left(plan) == left(progress)          # 左端はそろう
    plan_w = int(plan.find(f"{NS_XDR}sp/{NS_XDR}spPr/{NS_A}xfrm/{NS_A}ext").get("cx"))
    prog_w = int(progress.find(f"{NS_XDR}sp/{NS_XDR}spPr/{NS_A}xfrm/{NS_A}ext").get("cx"))
    assert 0 < prog_w < plan_w                   # 進捗 60% ぶんだけ短い


def test_bars_land_on_the_dates(exported):
    """予定バーの左端が開始日の列に載っていること。"""
    anchor = _anchors(exported)["plan-101"]
    node = anchor.find(f"{NS_XDR}from")
    # 2026-04-01 は表示開始日そのもの → 先頭列 (S = index 18) の左端
    assert int(node.find(f"{NS_XDR}col").text) == 18
    assert int(node.find(f"{NS_XDR}colOff").text) == 0


def test_every_shape_has_a_positive_extent(exported):
    for name, shape in _shapes(exported).items():
        ext = shape.find(f"{NS_XDR}spPr/{NS_A}xfrm/{NS_A}ext")
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        assert cx >= 0 and cy >= 0, name
        if name != "now-line":
            assert cx > 0 and cy > 0, f"{name} の大きさが 0"


def test_members_get_different_colours(exported):
    shapes = _shapes(exported)

    def colour(name):
        return shapes[name].find(f".//{NS_A}srgbClr").get("val")

    assert colour("plan-101") != colour("plan-201")    # 設計 と 製造
    assert colour("plan-101") == colour("plan-102")    # どちらも 設計


def test_no_now_line_outside_the_period(filled_book, tmp_path):
    path = tmp_path / "future.xlsx"
    export(read(filled_book), path, base_date=dt.date(2040, 1, 1))
    assert "now-line" not in _shapes(path)


def test_the_drawing_is_wired_into_the_package(exported):
    with zipfile.ZipFile(exported) as zf:
        names = zf.namelist()
        assert "xl/drawings/drawing1.xml" in names
        assert "../drawings/drawing1.xml" in zf.read(
            "xl/worksheets/_rels/sheet1.xml.rels").decode()
        sheet = zf.read("xl/worksheets/sheet1.xml").decode()
        assert "<drawing r:id=" in sheet
        ET.fromstring(sheet)          # 名前空間が解決できること


# ---------------------------------------------------------------- 表示単位
@pytest.mark.parametrize("unit", ["day", "week", "month"])
def test_every_unit_exports(filled_book, tmp_path, unit):
    imported = read(filled_book)
    imported.spec.unit = unit
    path = tmp_path / f"{unit}.xlsx"
    export(imported, path, base_date=BASE)
    assert "plan-101" in _shapes(path)

    ws = openpyxl.load_workbook(path)[SHEET_PLAN]
    first_row = 6 if unit == "day" else 5
    assert ws[f"E{first_row}"].value == "要件定義"


# ---------------------------------------------------------------- 往復
def test_the_exported_file_can_be_read_back(filled_book, exported):
    """書き出したファイルをもう一度読み込んでも、同じ内容・同じ期間になる。"""
    before = build_chart(read(filled_book), base_date=BASE)
    after = build_chart(read(exported), base_date=BASE)

    def summary(model):
        return [(r["no"], r["name"], r["start"], r["end"], r["actual_start"],
                 r["progress"], r["member"], r["status"], r["group"], r["subgroup"])
                for r in model["rows"]]

    assert summary(after) == summary(before)
    assert after["totals"] == before["totals"]
    assert after["timeline"]["start"] == before["timeline"]["start"]
