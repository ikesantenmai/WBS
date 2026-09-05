"""記入済み Excel の読み込みのテスト。"""

import datetime as dt
import time
import zipfile

import openpyxl
import pytest

from wbsgen.blank import SpecError, build
from wbsgen.importer import read
from wbsgen.workbook import SHEET_PLAN, write


# ---------------------------------------------------------------- 基本
def test_reads_the_rows(filled_book):
    imported = read(filled_book, "filled.xlsx")
    assert [r.name for r in imported.rows] == [
        "要件定義", "基本設計", "詳細設計", "コーディング", "結合テスト"]
    assert imported.warnings == []


def test_reads_the_values(filled_book):
    row = read(filled_book).rows[0]
    assert row.no == "101"
    assert row.name == "要件定義"
    assert row.start == dt.date(2026, 4, 1)
    assert row.days == 10
    assert row.end == dt.date(2026, 4, 14)
    assert row.actual_start == dt.date(2026, 4, 1)
    assert row.actual_end == dt.date(2026, 4, 14)
    assert row.progress == 1.0
    assert row.effort == 10
    assert row.member == "設計"
    assert row.status == "完了"


def test_group_and_subgroup_carry_down(filled_book):
    """大項目・中項目は変わった行にだけ書かれるので、下の行へ引き継ぐ。"""
    rows = read(filled_book).rows
    assert [(r.group, r.subgroup) for r in rows] == [
        ("開発", "要件"), ("開発", "要件"), ("開発", "製造"),
        ("開発", "製造"), ("テスト", "結合"),
    ]


def test_blank_rows_are_skipped(filled_book):
    """記入されていない空行は読み飛ばす。"""
    assert len(read(filled_book).rows) == 5      # 空行は 30 行ある


def test_title_comes_from_the_sheet(filled_book):
    assert read(filled_book).title == "2026年度 開発スケジュール"


# ---------------------------------------------------------------- 設定
def test_period_comes_from_the_config_sheet(filled_book):
    spec = read(filled_book).spec
    assert spec.start == dt.date(2026, 4, 1)
    assert spec.period_days == 365
    assert spec.unit == "week"


def test_members_come_from_the_member_sheet(filled_book):
    assert read(filled_book).spec.members == ["設計", "製造", "テスト"]


def test_holidays_come_from_the_config_sheet(filled_book):
    spec = read(filled_book).spec
    assert dt.date(2026, 4, 29) in spec.holidays      # 昭和の日
    assert not spec.calendar().is_workday(dt.date(2026, 4, 29))


def test_workdays_come_from_the_config_sheet(make_filled):
    path = make_filled("sat.xlsx", workdays=["mon", "tue", "wed", "thu", "fri", "sat"])
    spec = read(path).spec
    assert spec.workdays == ["mon", "tue", "wed", "thu", "fri", "sat"]
    assert spec.calendar().is_workday(dt.date(2026, 4, 18))   # 土


def test_day_unit_skips_the_weekday_row(make_filled):
    """日単位のファイルは見出しの下に曜日の行があるぶん、開始行がずれる。"""
    path = make_filled("day.xlsx", unit="day", end=None, months=6)
    imported = read(path)
    assert imported.spec.unit == "day"
    assert [r.name for r in imported.rows][:2] == ["要件定義", "基本設計"]


# ---------------------------------------------------------------- 列の探索
def test_columns_are_found_by_their_labels(make_filled):
    """列を入れ替えても見出しで探すので読める。"""
    path = make_filled("moved.xlsx")
    book = openpyxl.load_workbook(path)
    sheet = book[SHEET_PLAN]
    sheet.insert_cols(2)                     # 表の左に列を 1 本足す
    book.save(path)

    rows = read(path).rows
    assert [r.name for r in rows][:2] == ["要件定義", "基本設計"]
    assert rows[0].start == dt.date(2026, 4, 1)


def test_plan_and_actual_are_told_apart(filled_book):
    """「開始」「日数」は予定と実績の両方にあるので、上の帯で見分ける。"""
    row = read(filled_book).rows[1]
    assert (row.start, row.end) == (dt.date(2026, 4, 15), dt.date(2026, 5, 7))
    assert (row.actual_start, row.actual_end) == (dt.date(2026, 4, 15), dt.date(2026, 5, 12))


# ---------------------------------------------------------------- 値の解釈
@pytest.mark.parametrize("written,expected", [
    (0.8, 0.8), ("80%", 0.8), (80, 0.8), ("0.25", 0.25), (1, 1.0), (150, 1.0),
])
def test_progress_accepts_several_notations(make_filled, written, expected):
    path = make_filled("p.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 5, None,
         None, None, None, written, None, "", "設計", ""),
    ])
    assert read(path).rows[0].progress == expected


def test_days_written_with_a_unit_is_read(make_filled):
    path = make_filled("d.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), "5 日", None,
         None, None, None, None, None, "", "設計", ""),
    ])
    assert read(path).rows[0].days == 5


def test_an_unreadable_cell_does_not_drop_the_row(make_filled):
    """1 か所の書き間違いで行がまるごと消えないこと。"""
    path = make_filled("bad.xlsx")
    book = openpyxl.load_workbook(path)
    book[SHEET_PLAN]["F6"] = "来週くらい"       # 予定開始日
    book.save(path)

    imported = read(path)
    assert len(imported.rows) == 5              # 行は残る
    broken = imported.rows[1]
    assert broken.name == "基本設計"
    assert broken.start is None                 # 読めなかった項目だけ空
    assert broken.end == dt.date(2026, 5, 7)    # 他の項目は生きている
    assert any("6 行目" in w and "開始日" in w for w in imported.warnings)


# ---------------------------------------------------------------- 雑な記入
@pytest.mark.parametrize("column,written", [
    ("L", "-"), ("L", "ー"), ("L", "―"),        # 遅れに「無し」を表す記号
    ("G", "-"), ("J", "-"),                     # 日数に記号
    ("M", "-"),                                 # 進捗に記号
    ("K", "未"), ("K", "未定"), ("K", "なし"),    # 実績終了日に「まだ」
    ("K", "   "),                               # 空白だけ
])
def test_marks_that_mean_nothing_are_read_as_blank(make_filled, column, written):
    """手書きの表でよく出る「無し」の書き方は、空として扱う。"""
    path = make_filled("blank-mark.xlsx")
    book = openpyxl.load_workbook(path)
    book[SHEET_PLAN][f"{column}5"] = written
    book.save(path)

    imported = read(path)
    assert len(imported.rows) == 5
    assert imported.warnings == []


@pytest.mark.parametrize("written,expected", [
    ("８０％", 0.8), ("80%", 0.8), ("０．５", 0.5), ("100", 1.0),
])
def test_full_width_numbers_are_understood(make_filled, written, expected):
    """日本語入力では全角の数字や％がよく混ざる。"""
    path = make_filled("zenkaku.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 5, dt.date(2026, 4, 7),
         None, None, None, None, None, "", "設計", ""),
    ])
    book = openpyxl.load_workbook(path)
    book[SHEET_PLAN]["M5"] = written
    book.save(path)
    assert read(path).rows[0].progress == expected


@pytest.mark.parametrize("written,expected", [
    ("2026-04-10", (2026, 4, 10)),
    ("2026/4/10", (2026, 4, 10)),
    ("2026.4.10", (2026, 4, 10)),
    ("2026年4月10日", (2026, 4, 10)),
    ("２０２６/４/１０", (2026, 4, 10)),
])
def test_dates_written_in_several_ways_are_understood(make_filled, written, expected):
    path = make_filled("dates.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 5, dt.date(2026, 4, 7),
         dt.date(2026, 4, 1), None, None, None, None, "", "設計", ""),
    ])
    book = openpyxl.load_workbook(path)
    book[SHEET_PLAN]["K5"] = written               # 実績終了日
    book.save(path)

    row = read(path).rows[0]
    assert row.actual_end == dt.date(*expected)
    assert row.progress == 1.0                     # 終了日が読めれば 100%


@pytest.mark.parametrize("written,expected", [
    ("10 日", 10), ("10日", 10), ("１０日", 10), ("10 days", 10), ("10d", 10),
])
def test_days_with_a_unit_are_understood(make_filled, written, expected):
    path = make_filled("units.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), None, None,
         None, None, None, None, None, "", "設計", ""),
    ])
    book = openpyxl.load_workbook(path)
    book[SHEET_PLAN]["G5"] = written
    book.save(path)
    assert read(path).rows[0].days == expected


def test_a_row_stays_finished_even_if_another_cell_is_unreadable(make_filled):
    """進捗欄が読めなくても、実績終了日があれば 100% になる。"""
    path = make_filled("mixed.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 5, dt.date(2026, 4, 7),
         dt.date(2026, 4, 1), None, dt.date(2026, 4, 6), None, None, "", "設計", ""),
    ])
    book = openpyxl.load_workbook(path)
    book[SHEET_PLAN]["M5"] = "済"                  # 進捗に文字を書いてしまった
    book.save(path)

    imported = read(path)
    assert imported.rows[0].progress == 1.0
    assert any("進捗" in w for w in imported.warnings)


# ---------------------------------------------------------------- 異常系
def test_a_blank_workbook_has_no_rows(tmp_path):
    path = tmp_path / "blank.xlsx"
    write(build(dt.date(2026, 4, 1), months=3, rows=10), path)
    assert read(path).rows == []


def test_a_sheet_without_the_item_column_is_rejected(tmp_path):
    book = openpyxl.Workbook()
    book.active["A1"] = "なにか別の表"
    path = tmp_path / "other.xlsx"
    book.save(path)
    with pytest.raises(SpecError, match="見出し行が見つかりません"):
        read(path)


def test_a_non_excel_file_is_rejected(tmp_path):
    path = tmp_path / "not.xlsx"
    path.write_bytes(b"this is not a workbook")
    with pytest.raises(SpecError, match="Excel として読めません"):
        read(path)


# ---------------------------------------------------------------- 導出
def test_planned_days_are_counted_from_the_dates(make_filled):
    """予定の日数は、書かれた値ではなく開始日と終了日から数える。"""
    path = make_filled("d.xlsx", rows=[
        # 4/1〜4/14 は稼働日 10 日だが、わざと 99 と書いておく
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 99, dt.date(2026, 4, 14),
         None, None, None, None, None, "", "設計", ""),
    ])
    row = read(path).rows[0]
    assert row.days == 10
    assert "days" in row.derived


def test_actual_days_are_counted_from_the_dates(make_filled):
    path = make_filled("ad.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 10, dt.date(2026, 4, 14),
         dt.date(2026, 4, 1), 99, dt.date(2026, 4, 10), None, None, "", "設計", ""),
    ])
    row = read(path).rows[0]
    assert row.actual_days == 8          # 4/1〜4/10 の稼働日
    assert "actual_days" in row.derived


def test_an_actual_end_date_means_one_hundred_percent(make_filled):
    """実績の終了日が入っていれば、書かれた進捗より優先して 100% とする。"""
    path = make_filled("p.xlsx", rows=[
        ("開発", "", "1", "済", dt.date(2026, 4, 1), 10, dt.date(2026, 4, 14),
         dt.date(2026, 4, 1), None, dt.date(2026, 4, 10), 0.3, None, "", "設計", ""),
        ("開発", "", "2", "途中", dt.date(2026, 4, 1), 10, dt.date(2026, 4, 14),
         dt.date(2026, 4, 1), 5, None, 0.3, None, "", "設計", ""),
    ])
    done, running = read(path).rows
    assert done.progress == 1.0
    assert "progress" in done.derived
    assert running.progress == 0.3       # 終了日が無ければ触らない
    assert "progress" not in running.derived


def test_actual_days_without_an_end_date_are_cleared(make_filled):
    """実績の終了日が無いのに日数だけ残っている行は、日数を捨てる。

    終わっていない作業に残った日数は元データの書き間違いなので、
    そこから実績の終了日を作らない。
    """
    path = make_filled("r.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 10, dt.date(2026, 4, 14),
         dt.date(2026, 4, 1), 5, None, 0.5, None, "", "設計", ""),
    ])
    row = read(path).rows[0]
    assert row.actual_days is None
    assert "actual_days" in row.derived
    assert row.actual_end is None                    # 終了日は作らない
    assert row.actual_start == dt.date(2026, 4, 1)   # 着手済みなのは残る
    assert row.progress == 0.5                       # 完了扱いにはしない


def test_a_missing_end_date_is_derived_from_the_days(make_filled):
    path = make_filled("e.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 10, None,
         None, None, None, None, None, "", "設計", ""),
    ])
    row = read(path).rows[0]
    assert row.end == dt.date(2026, 4, 14)
    assert "end" in row.derived


def test_days_that_already_match_are_not_marked_as_derived(make_filled):
    path = make_filled("m.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 10, dt.date(2026, 4, 14),
         None, None, None, None, None, "", "設計", ""),
    ])
    # 状態は必ず計算するので、日数だけを見る
    assert "days" not in read(path).rows[0].derived


def test_days_are_cleared_when_they_cannot_be_counted(make_filled):
    """開始日が無く数え直せない行は、書かれた日数を残さず空にする。

    日数は開始日と終了日から数え直すものなので、片方しか無い行に
    書かれた数字は当てにできない。
    """
    path = make_filled("c1.xlsx", rows=[
        ("開発", "", "1", "開始日なし", None, 12, dt.date(2026, 4, 14),
         None, 12, dt.date(2026, 4, 14), None, None, "", "設計", ""),
    ])
    row = read(path, base_date=BASE).rows[0]
    assert row.days is None
    assert row.actual_days is None
    assert "days" in row.derived
    assert "actual_days" in row.derived


def test_days_alone_without_any_date_are_cleared(make_filled):
    path = make_filled("c2.xlsx", rows=[
        ("開発", "", "1", "日数だけ", None, 7, None,
         None, 7, None, None, None, "", "設計", ""),
    ])
    row = read(path, base_date=BASE).rows[0]
    assert row.days is None
    assert row.actual_days is None


def test_only_the_plan_derives_a_missing_end_date(make_filled):
    """終了日を日数から補うのは予定だけ。実績は補わず日数を捨てる。"""
    path = make_filled("c3.xlsx", rows=[
        ("開発", "", "1", "終了日なし", dt.date(2026, 4, 1), 10, None,
         dt.date(2026, 4, 1), 5, None, None, None, "", "設計", ""),
    ])
    row = read(path, base_date=BASE).rows[0]
    assert row.end == dt.date(2026, 4, 14)
    assert row.days == 10                # 補った終了日から数えても同じ
    assert row.actual_end is None
    assert row.actual_days is None


def test_recounted_days_are_written_out(make_filled, tmp_path):
    """書き出した Excel の日数欄にも、数え直した値が入る。"""
    from wbsgen.workbook import export

    path = make_filled("c4.xlsx", rows=[
        ("開発", "", "1", "でたらめな日数", dt.date(2026, 4, 1), 99,
         dt.date(2026, 4, 14), dt.date(2026, 4, 1), 99, dt.date(2026, 4, 10),
         None, None, "", "設計", ""),
        ("開発", "", "2", "数え直せない", None, 12, dt.date(2026, 4, 14),
         None, 12, dt.date(2026, 4, 14), None, None, "", "設計", ""),
    ])
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)

    sheet = openpyxl.load_workbook(out)[SHEET_PLAN]
    assert (sheet["G5"].value, sheet["J5"].value) == (10, 8)
    assert (sheet["G6"].value, sheet["J6"].value) == (None, None)


def test_the_working_calendar_is_used_for_counting(make_filled):
    """稼働曜日の設定が数え方に効く (土曜も稼働にすると日数が増える)。"""
    weekdays = ["mon", "tue", "wed", "thu", "fri", "sat"]
    rows = [("開発", "", "1", "A", dt.date(2026, 4, 1), None, dt.date(2026, 4, 14),
             None, None, None, None, None, "", "設計", "")]
    assert read(make_filled("w1.xlsx", rows=rows)).rows[0].days == 10
    assert read(make_filled("w2.xlsx", rows=rows, workdays=weekdays)).rows[0].days == 12


# ---------------------------------------------------------------- 足したシート
def _add_sheet(path, title, at=None, fill=None):
    book = openpyxl.load_workbook(path)
    sheet = book.create_sheet(title) if at is None else book.create_sheet(title, at)
    if fill:
        fill(sheet)
    book.save(path)
    return path


def test_an_extra_sheet_does_not_disturb_reading(make_filled):
    """足したシートがあっても、日程表はこれまでどおり読める。"""
    path = _add_sheet(make_filled("extra.xlsx"), "メモ", at=0,
                      fill=lambda s: s.__setitem__("A1", "打合せメモ"))
    imported = read(path, base_date=BASE)
    assert [r.name for r in imported.rows][:2] == ["要件定義", "基本設計"]
    # 足してあるシートがあるので、元のファイルを書き出しの土台として控える
    assert imported.source == path.read_bytes()


def test_an_extra_sheet_is_written_out_again(make_filled, tmp_path):
    """足したシートは、書き出したファイルにもそのまま残る。"""
    from wbsgen.workbook import export

    def fill(sheet):
        sheet["A1"] = "課題一覧"
        sheet["A1"].font = openpyxl.styles.Font(bold=True, size=14)
        sheet["A1"].fill = openpyxl.styles.PatternFill("solid", fgColor="FFFF00")
        sheet.merge_cells("A1:C1")
        sheet["B2"] = "=COUNTA(A4:A100)"
        sheet["A4"] = dt.date(2026, 8, 20)
        sheet["A4"].number_format = "yyyy/mm/dd"
        sheet.column_dimensions["A"].width = 28
        sheet.freeze_panes = "A4"

    path = _add_sheet(make_filled("extra2.xlsx"), "課題管理", fill=fill)
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)

    book = openpyxl.load_workbook(out)
    # 元のファイルの並びのまま、このツールが作るシートだけが差し替わる
    assert book.sheetnames[:3] == [SHEET_PLAN, "担当者一覧", "設定"]
    assert book.sheetnames[-1] == "課題管理"
    sheet = book["課題管理"]
    assert sheet["A1"].value == "課題一覧"
    assert sheet["A1"].font.b and sheet["A1"].font.sz == 14
    assert sheet["A1"].fill.fgColor.rgb.endswith("FFFF00")
    assert [str(r) for r in sheet.merged_cells.ranges] == ["A1:C1"]
    assert sheet["B2"].value == "=COUNTA(A4:A100)"      # 数式のまま残る
    assert sheet["A4"].number_format == "yyyy/mm/dd"
    assert sheet.column_dimensions["A"].width == 28
    assert sheet.freeze_panes == "A4"


def test_the_gantt_shapes_still_land_on_the_schedule_sheet(make_filled, tmp_path):
    """シートが増えて日程表が 1 枚目でなくなっても、図形はそこに入る。"""
    from wbsgen.inject import sheet_part
    from wbsgen.workbook import export

    path = _add_sheet(make_filled("extra3.xlsx"), "メモ", at=0,
                      fill=lambda s: s.__setitem__("A1", "x"))
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)

    # 足してあったシートの位置はそのまま (メモが 1 枚目)
    assert openpyxl.load_workbook(out).sheetnames[0] == "メモ"
    part = sheet_part(out, SHEET_PLAN)
    assert part != "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(out) as archive:
        assert "<drawing" in archive.read(part).decode()
        assert "<drawing" not in archive.read("xl/worksheets/sheet1.xml").decode()


def test_a_renamed_schedule_sheet_is_still_found(make_filled):
    """日程表のシート名を変えてあっても、「項目」の列で見つける。"""
    path = make_filled("renamed.xlsx")
    book = openpyxl.load_workbook(path)
    book[SHEET_PLAN].title = "ITb 工程表"
    cover = book.create_sheet("表紙", 0)
    cover["A1"] = "外部結合テスト"
    book.save(path)

    imported = read(path, base_date=BASE)
    assert [r.name for r in imported.rows][:1] == ["要件定義"]
    assert imported.source == path.read_bytes()


def test_every_sheet_from_the_import_is_written_out(make_filled, tmp_path):
    """読み込んだときにあったシートは、名前を変えてあるものも含めて残す。"""
    from wbsgen.workbook import export

    path = make_filled("keep-all.xlsx")
    book = openpyxl.load_workbook(path)
    book[SHEET_PLAN].title = "ITb 工程表"          # 日程表の名前を変えてある
    book.create_sheet("表紙", 0)["A1"] = "外部結合テスト"
    book.create_sheet("メモ")["A1"] = "打合せ"
    book.create_sheet("要員稼働チェック_12月")["A1"] = "手作りの表"
    before = book.sheetnames
    book.save(path)

    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)
    after = openpyxl.load_workbook(out).sheetnames
    # 読み込み時にあった名前がすべて残っている
    assert set(before) <= set(after), set(before) - set(after)
    # このツールが作るぶんは足されている
    assert {SHEET_PLAN, "担当者一覧", "設定"} <= set(after)
    # 作り直さないシートの中身はそのまま
    assert openpyxl.load_workbook(out)["要員稼働チェック_12月"]["A1"].value == "手作りの表"


def test_a_chart_and_an_image_on_an_extra_sheet_survive(make_filled, tmp_path):
    """足したシートのグラフ・画像・条件付き書式などがそのまま残る。

    書き出しは元のファイルを土台にして 3 シートだけ差し替えるので、
    足したシートには手が入らない。
    """
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from wbsgen.workbook import export

    def fill(sheet):
        sheet["A1"], sheet["B1"] = "項目", "件数"
        for i, (name, count) in enumerate([("設計", 5), ("製造", 12)], 2):
            sheet[f"A{i}"], sheet[f"B{i}"] = name, count
        chart = BarChart()
        chart.title = "工程別件数"
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3),
                       titles_from_data=True)
        sheet.add_chart(chart, "D2")
        sheet.conditional_formatting.add("B2:B3", CellIsRule(
            operator="greaterThan", formula=["10"],
            fill=PatternFill(bgColor="FFC7CE")))
        rule = DataValidation(type="list", formula1='"未,済"')
        sheet.add_data_validation(rule)
        rule.add("C2:C3")
        sheet.auto_filter.ref = "A1:C3"
        sheet.sheet_properties.tabColor = "FF9900"

    path = _add_sheet(make_filled("chart.xlsx"), "集計", fill=fill)
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)

    with zipfile.ZipFile(out) as archive:
        assert "xl/charts/chart1.xml" in archive.namelist()

    sheet = openpyxl.load_workbook(out)["集計"]
    assert len(list(sheet.conditional_formatting)) == 1
    assert len(sheet.data_validations.dataValidation) == 1
    assert sheet.auto_filter.ref == "A1:C3"
    assert sheet.sheet_properties.tabColor.rgb.endswith("FF9900")


def test_a_file_with_no_schedule_sheet_is_still_rejected(tmp_path):
    book = openpyxl.Workbook()
    book.active["A1"] = "表紙"
    book.create_sheet("メモ")["A1"] = "なにか"
    path = tmp_path / "no-plan.xlsx"
    book.save(path)
    with pytest.raises(SpecError, match="見出し行が見つかりません"):
        read(path)


# ---------------------------------------------------------------- 文字色
def _paint(path, cells, sheet=SHEET_PLAN):
    book = openpyxl.load_workbook(path)
    target = book[sheet]
    for coordinate, color in cells.items():
        cell = target[coordinate]
        cell.font = openpyxl.styles.Font(
            name=cell.font.name, size=cell.font.sz, color=color)
    book.save(path)
    return path


def test_the_font_colour_of_each_cell_is_read(make_filled):
    """記入した文字色を控える (書き出しでそのまま戻すため)。"""
    path = _paint(make_filled("color.xlsx"),
                  {"E5": "FFFF0000", "Q6": "FF0070C0", "F7": "FF00B050"})
    rows = read(path, base_date=BASE).rows
    assert rows[0].colors["name"] == "FFFF0000"
    assert rows[1].colors["member"] == "FF0070C0"
    assert rows[2].colors["start"] == "FF00B050"


def test_the_font_colour_survives_the_export(make_filled, tmp_path):
    """読み込んだ文字色を、書き出しで既定色に変えてしまわない。"""
    from wbsgen.workbook import export

    path = _paint(make_filled("color2.xlsx"),
                  {"E5": "FFFF0000", "Q6": "FF0070C0", "F7": "FF00B050"})
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)

    sheet = openpyxl.load_workbook(out)[SHEET_PLAN]
    assert sheet["E5"].font.color.rgb == "FFFF0000"
    assert sheet["Q6"].font.color.rgb == "FF0070C0"
    assert sheet["F7"].font.color.rgb == "FF00B050"


def test_cells_without_a_colour_keep_the_usual_one(make_filled, tmp_path):
    """色を指定していないセルは、これまでどおりの色で書く。"""
    from wbsgen.workbook import export

    out = export(read(make_filled("plain.xlsx"), base_date=BASE),
                 tmp_path / "out.xlsx", BASE)
    sheet = openpyxl.load_workbook(out)[SHEET_PLAN]
    assert sheet["F5"].font.color.rgb[-6:] == "000080"    # 予定は紺
    assert sheet["E5"].font.color.rgb[-6:] == "000000"    # ほかは黒


def test_the_status_column_uses_the_calculated_colour(make_filled, tmp_path):
    """状態の列だけは例外。文字そのものを計算し直すので、色も計算に合わせる。

    文字色をそのまま使うのは、大項目から担当まで (B〜Q 列)。
    """
    from wbsgen.workbook import export

    path = _paint(make_filled("status-color.xlsx"), {"R5": "FF7030A0"})
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)
    sheet = openpyxl.load_workbook(out)[SHEET_PLAN]
    assert sheet["R5"].value == "完了"
    assert sheet["R5"].font.color.rgb[-6:] != "7030A0"


def test_a_derived_value_keeps_the_written_colour(make_filled, tmp_path):
    """導き出した値のセルでも、記入した文字色から変えない。"""
    from wbsgen.workbook import export

    path = make_filled("derived-color.xlsx", rows=[
        # 終了日を書いていないので、日数から補う (= 導出)
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 10, None,
         None, None, None, None, None, "", "設計", ""),
    ])
    _paint(path, {"H5": "FF7030A0"})
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)

    sheet = openpyxl.load_workbook(out)[SHEET_PLAN]
    assert sheet["H5"].value == dt.datetime(2026, 4, 14)      # 補った値
    assert sheet["H5"].font.color.rgb == "FF7030A0"           # 色はそのまま


def test_no_cell_changes_colour_through_a_round_trip(make_filled, tmp_path):
    """記入済みの表を通したとき、文字色が変わるセルが 1 つも無いこと。"""
    from wbsgen.workbook import export

    path = _paint(make_filled("keep-all-colors.xlsx"), {
        "B5": "FF7030A0", "E5": "FFFF0000", "F5": "FF00B050", "G6": "FF0070C0",
        "K7": "FFFFC000", "M7": "FF7030A0", "O8": "FF00B0F0", "Q8": "FFFF00FF",
    })
    imported = read(path, base_date=BASE)
    out = export(imported, tmp_path / "out.xlsx", BASE)

    before = openpyxl.load_workbook(path)[SHEET_PLAN]
    after = openpyxl.load_workbook(out)[SHEET_PLAN]

    def colour(cell):
        found = cell.font.color
        if found is None or found.type != "rgb" or not isinstance(found.rgb, str):
            return None
        return found.rgb[-6:]                    # 透明度のバイトは見ない

    first = 5                                    # 週単位は見出しのすぐ下から
    changed = [
        (before.cell(row=row.row, column=column).coordinate, was, now)
        for index, row in enumerate(imported.rows)
        for column in range(2, 18)               # 大項目 (B) 〜 担当 (Q)
        for was, now in [(colour(before.cell(row=row.row, column=column)),
                          colour(after.cell(row=first + index, column=column)))]
        if was != now
    ]
    assert changed == []


def test_the_colour_survives_two_round_trips(make_filled, tmp_path):
    """書き出したものを読み直しても、色は変わらない。"""
    from wbsgen.workbook import export

    path = _paint(make_filled("color3.xlsx"), {"E5": "FFFF0000"})
    once = export(read(path, base_date=BASE), tmp_path / "a.xlsx", BASE)
    twice = export(read(once, base_date=BASE), tmp_path / "b.xlsx", BASE)
    assert openpyxl.load_workbook(twice)[SHEET_PLAN]["E5"].font.color.rgb == "FFFF0000"


# ---------------------------------------------------------------- 項目名の空白
def test_the_indent_in_a_task_name_is_kept(make_filled):
    """項目名の行頭の空白 (階層を表す字下げ) をそのまま読む。"""
    path = make_filled("indent.xlsx", rows=[
        ("開発", "", "1", "テスト実施", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), None, None, None, None, None, "", "設計", ""),
        ("", "", "2", "\u3000WEB口座開設システム_画像管理", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), None, None, None, None, None, "", "設計", ""),
        ("", "", "3", "  半角スペースでも同じ", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), None, None, None, None, None, "", "設計", ""),
    ])
    names = [r.name for r in read(path, base_date=BASE).rows]
    assert names == ["テスト実施", "\u3000WEB口座開設システム_画像管理",
                     "  半角スペースでも同じ"]


def test_a_name_of_only_spaces_is_still_empty(make_filled):
    """空白だけの項目名は、これまでどおり空として扱う。"""
    path = make_filled("spaces.xlsx", rows=[
        ("開発", "", "1", "\u3000\u3000  ", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), None, None, None, None, None, "", "設計", ""),
    ])
    assert read(path, base_date=BASE).rows[0].name == ""


def test_the_indent_survives_a_round_trip(make_filled, tmp_path):
    """書き出した Excel にも字下げが残る。"""
    from wbsgen.workbook import export

    path = make_filled("indent2.xlsx", rows=[
        ("開発", "", "1", "\u3000字下げした項目", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), None, None, None, None, None, "", "設計", ""),
    ])
    out = export(read(path, base_date=BASE), tmp_path / "out.xlsx", BASE)

    assert openpyxl.load_workbook(out)[SHEET_PLAN]["E5"].value == "\u3000字下げした項目"
    assert read(out, base_date=BASE).rows[0].name == "\u3000字下げした項目"


# ------------------------------------------------- 実績の終了日は必ず完了にする
def _write_plain_number(path, cell, value):
    """書式を「標準」にしたうえで数値を入れる (日付書式が付いていない状態)。"""
    book = openpyxl.load_workbook(path)
    target = book[SHEET_PLAN][cell]
    target.value = value
    target.number_format = "General"
    book.save(path)


def test_a_date_serial_number_is_read_as_a_date(make_filled):
    """書式が「標準」のままの日付 (シリアル値) も日付として読む。

    そのままだと数値として読めず空になり、完了にならない。
    """
    path = make_filled("serial.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 10, dt.date(2026, 4, 14),
         dt.date(2026, 4, 1), None, None, 0.3, None, "", "設計", "実行中"),
    ])
    _write_plain_number(path, "K5", 46173)   # 2026-05-31 のシリアル値

    imported = read(path, base_date=BASE)
    row = imported.rows[0]
    assert row.actual_end == dt.date(2026, 5, 31)
    assert row.progress == 1.0
    assert row.status == "完了"
    assert imported.warnings == []


def test_a_number_that_is_not_a_date_is_still_reported(make_filled):
    path = make_filled("serial2.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 1), 10, dt.date(2026, 4, 14),
         dt.date(2026, 4, 1), None, None, 0.3, None, "", "設計", ""),
    ])
    _write_plain_number(path, "K5", 5)       # 日付にしては小さすぎる

    imported = read(path, base_date=BASE)
    assert imported.rows[0].actual_end is None
    assert any("5 行目" in w for w in imported.warnings), imported.warnings


def test_a_row_with_only_an_actual_end_date_is_kept(make_filled):
    """項目名も予定も無くても、実績が入っていれば行を残して完了にする。"""
    path = make_filled("only-actual.xlsx", rows=[
        ("開発", "", "", "", None, None, None,
         None, None, dt.date(2026, 5, 1), None, None, "", "設計", ""),
    ])
    rows = read(path, base_date=BASE).rows
    assert len(rows) == 1
    assert rows[0].progress == 1.0
    assert rows[0].status == "完了"


def test_a_late_finish_is_still_done(make_filled):
    """予定の終了日より後に終わった行も完了とする。"""
    path = make_filled("late.xlsx", rows=[
        ("開発", "", "1", "遅れて完了", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), dt.date(2026, 4, 1), None, dt.date(2026, 9, 30),
         0.1, None, "", "設計", "遅れ 30 日"),
    ])
    row = read(path, base_date=BASE).rows[0]
    assert row.progress == 1.0
    assert row.status == "完了"
    assert row.delay is None


# ---------------------------------------------------------------- 状態
BASE = dt.date(2026, 6, 10)


def _one(make_filled, name, **cells):
    defaults = dict(start=dt.date(2026, 6, 1), days=10, end=dt.date(2026, 6, 12),
                    actual_start=None, actual_days=None, actual_end=None,
                    progress=None, status="")
    defaults.update(cells)
    path = make_filled(f"s-{name}.xlsx", rows=[(
        "開発", "", "1", name, defaults["start"], defaults["days"], defaults["end"],
        defaults["actual_start"], defaults["actual_days"], defaults["actual_end"],
        defaults["progress"], None, "", "設計", defaults["status"],
    )])
    return read(path, base_date=BASE).rows[0]


def test_a_finished_row_is_marked_done(make_filled):
    row = _one(make_filled, "done", actual_start=dt.date(2026, 6, 1),
               actual_end=dt.date(2026, 6, 8))
    assert row.status == "完了"
    assert row.delay is None


def test_a_started_row_counts_down_to_the_planned_end(make_filled):
    row = _one(make_filled, "remaining", actual_start=dt.date(2026, 6, 1))
    assert row.status == "残り 2 日"      # 6/10 から 6/12 まで


def test_an_unstarted_row_counts_down_to_the_planned_start(make_filled):
    row = _one(make_filled, "upcoming", start=dt.date(2026, 6, 22), days=5,
               end=dt.date(2026, 6, 26))
    assert row.status == "あと 8 日"


def test_a_row_past_its_planned_end_is_delayed(make_filled):
    row = _one(make_filled, "end-late", start=dt.date(2026, 5, 1), days=5,
               end=dt.date(2026, 5, 7), actual_start=dt.date(2026, 5, 1))
    assert row.delay == 24
    assert row.status == "遅れ 24 日"


def test_a_row_that_should_have_started_is_delayed(make_filled):
    """予定開始日を過ぎているのに未着手なら「開始遅れ」。"""
    row = _one(make_filled, "start-late", start=dt.date(2026, 6, 1), days=20,
               end=dt.date(2026, 6, 26))
    assert row.delay == 7                 # 6/1 から 6/10 まで
    assert row.status == "遅れ 7 日"


def test_the_countdown_is_zero_on_the_day_itself(make_filled):
    row = _one(make_filled, "today", actual_start=dt.date(2026, 6, 1),
               end=BASE, days=8)
    assert row.status == "残り 0 日"


def test_a_row_without_dates_keeps_the_written_status(make_filled):
    row = _one(make_filled, "kept", start=None, days=None, end=None, status="保留")
    assert row.status == "保留"


def test_the_status_follows_the_base_date(make_filled):
    """基準日を変えると状態も変わる (何度計算しても壊れない)。"""
    from wbsgen.importer import resolve

    path = make_filled("moving.xlsx", rows=[
        ("開発", "", "1", "A", dt.date(2026, 6, 1), 10, dt.date(2026, 6, 12),
         dt.date(2026, 6, 1), None, None, None, None, "", "設計", ""),
    ])
    imported = read(path, base_date=BASE)
    calendar = imported.spec.calendar()

    seen = []
    for day in (dt.date(2026, 6, 1), dt.date(2026, 6, 12), dt.date(2026, 7, 1)):
        resolve(imported.rows, calendar, day)
        seen.append(imported.rows[0].status)
    assert seen == ["残り 9 日", "残り 0 日", "遅れ 13 日"]

    # 同じ基準日で数え直しても結果は変わらない
    resolve(imported.rows, calendar, BASE)
    first = imported.rows[0].status
    resolve(imported.rows, calendar, BASE)
    assert imported.rows[0].status == first


@pytest.mark.parametrize("language,expected", [
    ("ja", "完了"), ("en", "Done"),
])
def test_the_status_is_written_in_the_language(make_filled, language, expected):
    path = make_filled(f"lang-{language}.xlsx", language=language, rows=[
        ("開発", "", "1", "A", dt.date(2026, 6, 1), 5, dt.date(2026, 6, 5),
         dt.date(2026, 6, 1), None, dt.date(2026, 6, 5), None, None, "", "設計", ""),
    ])
    assert read(path, language=language, base_date=BASE).rows[0].status == expected


# ---------------------------------------------------------------- 結合セル
def _merge(path, ranges):
    book = openpyxl.load_workbook(path)
    for cells in ranges:
        book[SHEET_PLAN].merge_cells(cells)
    book.save(path)


def test_a_vertically_merged_cell_reaches_every_row(make_filled):
    """縦に結合した列は、下の行からも同じ値が読めること。

    Excel は結合セルの値を左上にしか持たないが、画面では範囲すべてに
    表示される。手作りの表では日付を縦に結合してあることがある。
    """
    path = make_filled("merged.xlsx", rows=[
        ("開発", "", str(i), f"タスク{i}", dt.date(2026, 4, 1), 10, dt.date(2026, 4, 14),
         dt.date(2026, 4, 1), None, dt.date(2026, 4, 10) if i == 1 else None,
         None, None, "", "設計", "")
        for i in (1, 2, 3)
    ])
    _merge(path, ["K5:K7"])          # 実績終了日を 3 行ぶん結合

    rows = read(path, base_date=BASE).rows
    assert len(rows) == 3
    assert all(r.actual_end == dt.date(2026, 4, 10) for r in rows)
    assert all(r.progress == 1.0 for r in rows)
    assert all(r.status == "完了" for r in rows)


def test_merged_group_cells_do_not_create_empty_rows(make_filled):
    """大項目を広く結合してあっても、データの無い行は増やさない。"""
    path = make_filled("merged-group.xlsx", rows=[
        ("開発", "設計", str(i), f"タスク{i}", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), None, None, None, None, None, "", "設計", "")
        for i in (1, 2)
    ])
    _merge(path, ["B5:B20", "C5:C20"])   # データは 2 行しかない

    rows = read(path, base_date=BASE).rows
    assert len(rows) == 2
    assert [r.group for r in rows] == ["開発", "開発"]
    assert [r.subgroup for r in rows] == ["設計", "設計"]


# ---------------------------------------------------------------- 予定が無い行
def test_an_actual_end_date_completes_a_row_without_planned_dates(make_filled):
    """予定の日付が無くても、実績の終了日があれば完了とする。"""
    path = make_filled("no-plan.xlsx", rows=[
        ("開発", "", "1", "実績だけ記録", None, None, None,
         dt.date(2026, 4, 1), None, dt.date(2026, 4, 10), None, None, "", "設計", ""),
    ])
    row = read(path, base_date=BASE).rows[0]
    assert row.progress == 1.0
    assert row.status == "完了"


def test_a_row_with_neither_plan_nor_actual_end_keeps_its_status(make_filled):
    path = make_filled("no-plan2.xlsx", rows=[
        ("開発", "", "1", "保留中", None, None, None,
         None, None, None, None, None, "", "設計", "保留"),
    ])
    row = read(path, base_date=BASE).rows[0]
    assert row.status == "保留"


# ---------------------------------------------------------------- 数式のセル
def _put_formula(path, cell, formula, sheet=SHEET_PLAN):
    book = openpyxl.load_workbook(path)
    book[sheet][cell] = formula
    book.save(path)


def test_a_formula_without_a_stored_result_is_reported(make_filled):
    """数式の計算結果が入っていないセルは、空として扱ったうえで知らせる。

    Excel は数式の計算結果もファイルに残すが、スクリプトで作られた
    ファイルには入っていない。そのまま読むと空に見えるので、記入した
    つもりの実績終了日が反映されず、完了にならない。
    """
    path = make_filled("formula.xlsx", rows=[
        ("開発", "", "1", "数式で終了日", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), dt.date(2026, 4, 1), None, None, None, None,
         "", "設計", ""),
    ])
    _put_formula(path, "K5", "=I5+20")   # 実績の終了日

    imported = read(path, base_date=BASE)
    assert imported.rows[0].actual_end is None
    assert any("5 行目" in w and "数式" in w for w in imported.warnings), \
        imported.warnings


def test_a_formula_warning_is_translated(make_filled):
    path = make_filled("formula-en.xlsx", language="en", rows=[
        ("Dev", "", "1", "formula", dt.date(2026, 4, 1), 10,
         dt.date(2026, 4, 14), dt.date(2026, 4, 1), None, None, None, None,
         "", "Design", ""),
    ])
    _put_formula(path, "K5", "=I5+20", sheet="Schedule")

    imported = read(path, language="en", base_date=BASE)
    assert any("formula" in w for w in imported.warnings), imported.warnings


def test_reading_a_large_sheet_stays_quick(make_filled):
    """行数が増えても読み込みが重くならないこと。

    数式の確認でセルを 1 つずつ引くと、読み取り専用のシートは毎回
    先頭から読み直すため、行数の二乗で遅くなる (実測で 90 行 4 秒)。
    """
    rows = [
        ("開発", "", str(i), f"作業 {i}", dt.date(2026, 4, 1), None, None,
         None, None, None, None, None, "", "設計", "")
        for i in range(400)
    ]
    path = make_filled("big.xlsx", rows=rows)

    started = time.perf_counter()
    imported = read(path, base_date=BASE)
    elapsed = time.perf_counter() - started

    assert len(imported.rows) == 400
    assert elapsed < 5.0, f"読み込みに {elapsed:.1f} 秒かかりました"


# ---------------------------------------------------------------- 列が無い表
def test_a_missing_column_is_reported(filled_book):
    """見出しが違っていて列が見つからないときは、まとめて知らせる。"""
    book = openpyxl.load_workbook(filled_book)
    book[SHEET_PLAN]["K4"] = "おわり"    # 実績の「終了」を別の言葉にする
    book.save(filled_book)

    warnings = read(filled_book, base_date=BASE).warnings
    assert any("終了" in w for w in warnings), warnings


def test_a_complete_sheet_has_no_warnings(filled_book):
    assert read(filled_book, base_date=BASE).warnings == []
