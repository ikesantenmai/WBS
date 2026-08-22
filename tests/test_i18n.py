"""多言語対応のテスト。"""

import datetime as dt

import openpyxl
import pytest

from wbsgen import i18n
from wbsgen.blank import SpecError, build
from wbsgen.chart import build as build_chart
from wbsgen.importer import read
from wbsgen.timeline import Timeline
from wbsgen.workbook import export, write
from wbsgen.workcal import WorkCalendar


# ---------------------------------------------------------------- 基本
def test_japanese_is_the_default():
    assert i18n.DEFAULT_LANGUAGE == "ja"
    assert i18n.normalize(None) == "ja"
    assert i18n.labels().sheet_plan == "スケジュール"


@pytest.mark.parametrize("given,expected", [
    ("ja", "ja"), ("en", "en"), ("JA", "ja"), ("en-US", "en"), ("ja_JP", "ja"),
    ("fr", "ja"), ("", "ja"), (None, "ja"),
])
def test_language_codes_are_normalised(given, expected):
    assert i18n.normalize(given) == expected


def test_both_languages_define_every_message():
    ja, en = i18n.MESSAGES["ja"], i18n.MESSAGES["en"]
    assert set(ja) == set(en)
    assert set(i18n.CLI["ja"]) == set(i18n.CLI["en"])


def test_both_languages_define_every_column():
    assert set(i18n.JA.columns) == set(i18n.EN.columns)
    assert len(i18n.EN.weekdays) == 7
    assert len(i18n.EN.month_names) == 12


@pytest.mark.parametrize("status,kind", [
    ("完了", "done"), ("Done", "done"), ("実行中", "running"), ("In progress", "running"),
    ("遅れ 3 日", "delayed"), ("Delayed 3 d", "delayed"),
    ("あと 2 日", "upcoming"), ("残り 1 日", "remaining"),
    ("保留", None), ("", None), ("-", None),
])
def test_status_kind_reads_both_languages(status, kind):
    assert i18n.status_kind(status) == kind


# ---------------------------------------------------------------- Excel
@pytest.mark.parametrize("language,sheets,headers", [
    ("ja", ["スケジュール", "担当者一覧", "設定"], ["大項目", "中項目", "項番", "項目"]),
    ("en", ["Schedule", "Members", "Settings"], ["Group", "Sub-group", "No.", "Task"]),
])
def test_workbook_is_written_in_the_language(tmp_path, language, sheets, headers):
    path = tmp_path / f"{language}.xlsx"
    write(build(dt.date(2026, 4, 1), months=3, rows=5, language=language), path)

    book = openpyxl.load_workbook(path)
    assert book.sheetnames == sheets
    ws = book[sheets[0]]
    assert [ws[f"{c}4"].value for c in "BCDE"] == headers


@pytest.mark.parametrize("language,plan,actual", [
    ("ja", "予定", "実績"), ("en", "Planned", "Actual"),
])
def test_group_headers_are_translated(tmp_path, language, plan, actual):
    path = tmp_path / "g.xlsx"
    write(build(dt.date(2026, 4, 1), months=2, rows=2, language=language), path)
    ws = openpyxl.load_workbook(path)[i18n.labels(language).sheet_plan]
    assert ws["F3"].value == plan
    assert ws["I3"].value == actual


@pytest.mark.parametrize("language,unit,top,bottom", [
    ("ja", "week", 'm"月"', "m/d"),
    ("ja", "month", 'yyyy"年"', 'm"月"'),
    ("en", "week", "mmm", "m/d"),
    ("en", "month", "yyyy", "mmm"),
])
def test_header_formats_follow_the_language(tmp_path, language, unit, top, bottom):
    path = tmp_path / "f.xlsx"
    write(build(dt.date(2026, 4, 1), months=6, rows=2, unit=unit, language=language), path)
    ws = openpyxl.load_workbook(path)[i18n.labels(language).sheet_plan]
    assert ws["S3"].number_format == top
    assert ws["S4"].number_format == bottom


@pytest.mark.parametrize("language,expected", [
    ("ja", ["水", "木", "金", "土", "日"]),
    ("en", ["Wed", "Thu", "Fri", "Sat", "Sun"]),
])
def test_weekday_row_is_translated(tmp_path, language, expected):
    path = tmp_path / "w.xlsx"
    write(build(dt.date(2026, 4, 1), months=1, rows=2, unit="day", language=language), path)
    ws = openpyxl.load_workbook(path)[i18n.labels(language).sheet_plan]
    assert [ws.cell(row=5, column=c).value for c in range(19, 24)] == expected


@pytest.mark.parametrize("language,label,on", [
    ("ja", "月曜日", "出"), ("en", "Monday", "Work"),
])
def test_config_sheet_is_translated(tmp_path, language, label, on):
    path = tmp_path / "c.xlsx"
    write(build(dt.date(2026, 4, 1), months=2, rows=2, language=language), path)
    ws = openpyxl.load_workbook(path)[i18n.labels(language).sheet_config]
    pairs = {ws[f"B{r}"].value: ws[f"C{r}"].value for r in range(3, 20)}
    assert pairs[label] == on
    assert i18n.labels(language).config_start in pairs


def test_default_title_follows_the_language():
    assert build(dt.date(2026, 4, 1)).title == "2026年 スケジュール"
    assert build(dt.date(2026, 4, 1), language="en").title == "2026 Schedule"


# ---------------------------------------------------------------- 画面
@pytest.mark.parametrize("language,months,years", [
    ("ja", ["4月", "5月"], "2026年"),
    ("en", ["Apr", "May"], "2026"),
])
def test_chart_headings_follow_the_language(language, months, years):
    cal = WorkCalendar.build()
    week = Timeline(dt.date(2026, 4, 1), 90, "week", cal, language)
    assert [text for _, _, _, text in week.header_top()][:2] == months

    month = Timeline(dt.date(2026, 4, 1), 400, "month", cal, language)
    assert [text for _, _, _, text in month.header_top()][0] == years


# ---------------------------------------------------------------- メッセージ
@pytest.mark.parametrize("language,fragment", [
    ("ja", "終了日が開始日より前"), ("en", "end date is before"),
])
def test_errors_are_translated(language, fragment):
    with pytest.raises(SpecError, match=fragment):
        build(dt.date(2026, 4, 1), end=dt.date(2020, 1, 1), language=language)


@pytest.mark.parametrize("language,fragment", [
    ("ja", "Excel として読めません"), ("en", "Cannot read this as Excel"),
])
def test_import_errors_are_translated(tmp_path, language, fragment):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(SpecError, match=fragment):
        read(path, language=language)


# ---------------------------------------------------------------- 読み書きの往復
@pytest.mark.parametrize("written", ["ja", "en"])
@pytest.mark.parametrize("viewed", ["ja", "en"])
def test_a_file_of_either_language_can_be_read_in_either_language(
        tmp_path, make_filled, written, viewed):
    """日本語で作ったファイルを英語で開ける (逆も)。"""
    path = make_filled(f"{written}-{viewed}.xlsx", language=written)
    imported = read(path, language=viewed)

    assert [r.name for r in imported.rows][:2] == ["要件定義", "基本設計"]
    assert imported.rows[0].start == dt.date(2026, 4, 1)
    assert imported.rows[0].member == "設計"
    # 設定シートの期間・稼働曜日も言語をまたいで読める
    assert imported.spec.start == dt.date(2026, 4, 1)
    assert imported.spec.period_days == 365
    assert imported.spec.workdays == ["mon", "tue", "wed", "thu", "fri"]
    # 表示は開いた側の言語になる
    assert imported.spec.language == viewed


@pytest.mark.parametrize("language", ["ja", "en"])
def test_status_colours_work_across_languages(tmp_path, make_filled, language):
    path = make_filled("s.xlsx", language=language, rows=[
        ("開発", "", "1", "A", dt.date(2026, 4, 6), 5, dt.date(2026, 4, 10),
         None, None, None, None, None, "", "設計", "Done"),
        ("開発", "", "2", "B", dt.date(2026, 4, 6), 5, dt.date(2026, 4, 10),
         None, None, None, None, None, "", "設計", "完了"),
    ])
    model = build_chart(read(path, language=language), base_date=dt.date(2026, 4, 8))
    assert [r["status_bg"] for r in model["rows"]] == ["#C0C0C0", "#C0C0C0"]


@pytest.mark.parametrize("language", ["ja", "en"])
def test_export_keeps_the_language(tmp_path, make_filled, language):
    source = make_filled("src.xlsx", language=language)
    imported = read(source, language=language)
    path = tmp_path / "out.xlsx"
    export(imported, path, base_date=dt.date(2026, 6, 10))

    book = openpyxl.load_workbook(path)
    assert book.sheetnames[0] == i18n.labels(language).sheet_plan
    ws = book[book.sheetnames[0]]
    assert ws["E4"].value == i18n.labels(language).columns["name"]
    assert ws["E5"].value == "要件定義"
