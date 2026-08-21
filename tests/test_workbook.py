"""Excel 出力のテスト。

見出しの値・表示形式・配色は添付ファイルから実測したものと突き合わせる。
"""

import datetime as dt

import openpyxl
import pytest

from wbsgen.blank import build
from wbsgen.workbook import SHEET_CONFIG, SHEET_MEMBER, SHEET_PLAN, write


@pytest.fixture
def book(tmp_path):
    spec = build(dt.date(2026, 2, 1), period_days=364, rows=30,
                 title="2026年度 スケジュール", members=["設計", "製造"])
    path = tmp_path / "wbs.xlsx"
    write(spec, path)
    return openpyxl.load_workbook(path)


@pytest.fixture
def plan(book):
    return book[SHEET_PLAN]


# ---------------------------------------------------------------- 構成
def test_sheets(book):
    assert book.sheetnames == [SHEET_PLAN, SHEET_MEMBER, SHEET_CONFIG]


def test_title_and_table_header(plan):
    assert plan["B1"].value == "2026年度 スケジュール"
    assert plan["F3"].value == "予定"
    assert plan["I3"].value == "実績"
    assert [plan[f"{c}4"].value for c in "BCDE"] == ["大項目", "中項目", "項番", "項目"]
    assert [plan[f"{c}4"].value for c in "FGH"] == ["開始日", "日数", "終了日"]
    assert [plan[f"{c}4"].value for c in ("Q", "R")] == ["担当", "状態"]
    assert plan["B4"].fill.fgColor.rgb.endswith("99CCFF")


def test_panes_are_frozen_at_the_chart_origin(plan):
    assert plan.freeze_panes == "S5"
    assert plan.sheet_view.showGridLines is False


# ---------------------------------------------------------------- 見出し
def test_month_row_matches_the_reference_workbook(plan):
    """上段 (行3) は月。元ファイルは S3/W3/AB3/AF3/AK3 にだけ日付が入っている。"""
    assert plan["S3"].number_format == 'm"月"'
    assert [plan[f"{c}3"].value for c in ("S", "W", "AB", "AF", "AK")] == [
        dt.datetime(2026, 2, 1), dt.datetime(2026, 3, 1), dt.datetime(2026, 4, 5),
        dt.datetime(2026, 5, 3), dt.datetime(2026, 6, 7),
    ]
    assert [plan[f"{c}3"].value for c in ("T", "U", "V", "X", "AA")] == [None] * 5
    assert plan["S3"].fill.fgColor.rgb.endswith("FF9900")
    assert plan["S3"].font.bold is True


def test_week_row_matches_the_reference_workbook(plan):
    """下段 (行4) は週の開始日。表示開始日から 7 日刻み。"""
    assert plan["S4"].number_format == "m/d"
    assert [plan[f"{c}4"].value for c in ("S", "T", "U", "V", "W")] == [
        dt.datetime(2026, 2, 1), dt.datetime(2026, 2, 8), dt.datetime(2026, 2, 15),
        dt.datetime(2026, 2, 22), dt.datetime(2026, 3, 1),
    ]
    assert plan["S4"].fill.fgColor.rgb.endswith("FFFF99")


def test_day_unit_adds_a_weekday_row(tmp_path):
    spec = build(dt.date(2026, 4, 1), period_days=10, unit="day", rows=5)
    path = tmp_path / "day.xlsx"
    write(spec, path)
    ws = openpyxl.load_workbook(path)[SHEET_PLAN]

    assert ws["S3"].value == dt.datetime(2026, 4, 1)      # 月
    assert ws["S4"].value == dt.datetime(2026, 4, 1)      # 日
    assert ws["S4"].number_format == "d"
    assert [ws[f"{c}5"].value for c in ("S", "T", "U", "V", "W")] == [
        "水", "木", "金", "土", "日"]
    assert ws["V5"].fill.fgColor.rgb.endswith("EFF3FF")   # 土
    assert ws["W5"].fill.fgColor.rgb.endswith("FFEFEF")   # 日
    assert ws.freeze_panes == "S6"                        # 曜日行のぶん 1 行下がる


def test_month_unit_shows_years_and_months(tmp_path):
    spec = build(dt.date(2026, 11, 1), months=4, unit="month", rows=3)
    path = tmp_path / "month.xlsx"
    write(spec, path)
    ws = openpyxl.load_workbook(path)[SHEET_PLAN]
    assert ws["S3"].number_format == 'yyyy"年"'
    assert ws["S3"].value == dt.datetime(2026, 11, 1)
    assert ws["U3"].value == dt.datetime(2027, 1, 1)      # 年が変わる列
    assert ws["S4"].number_format == 'm"月"'


# ---------------------------------------------------------------- 空行
def test_blank_rows_are_empty_but_formatted(plan):
    # 5 行目から 30 行ぶん
    assert plan.max_row == 34
    for row in (5, 20, 34):
        assert all(plan[f"{c}{row}"].value is None for c in "BCDEFGHIJKLMNOQR")
        assert plan[f"E{row}"].border.left.style == "thin"
        assert plan[f"F{row}"].number_format == "m/dd"
        assert plan[f"G{row}"].number_format == '0\\ "日"'
        assert plan[f"M{row}"].number_format == "0%"
        assert plan[f"F{row}"].fill.fgColor.rgb.endswith("FFFFCC")   # 予定欄の色
        assert plan[f"I{row}"].fill.fgColor.rgb.endswith("FFFFFF")


def test_chart_area_is_ruled(plan):
    assert plan["S5"].border.left.style == "hair"
    assert plan["S34"].border.bottom.style == "hair"


def test_zero_rows_still_produces_a_sheet(tmp_path):
    spec = build(dt.date(2026, 4, 1), months=1, rows=0)
    path = tmp_path / "zero.xlsx"
    write(spec, path)
    ws = openpyxl.load_workbook(path)[SHEET_PLAN]
    assert ws["B4"].value == "大項目"
    assert ws["E5"].value is None


def test_nothing_is_filled_in(plan):
    """中身が入っていないこと (見出しと空行だけ)。"""
    filled = [
        cell.coordinate
        for row in plan.iter_rows(min_row=5)
        for cell in row
        if cell.value not in (None, "")
    ]
    assert filled == []


# ---------------------------------------------------------------- 他シート
def test_member_sheet_lists_the_names(book):
    ws = book[SHEET_MEMBER]
    assert ws["B2"].value == "◆担当者一覧"
    assert [ws["B3"].value, ws["C3"].value] == ["担当", "色"]
    assert [ws[f"B{r}"].value for r in (4, 5)] == ["設計", "製造"]
    assert ws["B6"].value is None                  # 追記用の空行
    assert ws["B6"].border.left.style == "thin"


def test_config_sheet_records_the_period(book):
    ws = book[SHEET_CONFIG]
    values = {ws[f"B{r}"].value: ws[f"C{r}"].value for r in range(3, 20)}
    assert values["チャート表示開始日"] == dt.datetime(2026, 2, 1)
    assert values["チャート表示期間(日)"] == 364
    assert values["チャート表示単位"] == "週単位"
    assert values["記入用の空行"] == 30
    assert values["月曜日"] == "出"
    assert values["土曜日"] == "休"


def test_config_sheet_lists_the_holidays(book):
    ws = book[SHEET_CONFIG]
    assert ws["E2"].value == "◆休日一覧"
    assert ws["E3"].value == dt.datetime(2026, 2, 11)   # 建国記念の日
