"""本日の状況 (日次チェック) のテスト。

読み込んだ WBS を基準日から見て、その日に手を打つべき行を拾えること。
"""

import datetime as dt

import pytest

from wbsgen.daily import build
from wbsgen.importer import Row, resolve
from wbsgen.workcal import WorkCalendar

D = dt.date
TODAY = D(2026, 9, 10)          # 木曜日


def _calendar():
    return WorkCalendar()


def _row(number, name="作業", **kwargs):
    row = Row(row=number, name=name, **kwargs)
    row.written = {key: getattr(row, key) for key in
                   ("days", "end", "actual_days", "actual_end", "progress",
                    "delay", "status")}
    return row


def _build(rows, base=TODAY):
    """基準日で解決してから数える (画面と同じ順序)。"""
    calendar = _calendar()
    resolve(rows, calendar, base)
    return build(rows, calendar, base)


# ---------------------------------------------------------------- 本日の予定
def test_tasks_planned_to_start_today_are_listed():
    rows = [
        _row(2, start=TODAY, end=D(2026, 9, 18)),
        _row(3, start=D(2026, 9, 9), end=D(2026, 9, 18)),
        _row(4, start=TODAY, end=TODAY),
    ]
    assert _build(rows).starting == [2, 4]


def test_tasks_planned_to_end_today_are_listed():
    rows = [
        _row(2, start=D(2026, 9, 1), end=TODAY),
        _row(3, start=D(2026, 9, 1), end=D(2026, 9, 11)),
    ]
    assert _build(rows).ending == [2]


def test_a_summary_row_counts_too():
    """大項目だけの行も数える (画面の集計と件数をそろえるため)。"""
    rows = [_row(2, name="", group="ITa", start=TODAY, end=TODAY)]
    assert _build(rows).starting == [2]


def test_a_row_with_only_a_number_still_counts():
    rows = [_row(2, name="", no="1", start=TODAY, end=TODAY)]
    assert _build(rows).starting == [2]


# ---------------------------------------------------------------- 遅延
def test_delayed_tasks_are_listed_worst_first():
    rows = [
        _row(2, start=D(2026, 9, 8), end=D(2026, 9, 30)),      # 開始遅れ 2 日
        _row(3, start=D(2026, 8, 3), end=D(2026, 8, 31)),      # もっと遅れ
        _row(4, start=TODAY, end=D(2026, 9, 30)),              # 遅れなし
    ]
    digest = _build(rows)
    assert digest.delayed == [3, 2]


def test_a_finished_task_is_not_delayed():
    rows = [_row(2, start=D(2026, 8, 3), end=D(2026, 8, 31),
                 actual_start=D(2026, 8, 3), actual_end=D(2026, 9, 4))]
    assert _build(rows).delayed == []


# ---------------------------------------------------------------- 未着手
def test_tasks_without_any_actual_are_not_started():
    rows = [
        _row(2, start=D(2026, 9, 20), end=D(2026, 9, 25)),
        _row(3, start=D(2026, 9, 1), end=D(2026, 9, 30),
             actual_start=D(2026, 9, 1)),
        _row(4, start=D(2026, 9, 14), end=D(2026, 9, 18)),
    ]
    # 予定の開始日が早い順
    assert _build(rows).not_started == [4, 2]


def test_a_row_with_no_plan_is_not_counted_as_not_started():
    """日付が 1 つも無い行は、着手の話ではなく書き漏れとして拾う。"""
    rows = [_row(2)]
    digest = _build(rows)
    assert digest.not_started == []
    assert _found(digest, "no_dates") == [2]


def test_a_finished_task_is_not_counted_as_not_started():
    """実績の終了日だけ書いてあることがある (開始日を書き忘れた行)。"""
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 3),
                 actual_end=D(2026, 9, 3))]
    assert _build(rows).not_started == []


# ---------------------------------------------------------------- 整合性
def _found(digest, kind):
    for check in digest.checks:
        if check.kind == kind:
            return check.rows
    raise AssertionError(f"{kind} を確かめていない")


def test_every_check_is_reported_even_when_nothing_is_wrong():
    """0 件のチェックも返す (何を確かめたかが画面で判るように)。"""
    digest = _build([_row(2, start=D(2026, 9, 1), end=D(2026, 9, 3),
                          actual_start=D(2026, 9, 1), actual_end=D(2026, 9, 3))])
    assert len(digest.checks) >= 10
    assert all(check.rows == [] for check in digest.checks)


def test_a_plan_that_ends_before_it_starts_is_found():
    rows = [_row(2, start=D(2026, 9, 10), end=D(2026, 9, 3))]
    assert _found(_build(rows), "end_before_start") == [2]


def test_an_actual_that_ends_before_it_starts_is_found():
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 30),
                 actual_start=D(2026, 9, 8), actual_end=D(2026, 9, 3))]
    assert _found(_build(rows), "actual_end_before_start") == [2]


def test_a_plan_with_only_one_date_is_found():
    """終了日が無く、日数からも決まらない行。"""
    rows = [_row(2, end=D(2026, 9, 30))]
    assert _found(_build(rows), "plan_incomplete") == [2]


def test_a_plan_completed_by_the_number_of_days_is_not_flagged():
    rows = [_row(2, start=D(2026, 9, 1), days=5)]
    assert _found(_build(rows), "plan_incomplete") == []


def test_an_actual_end_without_a_start_is_found():
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 3),
                 actual_end=D(2026, 9, 3))]
    assert _found(_build(rows), "actual_end_without_start") == [2]


def test_progress_without_an_actual_start_is_found():
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 30), progress=0.5)]
    assert _found(_build(rows), "progress_without_actual") == [2]


def test_full_progress_without_an_actual_end_is_found():
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 30),
                 actual_start=D(2026, 9, 1), progress=1.0)]
    assert _found(_build(rows), "done_without_actual_end") == [2]


def test_a_written_number_of_days_that_does_not_match_is_found():
    """開始日・終了日から数え直した日数と、書かれた日数が違う行。"""
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 4), days=99)]
    assert _found(_build(rows), "days_rewritten") == [2]


def test_a_matching_number_of_days_is_not_flagged():
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 4), days=4)]
    assert _found(_build(rows), "days_rewritten") == []


def test_an_actual_number_of_days_left_behind_is_found():
    """実績の終了日が無いのに日数だけ残っている (元データの書き間違い)。"""
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 30),
                 actual_start=D(2026, 9, 1), actual_days=3)]
    assert _found(_build(rows), "actual_days_rewritten") == [2]


def test_an_unknown_predecessor_is_found():
    rows = [
        _row(2, no="1", start=D(2026, 9, 1), end=D(2026, 9, 4)),
        _row(3, no="2", start=D(2026, 9, 7), end=D(2026, 9, 9), predecessor="9"),
    ]
    assert _found(_build(rows), "predecessor_missing") == [3]


def test_a_task_starting_before_its_predecessor_ends_is_found():
    rows = [
        _row(2, no="1", start=D(2026, 9, 1), end=D(2026, 9, 18)),
        _row(3, no="2", start=D(2026, 9, 7), end=D(2026, 9, 9), predecessor="1"),
        _row(4, no="3", start=D(2026, 9, 21), end=D(2026, 9, 24), predecessor="1"),
    ]
    digest = _build(rows)
    assert _found(digest, "predecessor_order") == [3]
    assert _found(digest, "predecessor_missing") == []


@pytest.mark.parametrize("text", ["1, 2", "1・2", "1 2", "1、2"])
def test_predecessors_can_be_written_in_several_ways(text):
    rows = [
        _row(2, no="1", start=D(2026, 9, 1), end=D(2026, 9, 4)),
        _row(3, no="2", start=D(2026, 9, 1), end=D(2026, 9, 4)),
        _row(4, no="3", start=D(2026, 9, 7), end=D(2026, 9, 9), predecessor=text),
    ]
    assert _found(_build(rows), "predecessor_missing") == []


# ---------------------------------------------------------------- 担当
def test_owners_without_a_task_today_are_listed():
    rows = [
        _row(2, start=D(2026, 9, 1), end=D(2026, 9, 30), member="高瀬"),
        _row(3, start=D(2026, 8, 1), end=D(2026, 8, 31), member="吉田"),
        _row(4, start=D(2026, 9, 20), end=D(2026, 9, 25), member="鈴木さん"),
    ]
    digest = _build(rows)
    assert digest.members == ["高瀬", "吉田", "鈴木"]
    assert digest.idle_members == ["吉田", "鈴木"]


def test_an_owner_named_in_a_shared_cell_counts_as_assigned():
    """「佐々木（高瀬）」のように 1 つのセルに複数書いてあっても数える。"""
    rows = [
        _row(2, start=D(2026, 9, 1), end=D(2026, 9, 30), member="佐々木（高瀬）"),
        _row(3, start=D(2026, 8, 1), end=D(2026, 8, 31), member="高瀬"),
    ]
    assert _build(rows).idle_members == []


def test_nobody_is_idle_when_no_owner_is_written():
    rows = [_row(2, start=D(2026, 9, 1), end=D(2026, 9, 30))]
    digest = _build(rows)
    assert digest.members == []
    assert digest.idle_members == []


# ---------------------------------------------------------------- 基準日
def test_the_base_date_is_reported():
    assert _build([_row(2, start=TODAY, end=TODAY)]).date == TODAY


def test_another_base_date_changes_what_is_due():
    rows = [_row(2, start=D(2026, 9, 14), end=D(2026, 9, 18))]
    assert _build(rows, base=D(2026, 9, 14)).starting == [2]
    assert _build(rows, base=D(2026, 9, 14)).delayed == []


# ---------------------------------------------------------------- Excel シート
def _sheet(make_filled, tmp_path, rows=None, base=TODAY, name="out.xlsx"):
    """記入済みファイルを書き出して、本日の状況シートを返す。"""
    import openpyxl
    from wbsgen.importer import read
    from wbsgen.workbook import export

    path = make_filled("daily.xlsx", rows=rows if rows is not None else _FILLED)
    out = export(read(path, base_date=base), tmp_path / name, base)
    return openpyxl.load_workbook(out)


#: 基準日 (木曜) から見て、開始・終了・遅れ・未着手が 1 つずつ出る中身
_FILLED = [
    ("開発", "要件", "1", "本日開始", TODAY, None, D(2026, 9, 18),
     None, None, None, None, None, "", "高瀬", ""),
    ("開発", "要件", "2", "本日終了", D(2026, 9, 1), None, TODAY,
     D(2026, 9, 1), None, None, None, None, "", "吉田", ""),
    ("開発", "製造", "3", "遅れている", D(2026, 8, 3), None, D(2026, 8, 31),
     None, None, None, None, None, "", "菊池", ""),
    ("開発", "製造", "4", "これから", D(2026, 10, 1), None, D(2026, 10, 9),
     None, None, None, None, None, "", "虎岩", ""),
]


def test_the_sheet_is_added_to_the_export(make_filled, tmp_path):
    book = _sheet(make_filled, tmp_path)
    assert "本日の状況" in book.sheetnames
    # スケジュール・担当者一覧・設定 の次に置く
    assert book.sheetnames[:4] == ["スケジュール", "担当者一覧", "設定", "本日の状況"]


def test_the_sheet_starts_with_the_base_date(make_filled, tmp_path):
    ws = _sheet(make_filled, tmp_path)["本日の状況"]
    assert ws["B1"].value == "◆本日の状況（基準日 2026/9/10）"
    assert ws["B3"].value == "◆まとめ"
    assert ws["B4"].value == "区分"
    assert ws["F4"].value == "件数"


def test_the_summary_counts_every_section(make_filled, tmp_path):
    ws = _sheet(make_filled, tmp_path)["本日の状況"]
    summary = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=6).value
               for r in range(5, 11)}
    assert summary == {
        "本日開始予定": 1,
        "本日終了予定": 1,
        "遅延タスク": 1,
        "未着手タスク": 3,        # 本日開始・遅れている・これから
        "スケジュール整合性チェック": 0,
        "本日アサインがない担当者": 2,   # 菊池・虎岩
    }


def test_each_section_lists_its_rows(make_filled, tmp_path):
    ws = _sheet(make_filled, tmp_path)["本日の状況"]
    names = {}
    section = None
    for r in range(1, ws.max_row + 1):
        head = ws.cell(row=r, column=2).value
        if isinstance(head, str) and head.startswith("【"):
            section = head.strip("【】")
        elif section and isinstance(head, str) and head not in ("大項目", "該当なし"):
            names.setdefault(section, []).append(ws.cell(row=r, column=5).value)

    assert names["本日開始予定"] == ["本日開始"]
    assert names["本日終了予定"] == ["本日終了"]
    assert names["遅延タスク"] == ["遅れている"]
    assert set(names["未着手タスク"]) == {"本日開始", "遅れている", "これから"}


def test_an_empty_section_says_so(make_filled, tmp_path):
    """該当が無い区分は「該当なし」とだけ書く。"""
    rows = [("開発", "", "1", "先の話", D(2026, 12, 1), None, D(2026, 12, 10),
             None, None, None, None, None, "", "高瀬", "")]
    ws = _sheet(make_filled, tmp_path, rows=rows)["本日の状況"]
    texts = [ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)]
    assert texts.count("該当なし") >= 2          # 本日開始予定・本日終了予定


def test_every_check_is_written_with_its_verdict(make_filled, tmp_path):
    ws = _sheet(make_filled, tmp_path)["本日の状況"]
    verdicts = {}
    for r in range(1, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        verdict = ws.cell(row=r, column=7).value
        if verdict in ("問題なし", "要確認"):
            verdicts[name] = (ws.cell(row=r, column=6).value, verdict)

    assert len(verdicts) == 11
    assert verdicts["予定の終了日が開始日より前"] == (0, "問題なし")


def test_a_contradiction_is_written_with_the_rows(make_filled, tmp_path):
    rows = [("開発", "", "1", "逆さま", D(2026, 9, 18), None, D(2026, 9, 1),
             None, None, None, None, None, "", "高瀬", "")]
    ws = _sheet(make_filled, tmp_path, rows=rows)["本日の状況"]
    found = [r for r in range(1, ws.max_row + 1)
             if ws.cell(row=r, column=2).value == "予定の終了日が開始日より前"]
    assert len(found) == 1
    at = found[0]
    assert (ws.cell(row=at, column=6).value, ws.cell(row=at, column=7).value) \
        == (1, "要確認")
    # すぐ下に見出しと、その行が並ぶ
    assert ws.cell(row=at + 1, column=2).value == "大項目"
    assert ws.cell(row=at + 2, column=5).value == "逆さま"


def test_the_idle_owners_are_listed(make_filled, tmp_path):
    ws = _sheet(make_filled, tmp_path)["本日の状況"]
    start = [r for r in range(1, ws.max_row + 1)
             if ws.cell(row=r, column=2).value == "担当者"][0]
    names = [ws.cell(row=r, column=2).value
             for r in range(start + 1, ws.max_row + 1)]
    assert names == ["菊池", "虎岩"]


def test_the_dates_keep_the_sheet_format(make_filled, tmp_path):
    ws = _sheet(make_filled, tmp_path)["本日の状況"]
    at = [r for r in range(1, ws.max_row + 1)
          if ws.cell(row=r, column=5).value == "本日開始"][0]
    assert ws.cell(row=at, column=7).value == dt.datetime(2026, 9, 10)
    assert ws.cell(row=at, column=7).number_format == "m/dd"
    assert ws.cell(row=at, column=11).number_format == "0%"


def test_the_status_keeps_its_colour(make_filled, tmp_path):
    """状態の色は、スケジュールシートと同じ配色にする。"""
    ws = _sheet(make_filled, tmp_path)["本日の状況"]
    at = [r for r in range(1, ws.max_row + 1)
          if ws.cell(row=r, column=5).value == "遅れている"][0]
    cell = ws.cell(row=at, column=12)
    assert cell.value.startswith("遅れ")
    assert cell.fill.fgColor.rgb.endswith("FF99CC")


def test_the_sheet_is_rebuilt_and_not_doubled(make_filled, tmp_path):
    """書き出したファイルをもう一度読んでも、シートは増えない。"""
    import openpyxl
    from wbsgen.importer import read
    from wbsgen.workbook import export

    path = make_filled("daily.xlsx", rows=_FILLED)
    once = export(read(path, base_date=TODAY), tmp_path / "1.xlsx", TODAY)
    twice = export(read(once, base_date=TODAY), tmp_path / "2.xlsx", TODAY)
    assert openpyxl.load_workbook(twice).sheetnames == \
        openpyxl.load_workbook(once).sheetnames


def test_the_sheet_is_translated(make_filled, tmp_path):
    from wbsgen.importer import read
    from wbsgen.workbook import export

    path = make_filled("en.xlsx", rows=_FILLED)
    imported = read(path, language="en", base_date=TODAY)
    imported.spec.language = "en"
    out = export(imported, tmp_path / "en-out.xlsx", TODAY)

    import openpyxl
    book = openpyxl.load_workbook(out)
    assert "Today" in book.sheetnames
    ws = book["Today"]
    assert ws["B1"].value == "◆Today's status (as of 2026/9/10)"
    assert ws["B4"].value == "Section"
    assert ws.cell(row=5, column=2).value == "Planned to start today"
