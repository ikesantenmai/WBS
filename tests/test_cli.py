import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from wbsgen.cli import main

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


def test_init_writes_a_buildable_template(tmp_path, capsys):
    source = tmp_path / "wbs.yaml"
    assert main(["init", "-o", str(source)]) == 0
    assert source.exists()

    output = tmp_path / "wbs.xlsx"
    assert main(["build", str(source), "-o", str(output)]) == 0
    assert output.exists()
    assert "生成しました" in capsys.readouterr().out


def test_init_refuses_to_overwrite_without_force(tmp_path, capsys):
    source = tmp_path / "wbs.yaml"
    main(["init", "-o", str(source)])
    assert main(["init", "-o", str(source)]) == 2
    assert "すでに存在します" in capsys.readouterr().err
    assert main(["init", "-o", str(source), "--force"]) == 0


@pytest.mark.parametrize("template", ["standard", "minimal"])
def test_every_template_builds(tmp_path, template):
    source = tmp_path / f"{template}.yaml"
    main(["init", "-t", template, "-o", str(source)])
    assert main(["build", str(source), "-o", str(tmp_path / "o.xlsx")]) == 0


def test_build_defaults_output_next_to_the_source(tmp_path):
    source = tmp_path / "plan.yaml"
    main(["init", "-o", str(source)])
    assert main(["build", str(source)]) == 0
    assert (tmp_path / "plan.xlsx").exists()


def test_build_overrides_unit_and_base_date(tmp_path):
    source = tmp_path / "wbs.yaml"
    main(["init", "-o", str(source)])
    output = tmp_path / "day.xlsx"
    assert main(["build", str(source), "-o", str(output),
                 "--unit", "day", "--base-date", "2026-06-01"]) == 0
    ws = openpyxl.load_workbook(output)["スケジュール"]
    # 日単位では 1 列 = 1 日なので、隣接する 2 列のラベルは連続した日付になる
    assert ws["S4"].value == "1"
    assert ws["T4"].value == "2"


def test_check_reports_delays(tmp_path, capsys):
    source = tmp_path / "late.yaml"
    source.write_text(
        "chart:\n  base_date: 2026-08-07\n"
        "tasks:\n"
        "  - {name: 遅れているタスク, start: 2026-04-20, days: 38, actual_start: 2026-05-06}\n"
        "  - {name: 完了したタスク, start: 2026-04-01, days: 5,"
        " actual_start: 2026-04-01, actual_days: 5, progress: 1.0}\n",
        encoding="utf-8",
    )
    assert main(["check", str(source)]) == 0
    out = capsys.readouterr().out
    assert "遅延 1" in out
    assert "遅れ 37 日" in out


def test_missing_file_returns_error_code(capsys):
    assert main(["build", "存在しない.yaml"]) == 2
    assert "見つかりません" in capsys.readouterr().err


def test_invalid_definition_returns_error_code(tmp_path, capsys):
    source = tmp_path / "bad.yaml"
    source.write_text("tasks:\n  - {name: A, progress: 500}\n", encoding="utf-8")
    assert main(["build", str(source)]) == 2
    assert "エラー" in capsys.readouterr().err


# ----------------------------------------------------------------------
def test_reference_example_builds_and_matches_the_source_workbook(tmp_path):
    """添付 WBS を写した例が、元ファイルと同じ日程・状態を再現すること。"""
    source = EXAMPLES / "sbi_web_wbs.yaml"
    output = tmp_path / "sbi.xlsx"
    assert main(["build", str(source), "-o", str(output)]) == 0

    ws = openpyxl.load_workbook(output)["スケジュール"]
    rows = {ws[f"D{r}"].value: r for r in range(5, ws.max_row + 1) if ws[f"D{r}"].value}

    # 元ファイルから実測した値 (項番 -> 予定終了日 / 遅れ / 状態)
    expected = {
        "101": (dt.datetime(2026, 5, 29), None, "完了"),
        "127": (dt.datetime(2026, 6, 16), 37, "遅れ 37 日"),
        "201": (dt.datetime(2026, 5, 29), None, "完了"),
        "161": (dt.datetime(2026, 8, 31), None, "あと 1 日"),
        "702": (dt.datetime(2026, 10, 16), None, "-"),
        "501": (dt.datetime(2027, 1, 22), None, "-"),
    }
    for no, (end, delay, status) in expected.items():
        row = rows[no]
        assert ws[f"H{row}"].value == end, f"項番 {no} の終了日"
        assert ws[f"L{row}"].value == delay, f"項番 {no} の遅れ"
        assert ws[f"R{row}"].value == status, f"項番 {no} の状態"


@pytest.mark.parametrize("name", ["sbi_web_wbs.yaml", "tasks.csv"])
def test_bundled_examples_build(tmp_path, name):
    assert main(["build", str(EXAMPLES / name), "-o", str(tmp_path / "o.xlsx")]) == 0
    assert openpyxl.load_workbook(tmp_path / "o.xlsx").sheetnames[0] == "スケジュール"
