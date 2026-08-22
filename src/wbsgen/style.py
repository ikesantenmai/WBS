"""配色・書式定義。

添付の WBS ファイルから抽出した実際の配色・数値書式をそのまま定数化している。
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT_NAME = "ＭＳ Ｐゴシック"
FONT_SIZE = 10

# ----------------------------------------------------------------------
# 配色 (元ファイルから抽出)
# ----------------------------------------------------------------------
C_HEADER = "99CCFF"        # 見出し行
C_TIMELINE = "FFFF99"      # 日付ヘッダ帯 (下段)
C_MONTH_BAND = "FF9900"    # 月ヘッダ帯 (上段)
C_GROUP = "CCFFCC"         # 大項目・中項目行
C_WHITE = "FFFFFF"
C_TITLE_FONT = "003300"    # タイトル文字
C_PLAN_CELL = "FFFFCC"     # 予定日付セル
C_PLAN_FONT = "000080"     # 予定日付文字
C_HOLIDAY = "FFEFEF"       # 休日列
C_SATURDAY = "EFF3FF"      # 土曜列
C_GRID = "BFBFBF"          # チャート罫線
C_NOWLINE = "E0203C"       # 現在日線

#: 状態表示の書式 (背景色, 文字色, 太字)。添付ファイルから抽出したもの。
#: キーは :func:`wbsgen.i18n.status_kind` が返す種類。
STATUS_STYLES = {
    "done": ("C0C0C0", "808080", False),
    "running": ("FFFF99", "FF9900", False),
    "remaining": ("FFCC00", "800000", False),
    "delayed": ("FF99CC", "FF0000", True),
    "upcoming": ("CCFFFF", "99CCFF", False),
}

#: 担当者色の既定パレット (未指定の担当者に順番に割り当てる)
MEMBER_PALETTE = [
    "4472C4", "ED7D31", "70AD47", "FFC000", "5B9BD5",
    "A5A5A5", "9E480E", "636363", "997300", "264478",
]

# ----------------------------------------------------------------------
# 数値書式 (元ファイルから抽出)
# ----------------------------------------------------------------------
FMT_DATE = "m/dd"
FMT_DAYS = '0\\ "日"'
FMT_PERCENT = "0%"
FMT_EFFORT = "0.0_);[Red]\\(0.0\\)"
FMT_UPDATED = 'yyyy/mm/dd\\ h:mm"更新"'
FMT_MONTH = 'm"月"'
FMT_DAY_AXIS = "m/d"
FMT_TEXT = "@"

# ----------------------------------------------------------------------
# 列レイアウト (幅は px)
# ----------------------------------------------------------------------
COL_MARGIN = "A"
COL_GROUP = "B"
COL_SUBGROUP = "C"
COL_NO = "D"
COL_NAME = "E"
COL_START = "F"
COL_DAYS = "G"
COL_END = "H"
COL_ASTART = "I"
COL_ADAYS = "J"
COL_AEND = "K"
COL_DELAY = "L"
COL_PROGRESS = "M"
COL_EFFORT = "N"
COL_PRED = "O"
COL_SHAPE = "P"
COL_MEMBER = "Q"
COL_STATUS = "R"
COL_CHART_FIRST = 19  # S 列

#: 表側の列幅 (px)。元ファイルの幅 (1/256 文字単位) を px 換算した値。
COLUMN_WIDTHS_PX = {
    "A": 25,
    "B": 110,
    "C": 110,
    "D": 52,
    "E": 220,
    "F": 56,
    "G": 48,
    "H": 56,
    "I": 45,
    "J": 48,
    "K": 45,
    "L": 57,
    "M": 57,
    "N": 47,
    "O": 52,
    "P": 16,
    "Q": 72,
    "R": 76,
}

#: チャート列の幅 (px)。表示単位ごとに切り替える。
CHART_COL_WIDTH_PX = {"day": 22, "week": 45, "month": 62}

ROW_HEIGHT_TITLE = 19.5
ROW_HEIGHT_BAND = 27.75
ROW_HEIGHT_HEADER = 16.5
ROW_HEIGHT_TASK = 16.5
ROW_HEIGHT_SPACER = 6.0

# ----------------------------------------------------------------------
# ヘルパ
# ----------------------------------------------------------------------
MDW = 7.0  # 標準フォントの最大数字幅 (px)


def px_to_width(px: float) -> float:
    """px 幅を Excel の列幅 (文字数) に変換する。"""
    return max((px - 5.0) / MDW, 0.0)


def width_to_px(width: float) -> int:
    """Excel の列幅 (文字数) を px に戻す (``px_to_width`` の逆変換)。"""
    return int(round(width * MDW)) + 5


def font(size: int = FONT_SIZE, bold: bool = False, color: str = "000000",
         italic: bool = False) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, italic=italic)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def thin(color: str = C_GRID) -> Side:
    return Side(style="thin", color=color)


def hair(color: str = C_GRID) -> Side:
    return Side(style="hair", color=color)


BORDER_CELL = Border(left=thin(), right=thin(), top=thin(), bottom=thin())
BORDER_CHART = Border(left=hair(), right=hair(), top=hair(), bottom=hair())

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_NAME = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)


def shade(color: str, factor: float) -> str:
    """``#RRGGBB`` を ``factor`` (0=黒, 1=元色, 2=白寄り) で調整する。"""
    color = color.lstrip("#")
    rgb = [int(color[i:i + 2], 16) for i in (0, 2, 4)]
    if factor <= 1:
        rgb = [int(c * factor) for c in rgb]
    else:
        t = factor - 1
        rgb = [int(c + (255 - c) * t) for c in rgb]
    return "%02X%02X%02X" % tuple(max(0, min(255, c)) for c in rgb)
