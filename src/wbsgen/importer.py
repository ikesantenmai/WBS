"""記入済み Excel の読み込み。

このツールが作った空の WBS に書き込んだものを読み戻す。
見出しの文字で列を探すので、列を足したり順を入れ替えたりしていても読める。
"""

from __future__ import annotations

import datetime as _dt
import unicodedata as _unicodedata
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from .blank import MAX_ROWS, BlankWBS, SpecError
from .i18n import DEFAULT_LANGUAGE, LABELS, labels, message, normalize
from .palette import theme_colors, to_rgb
from .timeline import UNIT_DAY, UNIT_WEEK, VALID_UNITS
from .workcal import WEEKDAY_KEYS

def _build_header_maps():
    """全言語の見出しから、``見出し -> 属性名`` の表を作る。

    「開始」「日数」は予定と実績の両方にあるので、単独では決められない。
    それらは :data:`AMBIGUOUS_HEADERS` に回す。
    """
    plain: Dict[str, str] = {}
    ambiguous: Dict[str, tuple] = {}
    grouped: Dict[tuple, str] = {}

    for text in LABELS.values():
        pairs = [
            (text.group_plan, "start", text.columns["start"]),
            (text.group_plan, "days", text.columns["days"]),
            (text.group_plan, "end", text.columns["end"]),
            (text.group_actual, "actual_start", text.columns["actual_start"]),
            (text.group_actual, "actual_days", text.columns["actual_days"]),
            (text.group_actual, "actual_end", text.columns["actual_end"]),
        ]
        for group, key, label in pairs:
            grouped[(group, label)] = key

        for plan_key, actual_key in (("start", "actual_start"),
                                     ("days", "actual_days"),
                                     ("end", "actual_end")):
            for label in (text.columns[plan_key], text.columns[actual_key]):
                ambiguous[label] = (plan_key, actual_key)

        for key, label in text.columns.items():
            if label not in ambiguous:
                plain.setdefault(label, key)

    return plain, ambiguous, grouped


#: 見出しの文字 -> 属性名 (予定/実績で重ならない列)
HEADERS, AMBIGUOUS_HEADERS, GROUPED_HEADERS = _build_header_maps()

#: 「項目」の列を探すときに見る見出し (言語ごと)
NAME_LABELS = {text.columns["name"] for text in LABELS.values()}

#: 設定シートの表示単位の書き方 -> 内部の値
UNIT_LABELS = {
    label: unit
    for text in LABELS.values()
    for unit, label in text.units.items()
}

#: 設定シートの見出し -> 意味 (言語ごとの書き方をまとめる)
CONFIG_KEYS = {
    "start": {text.config_start for text in LABELS.values()},
    "period": {text.config_period for text in LABELS.values()},
    "unit": {text.config_unit for text in LABELS.values()},
}

#: 稼働日の行 (曜日名 -> weekday 番号) と「出」の書き方
WEEKDAY_LABELS = {
    name: index
    for text in LABELS.values()
    for index, name in enumerate(text.weekday_names)
}
WORK_ON_LABELS = {text.work_on for text in LABELS.values()}

#: シート名の候補 (言語ごと)
PLAN_SHEETS = {text.sheet_plan for text in LABELS.values()}
CONFIG_SHEETS = {text.sheet_config for text in LABELS.values()}
MEMBER_SHEETS = {text.sheet_member for text in LABELS.values()}
MEMBER_SKIP = ({text.member_title for text in LABELS.values()}
               | {text.member_headers[0] for text in LABELS.values()})

#: 見出しを探す範囲
HEADER_SEARCH_ROWS = 20
HEADER_SEARCH_COLS = 40

#: 結合セルを展開する上限 (列まるごとの結合などで膨らまないように)
MAX_MERGED_CELLS = 20000


@dataclass
class Row:
    """記入された 1 行。読めた値をそのまま持つ。"""

    row: int = 0
    group: str = ""
    subgroup: str = ""
    no: str = ""
    name: str = ""
    start: Optional[_dt.date] = None
    days: Optional[int] = None
    end: Optional[_dt.date] = None
    actual_start: Optional[_dt.date] = None
    actual_days: Optional[int] = None
    actual_end: Optional[_dt.date] = None
    delay: Optional[int] = None
    progress: Optional[float] = None
    effort: Optional[float] = None
    predecessor: str = ""
    member: str = ""
    status: str = ""
    #: セルの文字色そのもの (属性名 -> openpyxl の色 / 色なしは None)。
    #: 書き出しではこれをそのまま書き戻す。
    colors: Dict[str, Any] = field(default_factory=dict)
    #: 上を画面用の色に直したもの (属性名 -> ``RRGGBB``)。
    #: テーマ色や色番号も、ここで RGB に直してある。
    ink: Dict[str, str] = field(default_factory=dict)
    #: セルの背景 (属性名 -> openpyxl の塗り / 塗りなしは None)
    fills: Dict[str, Any] = field(default_factory=dict)
    #: 上を画面用の色に直したもの (属性名 -> ``RRGGBB``)
    paper: Dict[str, str] = field(default_factory=dict)
    #: 記入内容から導き出した項目の名前 (画面で薄く見せるため)
    derived: set = field(default_factory=set)
    #: セルに書かれていた値。導出をやり直せるように控えておく。
    written: dict = field(default_factory=dict)

    @property
    def has_plan(self) -> bool:
        return self.start is not None and self.end is not None

    @property
    def has_actual(self) -> bool:
        return self.actual_start is not None

    @property
    def has_actual_record(self) -> bool:
        """実績の日付が 1 つでも入っているか。"""
        return self.actual_start is not None or self.actual_end is not None

    @property
    def is_finished(self) -> bool:
        """実績の終了日が**書かれている**か (完了の判定に使う)。

        実績日数から補った終了日は含めない。
        """
        return self.written.get("actual_end") is not None


@dataclass
class ImportedWBS:
    """読み込んだ WBS 一式。"""

    title: str
    spec: BlankWBS
    rows: List[Row] = field(default_factory=list)
    #: 読めなかったセルの説明 (画面に出して知らせる)
    warnings: List[str] = field(default_factory=list)
    #: 元のファイルの中身。足してあるシートがあるときだけ控える。書き出しは
    #: これを土台にするので、足したシートに手を入れずに済む。
    source: Optional[bytes] = None


# ----------------------------------------------------------------------
def read(source, filename: str = "", language: str = DEFAULT_LANGUAGE,
         base_date: Optional[_dt.date] = None) -> ImportedWBS:
    """``.xlsx`` を読み込む。``source`` はパスかファイルオブジェクト。

    日本語版・英語版のどちらのファイルも読める。``language`` は
    メッセージと、読み込んだあとの表示に使う言語。
    """
    lang = normalize(language)
    try:
        book = load_workbook(source, data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - openpyxl の例外は多岐にわたる
        raise SpecError(message(lang, "not_excel", reason=exc))

    sheet, header_row, columns = _pick_plan_sheet(book, lang)

    config = _read_config(book)
    first_row = header_row + (2 if config.get("unit") == UNIT_DAY else 1)
    rows, warnings = _read_rows(sheet, first_row, columns, lang,
                                theme_colors(book))
    warnings = (_missing_column_warnings(columns, lang)
                + _formula_warnings(source, sheet.title, columns, rows, lang)
                + warnings)

    spec = _build_spec(config, rows, book, filename, lang)
    resolve(rows, spec.calendar(), base_date, lang)
    title = _read_title(sheet, header_row) or spec.title
    spec.title = title
    return ImportedWBS(title=title, spec=spec, rows=rows, warnings=warnings,
                       source=_source_bytes(source) if _has_extra(book) else None)


#: 導出のたびに書かれた値へ戻す項目
WRITTEN_FIELDS = ("start", "end", "days", "actual_start", "actual_end",
                  "actual_days", "progress", "delay", "status")


def resolve(rows: List[Row], calendar, base_date: Optional[_dt.date] = None,
            language: str = DEFAULT_LANGUAGE) -> List[Row]:
    """記入内容から日数・終了日・進捗・状態を導き出す。

    - 予定の日数は、セルの値を捨てて予定の開始日と終了日から数え直す
      (稼働日、両端を含む)
    - 実績の日数も同じく、実績の開始日と終了日から数え直す
    - 実績の終了日が**記入されていれば**、進捗は 100% とみなす
    - 状態は「完了 / 遅れ n 日 / 残り n 日 / あと n 日」を基準日から求める

    予定の終了日が書かれていない行は、代わりに日数から終了日を求める。
    実績では求めず、終了日の無い行の実績日数は空にする。
    開始日と終了日が揃わず数え直せない行も、日数を空にする。

    書かれた値は :attr:`Row.written` に控えてあるので、基準日を変えて
    何度呼んでも同じ結果になる。
    """
    base = base_date or _dt.date.today()
    text = labels(language)

    for row in rows:
        # 前回の導出を巻き戻してから数え直す
        for key in WRITTEN_FIELDS:
            if key in row.written:
                setattr(row, key, row.written[key])
        row.derived = set()

        # --- 予定 ---
        # 終了日が無ければ、書かれた日数から求める (予定はそう書くため)
        _resolve_span(row, calendar, "start", "days", "end", derive_end=True)

        # --- 実績 ---
        finished = row.is_finished
        # 実績は終了日を補わない。終わっていない作業に日数だけ残っている
        # のは元データの書き間違いなので、その日数は捨てる。
        _resolve_span(row, calendar, "actual_start", "actual_days", "actual_end",
                      derive_end=False)

        # --- 進捗 ---
        if finished and row.progress != 1.0:
            row.progress = 1.0
            row.derived.add("progress")

        # --- 遅れと状態 ---
        _resolve_status(row, calendar, base, text, finished)
    return rows


def _resolve_span(row: Row, calendar, start_key: str, days_key: str,
                  end_key: str, derive_end: bool) -> None:
    """開始日・日数・終了日のつじつまを合わせる。

    日数はセルに書かれた値を使わず、**開始日と終了日から数え直す**
    (稼働日、両端を含む)。書かれた日数はいったん無かったものとして扱い、
    数え直せない行では空にする。

    ``derive_end`` は、終了日が無いときに日数から終了日を求めるかどうか。

    - **予定** は求める。「開始日 + n 日」と書くのが普通の書き方で、
      求めた終了日から数え直しても同じ日数になる。
    - **実績** は求めない。終わっていない作業に日数だけ残っているのは
      元データの書き間違いなので、その日数は捨てる。ありもしない
      終了日を作ると、バーも進捗も実際より進んで見えてしまう。
    """
    start = getattr(row, start_key)
    end = getattr(row, end_key)
    days = getattr(row, days_key)

    if start and end:
        _set(row, days_key, calendar.workdays_between(start, end))
    elif derive_end and start and days:
        _set(row, end_key, calendar.end_date(start, days))
    else:
        _set(row, days_key, None)



def _resolve_status(row: Row, calendar, base: _dt.date, text, finished: bool) -> None:
    """遅れの日数と状態を求めて書き込む。

    実績の終了日があれば、予定の日付が無くても完了とする。
    それ以外で予定の日付が無い行は判断できないので、書かれた状態を残す。
    """
    if finished:
        _set(row, "delay", None)
        _set(row, "status", text.status_done)
        return

    if row.start is None and row.end is None:
        return

    delay = _delay_days(row, calendar, base)
    _set(row, "delay", delay)
    if delay:
        _set(row, "status", text.status_delayed.format(days=delay))
    elif row.actual_start and row.end:
        # 着手済み: 予定の終了日まであと何稼働日か
        remaining = max(calendar.workdays_between(base, row.end) - 1, 0)
        _set(row, "status", text.status_remaining.format(days=remaining))
    elif not row.actual_start and row.start:
        # 未着手: 予定の開始日まであと何稼働日か
        upcoming = max(calendar.workdays_between(base, row.start) - 1, 0)
        _set(row, "status", text.status_upcoming.format(days=upcoming))
    else:
        _set(row, "status", text.status_none)


def _delay_days(row: Row, calendar, base: _dt.date) -> Optional[int]:
    """遅れの稼働日数。遅れていなければ ``None``。

    予定開始日を過ぎているのに未着手なら「開始遅れ」、
    予定終了日を過ぎているのに未完了なら「終了遅れ」を数える。
    """
    if row.start and row.start < base and row.actual_start is None:
        delay = calendar.workdays_between(row.start, base) - 1
        if delay > 0:
            return delay
    if row.end and row.end < base:
        delay = calendar.workdays_between(row.end, base) - 1
        if delay > 0:
            return delay
    return None


def _set(row: Row, key: str, value) -> None:
    """導出した値を書き込み、書かれた値と違えば印を付ける。"""
    setattr(row, key, value)
    if row.written.get(key) != value:
        row.derived.add(key)


#: 日数・進捗・状態の計算に効く列 (見つからなければ知らせる)
IMPORTANT_COLUMNS = ("start", "end", "days", "actual_start", "actual_end",
                     "actual_days", "progress")


def _missing_column_warnings(columns: Dict[str, int], language: str) -> List[str]:
    """計算に効く列が見つからなかったら知らせる。

    見出しの文字が違っていると、その列がまるごと空として読まれるため。
    """
    text = labels(language)
    missing = [text.columns[key] for key in IMPORTANT_COLUMNS if key not in columns]
    if not missing:
        return []
    joined = message(language, "list_separator").join(missing)
    return [message(language, "missing_columns", columns=joined)]


#: 数式の知らせは、これだけ出せば充分 (全部並べても読めない)
MAX_FORMULA_WARNINGS = 20


def _formula_warnings(source, sheet_name: str, columns: Dict[str, int],
                      rows: List[Row], language: str) -> List[str]:
    """空に見えるセルが実は数式なら、計算結果が無いことを知らせる。

    Excel は数式の計算結果をファイルに残すが、スクリプトなどで作られた
    ファイルには入っていないことがある。そのまま読むと空に見えるので、
    記入したつもりの日付や進捗が反映されない。

    読み取り専用のシートは ``cell()`` を呼ぶたびに先頭から読み直すので、
    セルごとに引かず、上から 1 回だけ流して必要な行を拾う。
    """
    blanks: Dict[int, List[tuple]] = {}
    for row in rows:
        for key in IMPORTANT_COLUMNS:
            if key in columns and getattr(row, key, None) is None:
                blanks.setdefault(row.row, []).append((columns[key], key))
    if not blanks:
        return []

    formulas = _formula_sheet(source, sheet_name)
    if formulas is None:
        return []

    text = labels(language)
    out: List[str] = []
    first, last = min(blanks), max(blanks)
    stream = formulas.iter_rows(min_row=first, max_row=last, min_col=1,
                                values_only=True)
    for offset, values in enumerate(stream):
        wanted = blanks.get(first + offset)
        if not wanted:
            continue
        for col, key in wanted:
            value = values[col - 1] if col <= len(values) else None
            if isinstance(value, str) and value.startswith("="):
                out.append(message(language, "cell_formula", row=first + offset,
                                   column=text.columns.get(key, key)))
                if len(out) >= MAX_FORMULA_WARNINGS:
                    return out
    return out


def _formula_sheet(source, sheet_name: str):
    """数式を読むためにもう一度開く。開けなければ ``None``。"""
    try:
        if hasattr(source, "seek"):
            source.seek(0)
        book = load_workbook(source, data_only=False, read_only=True)
    except Exception:  # noqa: BLE001 - 診断用なので、失敗したら黙って諦める
        return None
    return book[sheet_name] if sheet_name in book.sheetnames else book.worksheets[0]


def _pick_sheet(book, names):
    """言語ごとのシート名を順に探し、無ければ先頭のシートを使う。"""
    for name in book.sheetnames:
        if name in names:
            return book[name]
    return book.worksheets[0]


def _pick_plan_sheet(book, language: str):
    """日程表のシートと、その見出しの位置を決める。

    まず決まったシート名で探す。名前を変えてあるファイルもあるので、
    見つからなければ「項目」の列があるシートを上から探す
    (メモなどのシートが足してあっても、そちらを掴まないように)。
    """
    named = [book[name] for name in book.sheetnames if name in PLAN_SHEETS]
    found_header = False
    for sheet in named or book.worksheets:
        try:
            header_row, columns = _find_headers(sheet, language)
        except SpecError:
            continue        # 見出しの無いシート (メモや表紙など) は飛ばす
        found_header = True
        if "name" in columns:
            return sheet, header_row, columns
    # 見出しらしい行がどこにも無いのか、あっても「項目」が無いのかで分ける
    raise SpecError(message(language,
                            "no_task_column" if found_header else "no_header"))


def _has_extra(book) -> bool:
    """このツールが使わないシートが足してあるか。"""
    known = PLAN_SHEETS | CONFIG_SHEETS | MEMBER_SHEETS
    return any(name not in known for name in book.sheetnames)


def _source_bytes(source) -> Optional[bytes]:
    """元のファイルの中身をそのまま控える。読めなければ ``None``。

    足してあるシートに手を入れずに書き出すため、書き出しはこの中身を
    土台にして、このツールが使う 3 シートだけを差し替える。
    """
    try:
        if hasattr(source, "seek"):
            source.seek(0)
            return source.read()
        return Path(source).read_bytes()
    except Exception:  # noqa: BLE001 - 控えられなくても読み込みは続ける
        return None


# ----------------------------------------------------------------------
def _find_headers(sheet, language: str) -> "tuple[int, Dict[str, int]]":
    """見出し行を探し、``属性名 -> 列番号`` を返す。"""
    max_row = min(sheet.max_row or 1, HEADER_SEARCH_ROWS)
    max_col = min(sheet.max_column or 1, HEADER_SEARCH_COLS)
    for row in range(1, max_row + 1):
        labels = {}
        for col in range(1, max_col + 1):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, str) and value.strip():
                labels[col] = value.strip()
        if not NAME_LABELS & set(labels.values()):
            continue
        return row, _map_columns(sheet, row, labels)
    raise SpecError(message(language, "no_header"))


def _map_columns(sheet, header_row: int, labels: Dict[int, str]) -> Dict[str, int]:
    """見出しを属性名に対応づける。

    「日数」「開始」などは予定と実績の両方にあるので、1 行上の帯
    (予定 / 実績) と組にして見分ける。帯が無い・ずれている場合は、
    左から出てきた順に「予定 → 実績」とみなす。
    """
    groups = _band_groups(sheet, header_row - 1, max(labels) if labels else 1)
    columns: Dict[str, int] = {}
    seen: Dict[str, int] = {}

    for col, label in sorted(labels.items()):
        key = GROUPED_HEADERS.get((groups.get(col, ""), label))
        if key is None and label in AMBIGUOUS_HEADERS:
            pair = AMBIGUOUS_HEADERS[label]
            index = min(seen.get(pair[0], 0), len(pair) - 1)
            key = pair[index]
        if key is None:
            key = HEADERS.get(label)
        if key is None or key in columns:
            continue
        columns[key] = col
        if label in AMBIGUOUS_HEADERS:
            first = AMBIGUOUS_HEADERS[label][0]
            seen[first] = seen.get(first, 0) + 1
    return columns


def _band_groups(sheet, row: int, max_col: int) -> Dict[int, str]:
    """グループ見出しの行を読み、結合を考慮して列ごとの所属を返す。"""
    if row < 1:
        return {}
    spans = {}
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row:
            value = sheet.cell(row=merged.min_row, column=merged.min_col).value
            if isinstance(value, str) and value.strip():
                for col in range(merged.min_col, merged.max_col + 1):
                    spans[col] = value.strip()
    for col in range(1, max_col + 1):
        if col in spans:
            continue
        value = sheet.cell(row=row, column=col).value
        if isinstance(value, str) and value.strip():
            spans[col] = value.strip()
    return spans


def _read_title(sheet, header_row: int) -> str:
    """見出しより上にある最初の文字列をタイトルとみなす。"""
    for row in range(1, header_row):
        for col in range(1, 6):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, str) and value.strip():
                return value.strip()
    return ""


# ----------------------------------------------------------------------
def _merged_values(sheet) -> Dict[tuple, object]:
    """結合セルの中のどのセルからでも、左上の値を引けるようにする。

    Excel では結合セルの値は左上にしか入っていないが、画面では結合した
    範囲すべてに表示される。手作りの WBS では日付や大項目を縦に結合して
    あることがあり、そのままでは 2 行目以降が空に見えてしまう。
    """
    out: Dict[tuple, object] = {}
    for merged in sheet.merged_cells.ranges:
        rows = merged.max_row - merged.min_row + 1
        cols = merged.max_col - merged.min_col + 1
        if rows * cols > MAX_MERGED_CELLS or len(out) > MAX_MERGED_CELLS:
            continue
        value = sheet.cell(row=merged.min_row, column=merged.min_col).value
        if value is None:
            continue
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                out[(row, col)] = value
    return out


def _read_rows(sheet, first_row: int, columns: Dict[str, int], language: str,
               theme: List[str] = ()):
    """記入された行を読む。空行は飛ばす。

    読めないセルがあっても行は捨てず、その項目だけ空にして注意書きに残す
    (1 か所の書き間違いで行がまるごと消えないようにするため)。
    """
    rows: List[Row] = []
    warnings: List[str] = []
    group = subgroup = ""
    text = labels(language)
    merged = _merged_values(sheet)

    for index in range(first_row, (sheet.max_row or first_row) + 1):
        # 行が空かどうかは、その行自身に書かれた値だけで判断する
        # (結合セルの引き継ぎで、空行が埋まって見えないように)
        own = {key: sheet.cell(row=index, column=col).value
               for key, col in columns.items()}
        if all(_is_blank(v) for v in own.values()):
            continue
        raw = {
            key: (own[key] if not _is_blank(own[key]) else merged.get((index, col)))
            for key, col in columns.items()
        }

        row = Row(row=index)
        row.name = _name_text(raw.get("name"))
        # 大項目・中項目は書かれた行だけに入るので、下の行へ引き継ぐ
        group = _text(raw.get("group")) or group
        subgroup = _text(raw.get("subgroup")) or subgroup
        row.group, row.subgroup = group, subgroup
        row.no = _text(raw.get("no"))
        row.predecessor = _text(raw.get("predecessor"))
        row.member = _text(raw.get("member"))
        row.status = _text(raw.get("status"))

        def cell(key, convert):
            """1 つのセルを読む。読めなければ空にして、どのセルかを控える。"""
            try:
                return convert(raw.get(key))
            except ValueError as exc:
                warnings.append(message(
                    language, "cell_prefix", row=index,
                    column=text.columns.get(key, key),
                    reason=_reason(exc, language)))
                return None

        row.start = cell("start", _date)
        row.end = cell("end", _date)
        row.actual_start = cell("actual_start", _date)
        row.actual_end = cell("actual_end", _date)
        row.days = cell("days", _int)
        row.actual_days = cell("actual_days", _int)
        row.delay = cell("delay", _int)
        row.progress = cell("progress", _ratio)
        row.effort = cell("effort", _number)
        # 導出をやり直せるよう、書かれていた値を控える
        row.written = {key: getattr(row, key) for key in WRITTEN_FIELDS}
        # 文字色と背景は書き出しでそのまま戻すので、ここで控えておく
        row.colors = _cell_colors(sheet, index, columns)
        row.ink = {key: rgb for key, color in row.colors.items()
                   for rgb in [to_rgb(color, theme)] if rgb}
        row.fills = _cell_fills(sheet, index, columns)
        row.paper = {key: rgb for key, fill in row.fills.items()
                     for rgb in [to_rgb(_fill_color(fill), theme)] if rgb}

        # 項目名も予定も無い行でも、実績が入っていれば残す
        # (実績の終了日だけを記録してある行を落とさないため)
        if not row.name and not row.has_plan and not row.has_actual_record:
            continue
        rows.append(row)
        if len(rows) > MAX_ROWS:
            warnings.append(message(language, "too_many_rows", maximum=MAX_ROWS))
            break

    return rows, warnings


#: 文字色を控えない列。状態は文字そのものを計算し直すので、色も計算に合わせる。
COLOR_SKIP = frozenset({"status"})


def _cell_colors(sheet, index: int, columns: Dict[str, int]) -> Dict[str, Any]:
    """その行のセルの文字色を、そのまま控える (大項目〜担当)。

    書き出しでは表を作り直すので、控えておかないと記入した文字色が
    このツールの既定色 (予定は紺、ほかは黒) に置き換わってしまう。

    RGB だけでなくテーマ色・色番号も、openpyxl の色をそのまま持って
    書き戻す。色を指定していないセルは ``None`` を控える (書き出しでも
    色を指定しない = Excel の「自動」のまま)。
    """
    out: Dict[str, Any] = {}
    for key, column in columns.items():
        if key in COLOR_SKIP:
            continue
        color = sheet.cell(row=index, column=column).font.color
        out[key] = copy(color) if color is not None else None
    return out


def _cell_fills(sheet, index: int, columns: Dict[str, int]) -> Dict[str, Any]:
    """その行のセルの背景を、そのまま控える (大項目〜担当)。

    塗っていないセルは ``None`` を控える (書き出しでも塗らない)。
    """
    out: Dict[str, Any] = {}
    for key, column in columns.items():
        if key in COLOR_SKIP:
            continue
        fill = sheet.cell(row=index, column=column).fill
        out[key] = copy(fill) if fill is not None and fill.patternType else None
    return out


def _fill_color(fill):
    """塗りの色 (画面用)。ベタ塗り以外は色を決めない。"""
    if fill is None or fill.patternType != "solid":
        return None
    return fill.fgColor


# ----------------------------------------------------------------------
def _read_config(book) -> Dict[str, Any]:
    """設定シートから表示期間と稼働日を読む。無ければ空。"""
    sheet = next((book[name] for name in book.sheetnames if name in CONFIG_SHEETS), None)
    if sheet is None:
        return {}
    pairs: Dict[str, Any] = {}
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 40),
                              max_col=3, values_only=True):
        label = row[1] if len(row) > 1 else None
        value = row[2] if len(row) > 2 else None
        if isinstance(label, str) and label.strip():
            pairs[label.strip()] = value

    def first(kind):
        for label in CONFIG_KEYS[kind]:
            if label in pairs:
                return pairs[label]
        return None

    config: Dict[str, Any] = {}
    start = _date_or_none(first("start"))
    if start:
        config["start"] = start
    period = first("period")
    if isinstance(period, (int, float)):
        config["period_days"] = int(period)
    unit = first("unit")
    if isinstance(unit, str) and unit.strip() in UNIT_LABELS:
        config["unit"] = UNIT_LABELS[unit.strip()]

    # 稼働曜日 (曜日名は言語ごとに違うので、まとめた表で引く)
    working = set()
    for label, value in pairs.items():
        index = WEEKDAY_LABELS.get(label)
        if index is not None and str(value).strip() in WORK_ON_LABELS:
            working.add(index)
    if working:
        config["workdays"] = [WEEKDAY_KEYS[i] for i in sorted(working)]

    holidays = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 400),
                              min_col=5, max_col=5, values_only=True):
        day = _date_or_none(row[0])
        if day:
            holidays.append(day)
    config["holidays"] = holidays
    return config


def _read_members(book) -> List[str]:
    sheet = next((book[name] for name in book.sheetnames if name in MEMBER_SHEETS), None)
    if sheet is None:
        return []
    names = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 200),
                               min_col=2, max_col=2, values_only=True):
        value = _text(row[0])
        if value and value not in MEMBER_SKIP:
            names.append(value)
    return names


#: 設定シートの期間より後ろの日付があったとき、伸ばしてよい日数。
#: これより先は日付の書き間違いとみなし、設定のままにする
#: (西暦を打ち間違えた 1 行で、日程表が何年ぶんにも伸びないように)。
MAX_EXTEND_DAYS = 366


def _period_days(start: _dt.date, period, days: List[_dt.date]) -> int:
    """日程表に出す日数。**記入された日付が入る長さ**にする。

    設定シートの期間より後ろの日付が書かれていることがある (計画が延びても
    設定を直していないファイル)。そのままでは日程表からはみ出した行の
    バーが描かれず、その作業が消えて見えるので、最後の日付まで伸ばす。
    ただし :data:`MAX_EXTEND_DAYS` より先は伸ばさない。

    終わりは**その月の末日**にそろえる。日単位では月の途中で切れると、
    最終月が尻切れに見えるため。
    """
    latest = max(days) if days else None
    if not period:
        # 設定シートが無いファイル。書かれている日付だけが手がかり。
        return (_end_of_month(latest or start) - start).days + 1

    end = start + _dt.timedelta(days=max(period, 1) - 1)
    if latest and end < latest <= end + _dt.timedelta(days=MAX_EXTEND_DAYS):
        end = latest
    return (_end_of_month(end) - start).days + 1


def _end_of_month(day: _dt.date) -> _dt.date:
    return (_dt.date(day.year + 1, 1, 1) if day.month == 12
            else _dt.date(day.year, day.month + 1, 1)) - _dt.timedelta(days=1)


def _build_spec(config, rows: List[Row], book, filename: str, language: str) -> BlankWBS:
    """表示に使う期間を決める。設定シートが無ければ記入内容から割り出す。"""
    days = [d for row in rows
            for d in (row.start, row.end, row.actual_start, row.actual_end) if d]

    start = config.get("start")
    if start is None:
        start = min(days) if days else _dt.date.today()
        start = start.replace(day=1)
    period = _period_days(start, config.get("period_days"), days)

    title = Path(filename).stem if filename else message(language, "imported_title")
    spec = BlankWBS(
        title=title,
        start=start,
        period_days=min(max(period, 1), 3660),
        unit=config.get("unit") or UNIT_WEEK,
        rows=0,
        members=_read_members(book),
        workdays=config.get("workdays") or ["mon", "tue", "wed", "thu", "fri"],
        # 土日は稼働曜日から、祝日はこのツールが調べて休みにする。
        # 設定シートの「休日一覧」はそこへ足す (会社の休業日が書けるように)。
        # 一覧は書いた時点の期間ぶんしか無く、日程表を伸ばした先や、
        # 手で作ったファイルでは祝日が抜けてしまうため。
        japanese_holidays=True,
        holidays=config.get("holidays") or [],
        language=language,
    )
    if spec.unit not in VALID_UNITS:
        spec.unit = UNIT_WEEK
    return spec


# ----------------------------------------------------------------------
# セルの値の変換
# ----------------------------------------------------------------------
class _CellError(str):
    """セルを読めなかったときの理由。文言は行を組み立てるときに解決する。

    ``str`` を継承しているので、そのまま例外に載せても壊れない。
    """

    def __new__(cls, key: str, value):
        self = super().__new__(cls, f"{key}:{value!r}")
        self.key = key
        self.value = value
        return self


#: 「ここには何も無い」を表す書き方。手書きの表ではよく出るので空として扱う。
BLANK_MARKS = {
    "-", "‐", "‑", "–", "—", "ー", "―", "−", "─", "*", "n/a", "na", "tbd", "?", "？",
    "未", "未定", "未着手", "なし", "無し", "無", "未実施", "―",
}


def _clean(value) -> str:
    """セルの文字列を扱いやすい形にする。

    全角の数字・記号は半角に直す (日本語入力では ３０％ のような値がよく入る)。
    """
    text = _unicodedata.normalize("NFKC", str(value)).strip()
    # 全角スペースは NFKC で半角になるので、まとめて落とす
    return " ".join(text.split())


def _is_blank(value) -> bool:
    """空、または「無し」を表す書き方か。"""
    if value is None:
        return True
    if isinstance(value, str):
        text = _clean(value)
        return not text or text.lower() in BLANK_MARKS
    return False


def _reason(exc: Exception, language: str) -> str:
    """例外から、その言語での理由を組み立てる。"""
    detail = exc.args[0] if exc.args else ""
    if isinstance(detail, _CellError):
        return message(language, detail.key, value=detail.value)
    return str(detail)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _name_text(value) -> str:
    """項目名。書かれた空白をそのまま残す。

    「　WEB口座開設システム」のように、行頭の空白で階層を表す書き方が
    あるため、前後の空白を落とさずに読む (落とすと字下げが消えてしまう)。
    空白だけのセルは、これまでどおり空として扱う。
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value)
    return text if text.strip() else ""


#: 受け付ける日付の書き方
DATE_FORMATS = (
    "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
    "%Y年%m月%d日", "%m月%d日", "%m-%d",
)


def _date(value) -> Optional[_dt.date]:
    if _is_blank(value):
        return None
    day = _date_or_none(value)
    if day is None:
        raise ValueError(_CellError("cell_bad_date", value))
    return day


#: Excel の日付の起点 (1900 年のうるう年の扱いを合わせるため 12/30)
EXCEL_EPOCH = _dt.date(1899, 12, 30)
#: 日付として受け付けるシリアル値の範囲 (1901-01-01 〜 2199-12-31)
SERIAL_RANGE = ((_dt.date(1901, 1, 1) - EXCEL_EPOCH).days,
                (_dt.date(2199, 12, 31) - EXCEL_EPOCH).days)


def _date_or_none(value) -> Optional[_dt.date]:
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    if _is_blank(value):
        return None
    # 書式が「標準」のままだと、日付は数値 (シリアル値) として入っている
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        low, high = SERIAL_RANGE
        if low <= value <= high:
            return EXCEL_EPOCH + _dt.timedelta(days=int(value))
        return None
    text = _clean(value).replace("/", "-").replace(".", "-")
    for fmt in DATE_FORMATS:
        try:
            parsed = _dt.datetime.strptime(text, fmt.replace("/", "-"))
        except ValueError:
            continue
        if "%Y" not in fmt:
            parsed = parsed.replace(year=_dt.date.today().year)
        return parsed.date()
    return None


#: 日数の後ろに付きがちな単位
DAY_SUFFIXES = ("日間", "日", "days", "day", "d")


def _int(value) -> Optional[int]:
    if _is_blank(value):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = _clean(value)
    for suffix in DAY_SUFFIXES:
        if text.lower().endswith(suffix):
            text = text[: -len(suffix)].strip()
            break
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        raise ValueError(_CellError("cell_bad_days", value))


def _number(value) -> Optional[float]:
    if _is_blank(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(_clean(value))
    except ValueError:
        raise ValueError(_CellError("cell_bad_number", value))


def _ratio(value) -> Optional[float]:
    """進捗。``0.8`` でも ``80%`` でも ``80`` でも ``８０％`` でも受ける。"""
    if _is_blank(value):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = _clean(value)
        percent = text.endswith("%")
        try:
            number = float(text.rstrip("%").strip()) / (100.0 if percent else 1.0)
        except ValueError:
            raise ValueError(_CellError("cell_bad_progress", value))
    if number > 1.0:
        number /= 100.0
    return min(max(number, 0.0), 1.0)
