"""Excel の書き出し。

「スケジュール」「担当者一覧」「設定」の 3 シートを作る。

読み込んだファイルに他のシートが足してあるときは、新しく作らず
**元のファイルを土台にして 3 シートだけを差し替える**。足したシートには
一切手を触れないので、そのまま書き出される。

- :func:`write`  … 記入用の空行だけの WBS (中身なし)
- :func:`export` … 読み込んだ WBS を、ガントチャートの図形つきで書き出す

書き出しでは、担当ごとの「要員稼働チェック」を月ごとのシートとして足す。

図形は元の Excel ツールと同じく浮動図形 (DrawingML) として描くので、
週表示・月表示でもバーの端が日付どおりの位置に載る。
"""

from __future__ import annotations

import datetime as _dt
import io
from copy import copy
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import style
from .blank import BlankWBS
from .drawing import Drawing, Geometry, Shape
from .i18n import Labels, labels as get_labels, status_kind
from .inject import inject_drawing, sheet_part
from .timeline import UNIT_DAY, Timeline
from .workload import WorkloadMonth, build as build_workload

#: 既定 (日本語) のシート名。英語版では :class:`~wbsgen.i18n.Labels` の値を使う。
SHEET_PLAN = "スケジュール"
SHEET_MEMBER = "担当者一覧"
SHEET_CONFIG = "設定"

ROW_TITLE = 1
ROW_BAND = 3        # 上段: 月 (月表示のときは年)
ROW_HEADER = 4      # 下段: 週の開始日 (日表示は日、月表示は月) / 表側の見出し

#: 表側の列 (属性名, 列文字, グループ)。見出しの文字は言語ごとに差し替える。
TABLE_COLUMNS = [
    ("group", style.COL_GROUP, None),
    ("subgroup", style.COL_SUBGROUP, None),
    ("no", style.COL_NO, None),
    ("name", style.COL_NAME, None),
    ("start", style.COL_START, "plan"),
    ("days", style.COL_DAYS, "plan"),
    ("end", style.COL_END, "plan"),
    ("actual_start", style.COL_ASTART, "actual"),
    ("actual_days", style.COL_ADAYS, "actual"),
    ("actual_end", style.COL_AEND, "actual"),
    ("delay", style.COL_DELAY, "actual"),
    ("progress", style.COL_PROGRESS, "actual"),
    ("effort", style.COL_EFFORT, None),
    ("predecessor", style.COL_PRED, None),
    ("member", style.COL_MEMBER, None),
    ("status", style.COL_STATUS, None),
]

#: 表側の列文字だけ (書式を敷くのに使う)
TABLE_LETTERS = [letter for _key, letter, _group in TABLE_COLUMNS]

#: 列文字 -> 属性名 (記入されていた文字色を引き当てるのに使う)
COLUMN_KEYS = {letter: key for key, letter, _group in TABLE_COLUMNS}

#: 記入欄に入れておく表示形式 (``date`` / ``days`` は言語ごとに差し替える)
CELL_FORMATS = {
    style.COL_START: "date",
    style.COL_DAYS: "days",
    style.COL_END: "date",
    style.COL_ASTART: "date",
    style.COL_ADAYS: "days",
    style.COL_AEND: "date",
    style.COL_DELAY: "days",
    style.COL_PROGRESS: style.FMT_PERCENT,
    style.COL_EFFORT: style.FMT_EFFORT,
    style.COL_NO: style.FMT_TEXT,
    style.COL_PRED: style.FMT_TEXT,
}

#: 予定欄 (色を変えておく列)
PLAN_COLUMNS = (style.COL_START, style.COL_DAYS, style.COL_END)

MEMBER_SHEET_ROWS = 12

# 要員稼働チェックの行位置 (添付の元表と同じ)
LOAD_ROW_DATE = 4
LOAD_ROW_WEEKDAY = 5
LOAD_ROW_WORKING = 6
LOAD_ROW_FIRST = 7

#: 行の中のバーの縦位置 (0.0=上端, 1.0=下端)
PLAN_TOP_WITH_ACTUAL = 0.12
PLAN_BOTTOM_WITH_ACTUAL = 0.48
ACTUAL_TOP = 0.52
ACTUAL_BOTTOM = 0.88
PLAN_TOP_ALONE = 0.22
PLAN_BOTTOM_ALONE = 0.78


def write(spec: BlankWBS, path) -> Path:
    """``spec`` の空 WBS を ``path`` に書き出す。"""
    return _Writer(spec).save(path)


def export(imported, path, base_date=None) -> Path:
    """読み込んだ WBS を、ガントチャートの図形つきで書き出す。"""
    from .importer import resolve

    spec = imported.spec
    day = base_date or _dt.date.today()
    # 状態と遅れは基準日で決まるので、この日付で数え直してから書く
    resolve(imported.rows, spec.calendar(), day, spec.language)
    writer = _Writer(spec, rows=imported.rows, base_date=day)
    writer.workload = build_workload(imported.rows, spec.calendar())
    return writer.save(path, source=imported.source)


def _md(day: _dt.date, since: _dt.date = None) -> str:
    """タイトルに出す日付。同じ年なら年を省く (2026/9/3〜9/30)。"""
    if since is not None and since.year == day.year:
        return f"{day.month}/{day.day}"
    return f"{day.year}/{day.month}/{day.day}"


def _needs_year(months) -> bool:
    """同じ月が 2 度出てくるなら、シート名に年を入れる。"""
    seen = [m.month for m in months]
    return len(seen) != len(set(seen))


class _Writer:
    def __init__(self, spec: BlankWBS, rows=None, base_date=None):
        self.spec = spec
        self.labels: Labels = get_labels(spec.language)
        self.rows = list(rows or [])
        #: 要員稼働チェック (月ごと)。空なら作らない。
        self.workload: List[WorkloadMonth] = []
        self.base_date = base_date or _dt.date.today()
        self.calendar = spec.calendar()
        self.timeline = Timeline(spec.start, spec.period_days, spec.unit,
                                 self.calendar, spec.language)
        self.chart_col_px = style.CHART_COL_WIDTH_PX[spec.unit]
        # 日単位表示では見出しの下に曜日の行が入る
        self.first_row = ROW_HEADER + (2 if spec.unit == UNIT_DAY else 1)
        self.last_row = self.first_row + max(len(self.rows) + spec.rows, 1) - 1
        self._colors = self._member_colors()

    # ------------------------------------------------------------------
    def save(self, path, source: bytes = None) -> Path:
        """``source`` があれば、その中身を土台にして書き出す。

        **読み込んだときにあったシートは、すべて書き出しにも残す。**
        消すのは、これから同じ名前で作り直すシートだけ。足してある
        シートには触れないので、グラフや画像を含めてそのまま書き出される。
        ``source`` が無ければ、これまでどおり新しく作る。
        """
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = self._base(source)
        titles = self._titles()
        at = self._clear(workbook, titles)
        plan = workbook.create_sheet(titles[0], at)
        self._plan_sheet(plan)
        self._member_sheet(workbook.create_sheet(titles[1], at + 1))
        self._config_sheet(workbook.create_sheet(titles[2], at + 2))
        for offset, month in enumerate(self.workload):
            self._workload_sheet(
                workbook.create_sheet(titles[3 + offset], at + 3 + offset), month)
        workbook.save(path)

        drawing = self._gantt()
        if len(drawing):
            inject_drawing(path, sheet_part(path, plan.title), drawing.to_xml())
        return path

    def _base(self, source: bytes):
        """書き出しの土台になる本を用意する。"""
        if source:
            try:
                # 数式はそのまま残したいので、計算結果ではなく数式を読む
                return load_workbook(io.BytesIO(source), data_only=False)
            except Exception:  # noqa: BLE001 - 開けなければ新しく作る
                pass
        workbook = Workbook()
        workbook.remove(workbook.active)     # 既定の空シートは使わない
        return workbook

    def _titles(self) -> List[str]:
        """これから作るシートの名前を、書き出す順に並べる。"""
        dated = _needs_year(self.workload)
        template = (self.labels.load_sheet_dated if dated
                    else self.labels.load_sheet)
        return [self.labels.sheet_plan, self.labels.sheet_member,
                self.labels.sheet_config] + [
            template.format(year=month.year, month=month.month)
            for month in self.workload]

    def _clear(self, workbook, titles: List[str]) -> int:
        """作り直すシートだけを消し、そこへ置き直す位置を返す。

        **同じ名前で作り直すもの以外は消さない。** 読み込んだときに
        あったシートは、名前も中身もそのまま書き出しに残る。
        """
        doomed = [n for n in workbook.sheetnames if n in titles]
        at = workbook.sheetnames.index(doomed[0]) if doomed else 0
        for name in doomed:
            del workbook[name]
        return min(at, len(workbook.sheetnames))

    # ==================================================================
    # スケジュールシート
    # ==================================================================
    def _plan_sheet(self, ws: Worksheet) -> None:
        self._columns(ws)
        self._title(ws)
        self._table_header(ws)
        self._timeline_header(ws)
        self._task_rows(ws)
        self._blank_rows(ws)
        ws.freeze_panes = f"{get_column_letter(style.COL_CHART_FIRST)}{self.first_row}"
        ws.sheet_view.showGridLines = False

    def _cell_format(self, letter: str):
        """記入欄の表示形式。日付と日数は言語ごとに変わる。"""
        value = CELL_FORMATS.get(letter)
        if value == "date":
            return self.labels.date_format
        if value == "days":
            return self.labels.days_format
        return value

    def _columns(self, ws: Worksheet) -> None:
        for letter, px in style.COLUMN_WIDTHS_PX.items():
            ws.column_dimensions[letter].width = style.px_to_width(px)
        ws.column_dimensions[style.COL_SHAPE].hidden = True
        chart_width = style.px_to_width(self.chart_col_px)
        for i in range(len(self.timeline)):
            ws.column_dimensions[get_column_letter(style.COL_CHART_FIRST + i)].width = chart_width

    def _title(self, ws: Worksheet) -> None:
        ws.row_dimensions[ROW_TITLE].height = style.ROW_HEIGHT_TITLE
        ws.row_dimensions[2].height = 6.0
        cell = ws.cell(row=ROW_TITLE, column=2, value=self.spec.title)
        cell.font = style.font(12, bold=True, color=style.C_TITLE_FONT)
        cell.alignment = style.ALIGN_LEFT

    def _table_header(self, ws: Worksheet) -> None:
        ws.row_dimensions[ROW_BAND].height = style.ROW_HEIGHT_HEADER
        ws.row_dimensions[ROW_HEADER].height = style.ROW_HEIGHT_HEADER

        # 上段はグループ見出し (予定 / 実績) を結合して置く
        for letter in TABLE_LETTERS:
            self._header_style(ws[f"{letter}{ROW_BAND}"])
        self._header_style(ws[f"{style.COL_SHAPE}{ROW_BAND}"])
        for label, first, last in self._column_groups():
            ws.merge_cells(f"{first}{ROW_BAND}:{last}{ROW_BAND}")
            ws[f"{first}{ROW_BAND}"].value = label

        for key, letter, _group in TABLE_COLUMNS:
            cell = ws[f"{letter}{ROW_HEADER}"]
            cell.value = self.labels.columns[key]
            self._header_style(cell)
        self._header_style(ws[f"{style.COL_SHAPE}{ROW_HEADER}"])

    def _column_groups(self):
        """``(グループ見出し, 先頭の列文字, 末尾の列文字)`` を返す。"""
        names = {"plan": self.labels.group_plan, "actual": self.labels.group_actual}
        out = []
        for _key, letter, group in TABLE_COLUMNS:
            if not group:
                continue
            if out and out[-1][0] == group:
                out[-1][2] = letter
            else:
                out.append([group, letter, letter])
        return [(names[group], first, last) for group, first, last in out]

    def _header_style(self, cell) -> None:
        cell.fill = style.fill(style.C_HEADER)
        cell.font = style.font()
        cell.alignment = style.ALIGN_CENTER
        cell.border = style.BORDER_CELL

    # ------------------------------------------------------------------
    def _timeline_header(self, ws: Worksheet) -> None:
        """日程表の見出しを 2 段で書く。

        上段はセルを結合せず、区切りが変わる列にだけ日付を置き、
        ``m"月"`` などの表示形式で見せる (中身は日付のまま)。
        """
        first = style.COL_CHART_FIRST
        top_format, bottom_format = self.timeline.formats

        for col in self.timeline.columns:
            cell = ws.cell(row=ROW_BAND, column=first + col.index)
            cell.fill = style.fill(style.C_MONTH_BAND)
            cell.font = style.font(bold=True)
            cell.alignment = style.ALIGN_CENTER
            cell.number_format = top_format
            cell.border = style.BORDER_CELL
        for index, _span, day, _text in self.timeline.header_top():
            ws.cell(row=ROW_BAND, column=first + index, value=day)

        for index, day, _text in self.timeline.header_bottom():
            col = self.timeline.columns[index]
            cell = ws.cell(row=ROW_HEADER, column=first + index, value=day)
            cell.number_format = bottom_format
            rest = self.timeline.is_rest_column(col)
            cell.fill = style.fill(style.C_HOLIDAY if rest else style.C_TIMELINE)
            cell.font = style.font(9, color="C00000" if rest else "000000")
            cell.alignment = style.ALIGN_CENTER
            cell.border = style.BORDER_CELL

        if self.timeline.unit == UNIT_DAY:
            self._weekday_row(ws, first)

    def _weekday_row(self, ws: Worksheet, first: int) -> None:
        """日単位表示の曜日行 (見出しの直下)。"""
        row = ROW_HEADER + 1
        ws.row_dimensions[row].height = style.ROW_HEIGHT_HEADER
        for column in range(2, first):
            cell = ws.cell(row=row, column=column)
            cell.fill = style.fill(style.C_HEADER)
            cell.border = style.BORDER_CELL
        for col in self.timeline.columns:
            rest = self.timeline.is_rest_column(col)
            sat = self.timeline.is_saturday_column(col)
            cell = ws.cell(row=row, column=first + col.index,
                           value=self._weekday(col))
            cell.number_format = style.FMT_TEXT
            cell.fill = style.fill(
                style.C_SATURDAY if sat else (style.C_HOLIDAY if rest else style.C_TIMELINE))
            cell.font = style.font(
                9, color="0070C0" if sat else ("C00000" if rest else "000000"))
            cell.alignment = style.ALIGN_CENTER
            cell.border = style.BORDER_CELL

    def _weekday(self, col) -> str:
        """曜日の見出し (言語ごと)。"""
        if self.timeline.unit != UNIT_DAY:
            return ""
        return self.labels.weekdays[col.start.weekday()]

    # ------------------------------------------------------------------
    def _blank_rows(self, ws: Worksheet) -> None:
        """記入用の空行。罫線と表示形式だけを入れておく。"""
        first = style.COL_CHART_FIRST
        for row in range(self.first_row + len(self.rows), self.last_row + 1):
            ws.row_dimensions[row].height = style.ROW_HEIGHT_TASK
            self._format_row(ws, row)
            ws[f"{style.COL_NAME}{row}"].alignment = style.ALIGN_NAME
            ws[f"{style.COL_GROUP}{row}"].alignment = style.ALIGN_LEFT
            ws[f"{style.COL_SUBGROUP}{row}"].alignment = style.ALIGN_LEFT
            ws[f"{style.COL_MEMBER}{row}"].alignment = style.ALIGN_LEFT

            for col in self.timeline.columns:
                cell = ws.cell(row=row, column=first + col.index)
                rest = self.timeline.is_rest_column(col)
                sat = self.timeline.is_saturday_column(col)
                cell.fill = style.fill(
                    style.C_SATURDAY if sat else (style.C_HOLIDAY if rest else style.C_WHITE))
                cell.border = style.BORDER_CHART

    # ------------------------------------------------------------------
    def _task_rows(self, ws: Worksheet) -> None:
        """読み込んだ行を書き込む。空行と同じ書式のうえに値を載せる。"""
        first = style.COL_CHART_FIRST
        previous = None
        for index, row in enumerate(self.rows):
            at = self.first_row + index
            ws.row_dimensions[at].height = style.ROW_HEIGHT_TASK
            self._format_row(ws, at, row)
            for col in self.timeline.columns:
                cell = ws.cell(row=at, column=first + col.index)
                rest = self.timeline.is_rest_column(col)
                sat = self.timeline.is_saturday_column(col)
                cell.fill = style.fill(
                    style.C_SATURDAY if sat else (style.C_HOLIDAY if rest else style.C_WHITE))
                cell.border = style.BORDER_CHART

            # 大項目・中項目は変わった行にだけ出す (元ファイルと同じ見え方)
            same = previous is not None and previous.group == row.group
            ws[f"{style.COL_GROUP}{at}"] = None if same else (row.group or None)
            if not (same and previous.subgroup == row.subgroup):
                ws[f"{style.COL_SUBGROUP}{at}"] = row.subgroup or None

            # 実績の終了日は書かれていたものだけ (導出はしない)。日数から
            # 補った値を書き戻すと、読み直したときに「完了」と誤解される。
            for letter, value in (
                (style.COL_NO, row.no), (style.COL_NAME, row.name),
                (style.COL_START, row.start), (style.COL_DAYS, row.days),
                (style.COL_END, row.end),
                (style.COL_ASTART, row.actual_start), (style.COL_ADAYS, row.actual_days),
                (style.COL_AEND, row.actual_end), (style.COL_DELAY, row.delay),
                (style.COL_PROGRESS, row.progress), (style.COL_EFFORT, row.effort),
                (style.COL_PRED, row.predecessor), (style.COL_MEMBER, row.member),
                (style.COL_STATUS, row.status),
            ):
                if value not in (None, ""):
                    ws[f"{letter}{at}"] = value

            if row.status:
                self._status_style(ws[f"{style.COL_STATUS}{at}"], row.status)
            previous = row

    def _format_row(self, ws: Worksheet, row: int, source=None) -> None:
        """1 行ぶんの罫線・色・表示形式を入れる。

        ``source`` があれば、そのセルに指定されていた文字色をそのまま使う
        (読み込んだ色を、書き出しで既定色に変えてしまわないため)。
        """
        written = getattr(source, "colors", None) or {}
        painted = getattr(source, "fills", None) or {}
        for letter in TABLE_LETTERS:
            cell = ws[f"{letter}{row}"]
            cell.border = style.BORDER_CELL
            key = COLUMN_KEYS[letter]
            # 読み込んだセルの文字色と背景は、指定なしも含めてそのまま使う
            if key in painted:
                if painted[key] is not None:
                    cell.fill = copy(painted[key])
            else:
                cell.fill = style.fill(
                    style.C_PLAN_CELL if letter in PLAN_COLUMNS else style.C_WHITE)
            color = (written[key] if key in written else
                     (style.C_PLAN_FONT if letter in PLAN_COLUMNS else "000000"))
            cell.font = style.font(color=color)
            cell.alignment = style.ALIGN_CENTER
            number_format = self._cell_format(letter)
            if number_format:
                cell.number_format = number_format
        ws[f"{style.COL_NAME}{row}"].alignment = style.ALIGN_NAME
        for letter in (style.COL_GROUP, style.COL_SUBGROUP, style.COL_MEMBER):
            ws[f"{letter}{row}"].alignment = style.ALIGN_LEFT

    def _status_style(self, cell, status: str) -> None:
        """状態の色。書かれた言葉から種類を判定する (日英どちらでも効く)。

        この列だけは文字そのものを計算し直すので、色も計算に合わせる。
        大項目〜担当の文字色は記入されたものをそのまま使う。
        """
        kind = status_kind(status)
        if kind is None:
            return
        background, foreground, bold = style.STATUS_STYLES[kind]
        cell.fill = style.fill(background)
        cell.font = style.font(color=foreground, bold=bold)

    # ==================================================================
    # ガントチャートの図形
    # ==================================================================
    def _member_colors(self) -> Dict[str, str]:
        palette = style.MEMBER_PALETTE
        names = list(self.spec.members)
        for row in self.rows:
            if row.member and row.member not in names:
                names.append(row.member)
        return {name: palette[i % len(palette)] for i, name in enumerate(names)}

    def _color(self, member: str) -> str:
        return self._colors.get(member, style.MEMBER_PALETTE[0])

    def _geometry(self) -> Geometry:
        table = style.COLUMN_WIDTHS_PX
        chart_px = self.chart_col_px

        def col_px(col0: int) -> float:
            letter = get_column_letter(col0 + 1)
            return table.get(letter, chart_px)

        def row_pt(row0: int) -> float:
            row = row0 + 1
            if row == ROW_TITLE:
                return style.ROW_HEIGHT_TITLE
            if row == 2:
                return 6.0
            if row in (ROW_BAND, ROW_HEADER) or row == ROW_HEADER + 1:
                return style.ROW_HEIGHT_HEADER
            return style.ROW_HEIGHT_TASK

        return Geometry(col_px, row_pt)

    def _gantt(self) -> Drawing:
        """予定・実績のバーと現在日線を図形として組み立てる。"""
        drawing = Drawing()
        if not self.rows:
            return drawing

        geometry = self._geometry()
        first_col = style.COL_CHART_FIRST - 1

        def rect(left, right, row0, top, bottom, **kwargs):
            if right <= left:
                right = left + 0.06
            drawing.add(Shape(
                frm=geometry.anchor(left, row0, top, first_col),
                to=geometry.anchor(right, row0, bottom, first_col),
                **kwargs))

        for index, row in enumerate(self.rows):
            row0 = self.first_row + index - 1
            color = self._color(row.member)
            plan = self._span(row.start, row.end)
            actual = self._span(row.actual_start, row.actual_end)

            if plan:
                top, bottom = ((PLAN_TOP_WITH_ACTUAL, PLAN_BOTTOM_WITH_ACTUAL) if actual
                               else (PLAN_TOP_ALONE, PLAN_BOTTOM_ALONE))
                rect(plan[0], plan[1], row0, top, bottom,
                     preset="rect", fill=color,
                     line_color=style.shade(color, 0.5),
                     name=f"plan-{row.no or row.name}")
                if row.progress:
                    width = (plan[1] - plan[0]) * min(row.progress, 1.0)
                    rect(plan[0], plan[0] + width, row0, top + 0.06, bottom - 0.06,
                         preset="rect", fill=style.shade(color, 0.55),
                         name=f"progress-{row.no or row.name}")
            if actual:
                done = row.progress is not None and row.progress >= 1.0
                rect(actual[0], actual[1], row0, ACTUAL_TOP, ACTUAL_BOTTOM,
                     preset="rect", fill=style.shade(color, 1.35),
                     line_color=style.shade(color, 0.5),
                     alpha=100000 if done else 65000,
                     name=f"actual-{row.no or row.name}")

        self._now_line(drawing, geometry, first_col)
        return drawing.finalize(geometry)

    def _span(self, start, end):
        if start is None or end is None or end < start:
            return None
        if not self.timeline.overlaps(start, end):
            return None
        return (self.timeline.position(self.timeline.clamp(start)),
                self.timeline.position(self.timeline.clamp(end), end_of_day=True))

    def _now_line(self, drawing: Drawing, geometry: Geometry, first_col: int) -> None:
        day = self.base_date
        if not self.timeline.overlaps(day, day):
            return
        x = self.timeline.position(day)
        drawing.add(Shape(
            frm=geometry.anchor(x, ROW_BAND - 1, 0.0, first_col),
            to=geometry.anchor(x, self.last_row - 1, 1.0, first_col),
            preset="line", fill=None, line_color=style.C_NOWLINE,
            line_width_px=1.25, dash="dash", name="now-line",
        ))

    # ==================================================================
    # 担当者一覧
    # ==================================================================
    def _member_sheet(self, ws: Worksheet) -> None:
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = style.px_to_width(25)
        for letter, px in (("B", 150), ("C", 70), ("D", 220)):
            ws.column_dimensions[letter].width = style.px_to_width(px)

        ws["B2"] = self.labels.member_title
        ws["B2"].font = style.font(11, bold=True, color=style.C_TITLE_FONT)

        for i, label in enumerate(self.labels.member_headers):
            cell = ws.cell(row=3, column=2 + i, value=label)
            self._header_style(cell)

        palette = style.MEMBER_PALETTE
        rows = max(len(self.spec.members), MEMBER_SHEET_ROWS)
        for i in range(rows):
            row = 4 + i
            name = self.spec.members[i] if i < len(self.spec.members) else None
            for column in range(2, 5):
                cell = ws.cell(row=row, column=column)
                cell.border = style.BORDER_CELL
                cell.font = style.font()
                cell.alignment = style.ALIGN_LEFT
            ws.cell(row=row, column=2, value=name)
            if name:
                swatch = ws.cell(row=row, column=3)
                swatch.fill = style.fill(palette[i % len(palette)])

    # ==================================================================
    # 要員稼働チェック
    # ==================================================================
    def _workload_sheet(self, ws: Worksheet, month: WorkloadMonth) -> None:
        """1 か月ぶんの稼働チェックを書く。

        添付いただいた元表と同じ並び。日ごとの件数は数式ではなく
        数え終えた値を入れる (このツールの稼働日の数え方と必ず揃うように)。
        """
        text = self.labels
        first = style.COL_LOAD_FIRST
        last = first + len(month.days) - 1
        total = last + 2                        # 集計欄は 1 列空けた右

        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = style.px_to_width(25)
        ws.column_dimensions["B"].width = style.px_to_width(90)
        for i in range(len(month.days)):
            ws.column_dimensions[get_column_letter(first + i)].width = \
                style.px_to_width(style.LOAD_DAY_WIDTH_PX)
        for i, px in enumerate((72, 96, 64, 380)):
            ws.column_dimensions[get_column_letter(total + i)].width = \
                style.px_to_width(px)

        ws["B1"] = text.load_title.format(
            start=_md(month.start), end=_md(month.end, month.start))
        ws["B1"].font = style.font(12, bold=True, color=style.C_TITLE_FONT)
        ws["B2"] = text.load_note.format(sheet=text.sheet_plan)
        ws["B2"].font = style.font(9, color=style.C_NOTE_FONT)

        self._workload_head(ws, month, first)
        for offset, load in enumerate(month.members):
            self._workload_row(ws, month, load, LOAD_ROW_FIRST + offset, first, total)

        end_row = LOAD_ROW_FIRST + max(len(month.members), 1) - 1
        legend = ws.cell(row=end_row + 2, column=2, value=text.load_legend)
        legend.font = style.font(9, bold=True, color=style.C_NOTE_FONT)
        ws.freeze_panes = f"{style.COL_LOAD_FIRST_LETTER}{LOAD_ROW_FIRST}"
        self._workload_colors(ws, month, first, last, end_row)

    def _workload_head(self, ws: Worksheet, month: WorkloadMonth, first: int) -> None:
        """日付・曜日・稼働判定の 3 行。"""
        text = self.labels
        rows = (
            (LOAD_ROW_DATE, text.load_date,
             [(day, style.DATE_MD) for day in month.days]),
            (LOAD_ROW_WEEKDAY, text.load_weekday,
             [(text.weekdays[day.weekday()], None) for day in month.days]),
            (LOAD_ROW_WORKING, text.load_working,
             [(text.load_on if on else text.load_off, None) for on in month.workdays]),
        )
        for row, label, values in rows:
            head = ws.cell(row=row, column=2, value=label)
            head.font = style.font(10, bold=True)
            head.alignment = style.ALIGN_CENTER
            head.fill = style.fill(style.C_LOAD_HEADER)
            head.border = style.BORDER_CELL
            for i, (value, number_format) in enumerate(values):
                cell = ws.cell(row=row, column=first + i, value=value)
                cell.font = style.font(9)
                cell.alignment = style.ALIGN_CENTER
                cell.fill = style.fill(style.C_LOAD_HEADER)
                cell.border = style.BORDER_CELL
                if number_format:
                    cell.number_format = number_format

    def _workload_row(self, ws: Worksheet, month: WorkloadMonth, load,
                      row: int, first: int, total: int) -> None:
        """1 人ぶんの行と、その右の集計。"""
        text = self.labels
        name = ws.cell(row=row, column=2, value=load.name)
        name.font = style.font(10, bold=True)
        name.alignment = style.ALIGN_CENTER
        name.border = style.BORDER_CELL

        for i, count in enumerate(load.counts):
            cell = ws.cell(row=row, column=first + i, value=count)
            cell.font = style.font(9)
            cell.alignment = style.ALIGN_CENTER
            cell.border = style.BORDER_CELL
            cell.number_format = style.LOAD_COUNT_FORMAT

        free = ", ".join(
            text.load_free_item.format(month=day.month, day=day.day,
                                       weekday=text.weekdays[day.weekday()])
            for day in load.free_days)
        summary = (month.workday_count, load.busy_days,
                   len(load.free_days), free or None)
        for i, value in enumerate(summary):
            cell = ws.cell(row=row, column=total + i, value=value)
            cell.font = style.font(10)
            cell.border = style.BORDER_CELL
            cell.alignment = style.ALIGN_LEFT if i == 3 else style.ALIGN_CENTER
        # 集計の見出しは 1 人目のときに置く
        if row == LOAD_ROW_FIRST:
            heads = (text.load_workdays, text.load_busy,
                     text.load_free, text.load_free_list)
            for i, label in enumerate(heads):
                cell = ws.cell(row=LOAD_ROW_DATE, column=total + i, value=label)
                cell.font = style.font(10, bold=True)
                cell.alignment = style.ALIGN_CENTER
                cell.fill = style.fill(style.C_LOAD_HEADER)
                cell.border = style.BORDER_CELL

    def _workload_colors(self, ws: Worksheet, month: WorkloadMonth,
                         first: int, last: int, end_row: int) -> None:
        """件数に応じた色。値を書き換えても効くよう条件付き書式で入れる。"""
        if not month.members:
            return
        head = get_column_letter(first)
        span = (f"{head}{LOAD_ROW_FIRST}:"
                f"{get_column_letter(last)}{end_row}")
        on, off = self.labels.load_on, self.labels.load_off
        rules = (
            (f'{head}${LOAD_ROW_WORKING}="{off}"', style.C_LOAD_OFF),
            (f'AND({head}${LOAD_ROW_WORKING}="{on}",{head}{LOAD_ROW_FIRST}=0)',
             style.C_LOAD_NONE),
            (f'AND({head}${LOAD_ROW_WORKING}="{on}",{head}{LOAD_ROW_FIRST}>=3)',
             style.C_LOAD_HEAVY),
            (f'AND({head}${LOAD_ROW_WORKING}="{on}",{head}{LOAD_ROW_FIRST}>0)',
             style.C_LOAD_OK),
        )
        for formula, color in rules:
            ws.conditional_formatting.add(span, FormulaRule(
                formula=[formula], fill=style.rule_fill(color), stopIfTrue=True))

        free_column = get_column_letter(last + 4)
        ws.conditional_formatting.add(
            f"{free_column}{LOAD_ROW_FIRST}:{free_column}{end_row}",
            FormulaRule(formula=[f"{free_column}{LOAD_ROW_FIRST}>0"],
                        fill=style.rule_fill(style.C_LOAD_NONE)))

    # ==================================================================
    # 設定
    # ==================================================================
    def _config_sheet(self, ws: Worksheet) -> None:
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = style.px_to_width(25)
        for letter, px in (("B", 170), ("C", 170), ("E", 130)):
            ws.column_dimensions[letter].width = style.px_to_width(px)

        text = self.labels
        entries = [
            (text.config_start, self.timeline.start),
            (text.config_end, self.timeline.end),
            (text.config_period, self.spec.period_days),
            (text.config_unit, text.units[self.spec.unit]),
            (text.config_rows, self.spec.rows),
        ]
        ws["B2"] = text.config_chart
        ws["B2"].font = style.font(11, bold=True, color=style.C_TITLE_FONT)
        row = 3
        for label, value in entries:
            self._config_row(ws, row, label, value)
            row += 1

        row += 1
        ws.cell(row=row, column=2, value=text.config_workdays).font = style.font(
            11, bold=True, color=style.C_TITLE_FONT)
        row += 1
        for index, name in enumerate(text.weekday_names):
            self._config_row(ws, row, name,
                             text.work_on if index in self.calendar.workdays
                             else text.work_off)
            row += 1

        # 休日一覧
        ws.cell(row=2, column=5, value=text.config_holidays).font = style.font(
            11, bold=True, color=style.C_TITLE_FONT)
        holidays = self.calendar.holidays_between(self.timeline.start, self.timeline.end)
        for i, day in enumerate(holidays):
            cell = ws.cell(row=3 + i, column=5, value=day)
            cell.number_format = "yyyy/mm/dd"
            cell.font = style.font()
            cell.alignment = style.ALIGN_CENTER
            cell.border = style.BORDER_CELL

    def _config_row(self, ws: Worksheet, row: int, label: str, value) -> None:
        name = ws.cell(row=row, column=2, value=label)
        name.font = style.font()
        name.alignment = style.ALIGN_LEFT
        name.border = style.BORDER_CELL
        target = ws.cell(row=row, column=3, value=value)
        target.font = style.font()
        target.alignment = style.ALIGN_LEFT
        target.border = style.BORDER_CELL
        if isinstance(value, _dt.date):
            target.number_format = "yyyy/mm/dd"
